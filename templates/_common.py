"""Shared styling helpers for template modules."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.properties import PageSetupProperties
F = "Arial"; INK = "1F2A37"; INPUT = "FFF8DB"; GREY = "F3F4F6"; LINE = "D7DEE3"
thin = Side(style="thin", color=LINE); B = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '#,##0.00;(#,##0.00);"-"'; DATE = 'yyyy-mm-dd'
def font(**k): return Font(name=F, **{"color": INK, **k})
def fill(c): return PatternFill("solid", fgColor=c)
class Kit:
    def __init__(s, main, soft): s.M = main; s.S = soft
    def hdr(s, ws, row, cols, widths=None, col0=1):
        for i, t in enumerate(cols):
            c = ws.cell(row=row, column=col0 + i, value=t)
            c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = fill(s.M)
            c.alignment = Alignment(vertical="center", wrap_text=True); c.border = B
        ws.row_dimensions[row].height = 30
        if widths:
            for i, w in enumerate(widths): ws.column_dimensions[L(col0 + i)].width = w
    def title(s, ws, text, sub):
        ws["A1"] = text; ws["A1"].font = Font(name=F, size=18, bold=True, color=s.M)
        ws["A2"] = sub; ws["A2"].font = Font(name=F, size=10, italic=True, color="5B6770")
        ws.sheet_view.showGridLines = False
    def section(s, ws, ref, text):
        ws[ref] = text; ws[ref].font = Font(name=F, bold=True, size=12, color=s.M)
    def kpi(s, ws, row, col, label, formula, fmt):
        lc = ws.cell(row=row, column=col, value=label); lc.font = Font(name=F, size=9, color="5B6770", bold=True)
        vc = ws.cell(row=row + 1, column=col, value=formula); vc.font = Font(name=F, size=18, bold=True, color=s.M); vc.number_format = fmt
        for rr in (row, row + 1):
            for cc in (col, col + 1): ws.cell(row=rr, column=cc).fill = fill(s.S)
    def start(s, ws, ttl, sub, lines):
        s.title(ws, ttl, sub)
        ws.column_dimensions["A"].width = 4; ws.column_dimensions["B"].width = 100
        ws["B1"] = ws["A1"].value; ws["B1"].font = Font(name=F, size=18, bold=True, color=s.M); ws["A1"] = None
        ws["B2"] = ws["A2"].value; ws["B2"].font = Font(name=F, size=10, italic=True, color="5B6770"); ws["A2"] = None
        for i, (t, b) in enumerate(lines, 4):
            c = ws.cell(row=i, column=2, value=t); c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = Font(name=F, bold=b, size=12 if b else 10, color=s.M if b else INK)
def style_range(ws, r1, r2, c1, c2, inputs=(), fmt=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c); cell.border = B; cell.font = font()
            cell.fill = fill(INPUT if c in inputs else GREY)
            if fmt and c in fmt: cell.number_format = fmt[c]
def input_cell(c, v=None):
    if v is not None: c.value = v
    c.font = Font(name=F, color="0000FF"); c.fill = fill(INPUT); c.border = B
def finish(wb, sample, path):
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    wb.active = 1 if sample else 0
    wb.save(path)
