"""Budget & Expense Planner (sample + blank)."""
import random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter as L

F="Arial"; INK="1F2A37"; ACC="4338CA"; SOFT="EEF2FF"; INPUT="FFF8DB"; GREY="F3F4F6"; LINE="D7DEE3"
thin=Side(style="thin",color=LINE); B=Border(left=thin,right=thin,top=thin,bottom=thin)
MONEY='#,##0.00;(#,##0.00);"-"'; DATE='yyyy-mm-dd'
TR=1500
def font(**k): return Font(name=F, **{"color":INK, **k})
def fill(c): return PatternFill("solid", fgColor=c)
def hdr(ws,row,cols,widths=None,c1=1):
    for i,t in enumerate(cols):
        c=ws.cell(row=row,column=c1+i,value=t); c.font=Font(name=F,bold=True,color="FFFFFF"); c.fill=fill(ACC)
        c.alignment=Alignment(vertical="center",wrap_text=True); c.border=B
    ws.row_dimensions[row].height=30
    if widths:
        for i,w in enumerate(widths): ws.column_dimensions[L(c1+i)].width=w
def title(ws,t,s,col="A"):
    ws[f"{col}1"]=t; ws[f"{col}1"].font=Font(name=F,size=18,bold=True,color=ACC)
    ws[f"{col}2"]=s; ws[f"{col}2"].font=Font(name=F,size=10,italic=True,color="5B6770")
    ws.sheet_view.showGridLines=False
def box(ws,r1,r2,c1,c2,inputs=(),fmt=None):
    for r in range(r1,r2+1):
        for c in range(c1,c2+1):
            x=ws.cell(row=r,column=c); x.border=B; x.font=font(); x.fill=fill(INPUT if c in inputs else GREY)
            if fmt and c in fmt: x.number_format=fmt[c]

INC=["Salary","Side income","Freelance","Interest & dividends","Gifts","Other income"]
EXP=[("Rent / mortgage",1400),("Utilities",180),("Groceries",450),("Dining out",150),("Transport",220),("Insurance",160),("Phone & internet",90),
     ("Subscriptions",45),("Health",80),("Shopping",150),("Entertainment",100),("Kids & family",120),("Debt payments",300),("Savings & investing",500),("Gifts & giving",50),("Other",60)]

def build(sample,path):
    wb=Workbook(); st=wb.active; st.title="Start Here"
    dash=wb.create_sheet("Dashboard"); tx=wb.create_sheet("Transactions"); bud=wb.create_sheet("Monthly Budget")
    yr=wb.create_sheet("Year Overview"); goals=wb.create_sheet("Savings Goals"); debt=wb.create_sheet("Debt Payoff"); s=wb.create_sheet("Settings")

    # Settings
    title(s,"Settings","Set these once. Yellow cells are yours to edit.")
    for i,(k,v) in enumerate([("Currency label","USD"),("Year",2026),("Month to review (1-12)",10)],4):
        s.cell(row=i,column=1,value=k).font=font(bold=True)
        c=s.cell(row=i,column=2,value=v); c.font=Font(name=F,color="0000FF"); c.fill=fill(INPUT); c.border=B
    s["A8"]="Income categories"; s["A8"].font=Font(name=F,bold=True,color=ACC,size=12)
    hdr(s,9,["Income category"])
    for i in range(10):
        c=s.cell(row=10+i,column=1,value=INC[i] if i<len(INC) else None); c.fill=fill(INPUT); c.border=B; c.font=font()
    s["C8"]="Expense categories & monthly budget"; s["C8"].font=Font(name=F,bold=True,color=ACC,size=12)
    hdr(s,9,["Expense category","Monthly budget"],c1=3)
    for i in range(25):
        r=10+i
        if i<len(EXP): s.cell(row=r,column=3,value=EXP[i][0]); s.cell(row=r,column=4,value=EXP[i][1] if sample else None)
    box(s,10,34,3,4,inputs=(3,4),fmt={4:MONEY})
    s["C35"]="Total monthly budget"; s["C35"].font=font(bold=True); s["D35"]="=SUM(D10:D34)"; s["D35"].number_format=MONEY; s["D35"].font=font(bold=True)
    s["F8"]="Accounts"; s["F8"].font=Font(name=F,bold=True,color=ACC,size=12); hdr(s,9,["Account"],c1=6)
    for i,a in enumerate(["Checking","Credit card","Cash","Savings","","",""]):
        c=s.cell(row=10+i,column=6,value=a or None); c.fill=fill(INPUT); c.border=B; c.font=font()
    s["H9"]="Type"; s["H10"]="Income"; s["H11"]="Expense"
    s["A37"]="Helper: all categories"; s["A37"].font=font(italic=True,size=9)
    for i in range(10): s.cell(row=38+i,column=1,value=f'=IF(A{10+i}="","",A{10+i})')
    for i in range(25): s.cell(row=48+i,column=1,value=f'=IF(C{10+i}="","",C{10+i})')
    for c,w in zip("ABCDEFGH",[24,14,24,16,3,16,3,10]): s.column_dimensions[c].width=w
    YR="Settings!$B$5"; MO="Settings!$B$6"

    # Transactions
    title(tx,"Transactions","Log every income and expense. Type, category and account come from drop-down lists.")
    hdr(tx,4,["Date","Type","Category","Description","Amount","Account","Year","Month","Check"],[12,10,22,30,12,14,7,7,16])
    box(tx,5,4+TR,1,9,inputs=range(1,7),fmt={1:DATE,5:MONEY})
    for r in range(5,5+TR):
        tx[f"G{r}"]=f'=IF(A{r}="","",YEAR(A{r}))'; tx[f"H{r}"]=f'=IF(A{r}="","",MONTH(A{r}))'
        tx[f"I{r}"]=(f'=IF(A{r}="","",IF(OR(B{r}="",C{r}="",E{r}=""),"Missing info",IF(AND(B{r}="Income",COUNTIF(Settings!$A$10:$A$19,C{r})=0),"Not an income category",'
                     f'IF(AND(B{r}="Expense",COUNTIF(Settings!$C$10:$C$34,C{r})=0),"Not an expense category","OK"))))')
    for rng,f in [(f"B5:B{4+TR}","=Settings!$H$10:$H$11"),(f"C5:C{4+TR}","=Settings!$A$38:$A$72"),(f"F5:F{4+TR}","=Settings!$F$10:$F$16")]:
        dv=DataValidation(type="list",formula1=f,allow_blank=True); tx.add_data_validation(dv); dv.add(rng)
    tx.conditional_formatting.add(f"I5:I{4+TR}",FormulaRule(formula=['AND(I5<>"",I5<>"OK")'],fill=fill("FEE2E2"),font=Font(name=F,bold=True,color="B91C1C")))
    tx.freeze_panes="A5"
    if sample:
        random.seed(3); r=5; rows=[]
        for m in range(1,11):
            rows.append((dt.date(2026,m,1),"Income","Salary","Monthly salary",4200,"Checking"))
            if m%2==0: rows.append((dt.date(2026,m,15),"Income","Freelance","Design project",random.choice([350,600,800]),"Checking"))
            rows.append((dt.date(2026,m,1),"Expense","Rent / mortgage","Rent",1400,"Checking"))
            rows.append((dt.date(2026,m,5),"Expense","Savings & investing","Transfer to savings",500,"Checking"))
            rows.append((dt.date(2026,m,8),"Expense","Debt payments","Car loan",300,"Checking"))
            rows.append((dt.date(2026,m,10),"Expense","Utilities","Electricity & water",random.randint(140,230),"Checking"))
            rows.append((dt.date(2026,m,12),"Expense","Phone & internet","Phone + fiber",90,"Credit card"))
            rows.append((dt.date(2026,m,3),"Expense","Insurance","Car + health",160,"Checking"))
            rows.append((dt.date(2026,m,2),"Expense","Subscriptions","Streaming & apps",random.choice([38,45,52]),"Credit card"))
            for _ in range(5): rows.append((dt.date(2026,m,random.randint(1,28)),"Expense","Groceries","Supermarket",random.randint(60,120),"Credit card"))
            for _ in range(3): rows.append((dt.date(2026,m,random.randint(1,28)),"Expense","Dining out",random.choice(["Pizza night","Lunch","Cafe"]),random.randint(18,70),"Credit card"))
            for _ in range(2): rows.append((dt.date(2026,m,random.randint(1,28)),"Expense","Transport",random.choice(["Fuel","Train pass"]),random.randint(60,120),"Credit card"))
            rows.append((dt.date(2026,m,random.randint(1,28)),"Expense","Shopping",random.choice(["Clothes","Home items","Shoes"]),random.randint(40,220),"Credit card"))
            rows.append((dt.date(2026,m,random.randint(1,28)),"Expense","Entertainment",random.choice(["Cinema","Concert","Games"]),random.randint(20,140),"Credit card"))
            if m in (3,7,10): rows.append((dt.date(2026,m,random.randint(1,28)),"Expense","Health","Dentist",random.randint(80,240),"Credit card"))
        rows.sort()
        for row in rows:
            for j,v in enumerate(row,1): tx.cell(row=r,column=j,value=v)
            r+=1

    T=lambda col: f"Transactions!${col}$5:${col}${4+TR}"
    # Monthly Budget
    title(bud,"Monthly Budget vs Actual","Shows the month chosen on the Settings tab.")
    bud["A3"]='="Month: "&TEXT(DATE(Settings!B5,Settings!B6,1),"mmmm yyyy")&"  ·  Amounts in "&Settings!B4'; bud["A3"].font=font(bold=True)
    hdr(bud,5,["Expense category","Budget","Actual","Difference","% used"],[24,14,14,14,10])
    for i in range(25):
        r=6+i; sr=10+i
        bud[f"A{r}"]=f'=IF(Settings!C{sr}="","",Settings!C{sr})'
        bud[f"B{r}"]=f'=IF(A{r}="","",N(Settings!D{sr}))'
        bud[f"C{r}"]=f'=IF(A{r}="","",SUMIFS({T("E")},{T("B")},"Expense",{T("C")},A{r},{T("G")},{YR},{T("H")},{MO}))'
        bud[f"D{r}"]=f'=IF(A{r}="","",B{r}-C{r})'
        bud[f"E{r}"]=f'=IF(OR(A{r}="",N(B{r})=0),"",C{r}/B{r})'
    box(bud,6,30,1,5,fmt={2:MONEY,3:MONEY,4:MONEY,5:'0%'})
    bud["A31"]="Total"; 
    for c in "BCD": bud[f"{c}31"]=f"=SUM({c}6:{c}30)"; bud[f"{c}31"].number_format=MONEY
    bud["E31"]='=IF(B31=0,"",C31/B31)'; bud["E31"].number_format='0%'
    for c in "ABCDE": bud[f"{c}31"].font=font(bold=True); bud[f"{c}31"].border=B; bud[f"{c}31"].fill=fill(SOFT)
    bud.conditional_formatting.add("D6:D30",CellIsRule(operator="lessThan",formula=["0"],font=Font(name=F,bold=True,color="B91C1C"),fill=fill("FEE2E2")))
    bud.conditional_formatting.add("E6:E30",DataBarRule(start_type="num",start_value=0,end_type="num",end_value=1,color="818CF8"))
    bud["G5"]="Income this month"; bud["G5"].font=Font(name=F,bold=True,color=ACC,size=12)
    hdr(bud,6,["Income category","Received"],[24,14],c1=7)
    for i in range(10):
        r=7+i; bud[f"G{r}"]=f'=IF(Settings!A{10+i}="","",Settings!A{10+i})'
        bud[f"H{r}"]=f'=IF(G{r}="","",SUMIFS({T("E")},{T("B")},"Income",{T("C")},G{r},{T("G")},{YR},{T("H")},{MO}))'
    box(bud,7,16,7,8,fmt={8:MONEY})
    bud["G17"]="Total income"; bud["H17"]="=SUM(H7:H16)"; bud["G18"]="Total spent"; bud["H18"]="=C31"; bud["G19"]="Left over"; bud["H19"]="=H17-H18"
    for r in (17,18,19): bud[f"G{r}"].font=font(bold=True); bud[f"H{r}"].font=font(bold=True); bud[f"H{r}"].number_format=MONEY; bud[f"H{r}"].border=B
    bud["G21"]="Savings rate"; bud["G21"].font=font(bold=True)
    bud["H21"]='=IF(H17=0,"",(H19+SUMIFS(C6:C30,A6:A30,"Savings & investing"))/H17)'; bud["H21"].number_format='0%'; bud["H21"].font=font(bold=True)
    bud["G22"]="(left over plus money moved to 'Savings & investing', divided by income)"; bud["G22"].font=font(italic=True,size=8)

    # Year Overview
    title(yr,"Year Overview","Every category, every month of the year chosen on the Settings tab.")
    months=["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
    hdr(yr,4,["Category"]+months+["Total","Avg / month"],[24]+[10]*12+[12,12])
    yr["A5"]="INCOME"; yr["A5"].font=Font(name=F,bold=True,color=ACC)
    def yrow(r,label_formula,typ):
        yr[f"A{r}"]=label_formula
        for m in range(1,13):
            yr.cell(row=r,column=1+m,value=f'=IF($A{r}="","",SUMIFS({T("E")},{T("B")},"{typ}",{T("C")},$A{r},{T("G")},{YR},{T("H")},{m}))')
        yr[f"N{r}"]=f'=IF($A{r}="","",SUM(B{r}:M{r}))'; yr[f"O{r}"]=f'=IF($A{r}="","",N{r}/MAX(1,COUNT(Dashboard!$B$15:$B$26)))'
    for i in range(10): yrow(6+i,f'=IF(Settings!A{10+i}="","",Settings!A{10+i})',"Income")
    yr["A16"]="Total income"
    yr["A17"]="EXPENSES"; yr["A17"].font=Font(name=F,bold=True,color=ACC)
    for i in range(25): yrow(18+i,f'=IF(Settings!C{10+i}="","",Settings!C{10+i})',"Expense")
    yr["A43"]="Total expenses"; yr["A44"]="Net (income − expenses)"
    for m in range(1,15):
        col=L(1+m)
        yr[f"{col}16"]=f"=SUM({col}6:{col}15)"; yr[f"{col}43"]=f"=SUM({col}18:{col}42)"; yr[f"{col}44"]=f"={col}16-{col}43"
    for r in list(range(6,17))+list(range(18,45)):
        for c in range(1,16):
            x=yr.cell(row=r,column=c); x.border=B; x.font=font(bold=(r in (16,43,44)))
            if c>1: x.number_format=MONEY
            if r in (16,43,44): x.fill=fill(SOFT)
    yr.freeze_panes="B5"

    # Savings Goals
    title(goals,"Savings Goals","Track what you are saving for. Update 'Saved so far' whenever you add money.")
    hdr(goals,4,["Goal","Target amount","Saved so far","Target date","Left to save","Progress","Months left","Needed per month"],[24,14,14,12,14,10,10,16])
    box(goals,5,19,1,8,inputs=(1,2,3,4),fmt={2:MONEY,3:MONEY,4:DATE,5:MONEY,6:'0%',8:MONEY})
    if sample:
        for i,g in enumerate([("Emergency fund",10000,6200,dt.date(2027,6,30)),("Summer holiday",3000,1150,dt.date(2027,5,31)),("New laptop",1800,900,dt.date(2027,1,31))]):
            for j,v in enumerate(g,1): goals.cell(row=5+i,column=j,value=v)
    for r in range(5,20):
        goals[f"E{r}"]=f'=IF(A{r}="","",MAX(0,B{r}-C{r}))'
        goals[f"F{r}"]=f'=IF(OR(A{r}="",N(B{r})=0),"",MIN(1,C{r}/B{r}))'
        goals[f"G{r}"]=f'=IF(OR(A{r}="",D{r}=""),"",MAX(0,(YEAR(D{r})-YEAR(TODAY()))*12+MONTH(D{r})-MONTH(TODAY())))'
        goals[f"H{r}"]=f'=IF(OR(A{r}="",G{r}=""),"",IF(G{r}=0,E{r},E{r}/G{r}))'
    goals.conditional_formatting.add("F5:F19",DataBarRule(start_type="num",start_value=0,end_type="num",end_value=1,color="34D399"))

    # Debt Payoff
    title(debt,"Debt Payoff","List your debts. Months to pay off uses your monthly payment and the yearly interest rate.")
    hdr(debt,4,["Debt","Balance","Interest rate (yearly)","Monthly payment","Months to pay off","Payoff date","Total interest","Status"],[22,14,12,14,12,13,14,22])
    box(debt,5,16,1,8,inputs=(1,2,3,4),fmt={2:MONEY,3:'0.0%',4:MONEY,6:'mmm yyyy',7:MONEY})
    if sample:
        for i,d in enumerate([("Car loan",8400,0.069,300),("Credit card",2100,0.229,150),("Student loan",14500,0.045,180)]):
            for j,v in enumerate(d,1): debt.cell(row=5+i,column=j,value=v)
    for r in range(5,17):
        ok=f'AND(N(B{r})>0,N(D{r})>B{r}*N(C{r})/12)'
        debt[f"E{r}"]=f'=IF(A{r}="","",IF({ok},ROUNDUP(IF(N(C{r})=0,B{r}/D{r},NPER(C{r}/12,-D{r},B{r})),0),""))'
        debt[f"F{r}"]=f'=IF(OR(A{r}="",E{r}=""),"",EDATE(TODAY(),E{r}))'
        debt[f"G{r}"]=f'=IF(OR(A{r}="",E{r}=""),"",MAX(0,E{r}*D{r}-B{r}))'
        debt[f"H{r}"]=f'=IF(A{r}="","",IF(N(B{r})<=0,"Paid off",IF({ok},"On track","Payment too low")))'
    debt.conditional_formatting.add("H5:H16",CellIsRule(operator="equal",formula=['"Payment too low"'],fill=fill("FEE2E2"),font=Font(name=F,bold=True,color="B91C1C")))
    debt["A18"]="Total"; debt["B18"]="=SUM(B5:B16)"; debt["D18"]="=SUM(D5:D16)"; debt["G18"]="=SUM(G5:G16)"
    for c in "ABDG": debt[f"{c}18"].font=font(bold=True); debt[f"{c}18"].number_format=MONEY if c!="A" else "General"
    debt["A20"]="Total interest is an estimate: it assumes the same payment and rate every month."; debt["A20"].font=font(italic=True,size=9)

    # Dashboard
    title(dash,"Budget Dashboard","Updates by itself from your Transactions. Change the year and month on the Settings tab.")
    dash["A3"]='="Year "&Settings!B5&"  ·  Month "&TEXT(DATE(Settings!B5,Settings!B6,1),"mmmm")&"  ·  Amounts in "&Settings!B4'; dash["A3"].font=font(bold=True)
    kp=[("Income (this month)","='Monthly Budget'!H17",MONEY),("Spent (this month)","='Monthly Budget'!H18",MONEY),("Left over (this month)","='Monthly Budget'!H19",MONEY),
        ("Budget used (this month)","='Monthly Budget'!E31",'0%'),("Savings rate (this month)","='Monthly Budget'!H21",'0%'),("Net saved (year)","='Year Overview'!N44",MONEY)]
    for i,(lab,f,fmt) in enumerate(kp):
        col=1+(i%3)*2; row=5+(i//3)*3
        a=dash.cell(row=row,column=col,value=lab); a.font=Font(name=F,size=9,bold=True,color="5B6770")
        v=dash.cell(row=row+1,column=col,value=f); v.font=Font(name=F,size=18,bold=True,color=ACC); v.number_format=fmt
        for rr in (row,row+1):
            for cc in (col,col+1): dash.cell(row=rr,column=cc).fill=fill(SOFT)
    for c,w in zip("ABCDEFGHIJ",[17,14,16,13,18,13,4,24,14,4]): dash.column_dimensions[c].width=w
    dash["A13"]="Month by month"; dash["A13"].font=Font(name=F,bold=True,size=12,color=ACC)
    for i,h in enumerate(["Month","Income","Expenses","Net"],1):
        c=dash.cell(row=14,column=i,value=h); c.font=Font(name=F,bold=True,color="FFFFFF"); c.fill=fill(ACC); c.border=B
    for m in range(1,13):
        r=14+m; col=L(1+m)
        dash[f"A{r}"]=months[m-1]
        dash[f"B{r}"]=f"=IF(AND('Year Overview'!{col}16=0,'Year Overview'!{col}43=0),\"\",'Year Overview'!{col}16)"
        dash[f"C{r}"]=f"=IF(B{r}=\"\",\"\",'Year Overview'!{col}43)"
        dash[f"D{r}"]=f'=IF(B{r}="","",B{r}-C{r})'
        for c in "ABCD":
            dash[f"{c}{r}"].border=B; dash[f"{c}{r}"].font=font()
            if c!="A": dash[f"{c}{r}"].number_format=MONEY
    dash.conditional_formatting.add("D15:D26",CellIsRule(operator="lessThan",formula=["0"],font=Font(name=F,bold=True,color="B91C1C")))
    dash["H13"]="Spending by category (year)"; dash["H13"].font=Font(name=F,bold=True,size=12,color=ACC)
    for i,h in enumerate(["Category","Spent"]):
        c=dash.cell(row=14,column=8+i,value=h); c.font=Font(name=F,bold=True,color="FFFFFF"); c.fill=fill(ACC); c.border=B
    for i in range(25):
        r=15+i
        dash[f"H{r}"]=f"='Year Overview'!A{18+i}"; dash[f"I{r}"]=f"=IF(H{r}=\"\",\"\",'Year Overview'!N{18+i})"
        dash[f"H{r}"].font=font(); dash[f"I{r}"].font=font(); dash[f"I{r}"].number_format=MONEY; dash[f"H{r}"].border=B; dash[f"I{r}"].border=B
    ch=BarChart(); ch.type="col"; ch.title="Income vs expenses"; ch.height=7.5; ch.width=15
    ch.add_data(Reference(dash,min_col=2,min_row=14,max_col=3,max_row=26),titles_from_data=True); ch.set_categories(Reference(dash,min_col=1,min_row=15,max_row=26))
    ch.y_axis.majorGridlines=None
    for ser,col in zip(ch.series,("4338CA","F59E0B")): ser.graphicalProperties=GraphicalProperties(solidFill=col); ser.graphicalProperties.line.solidFill=col
    dash.add_chart(ch,"A29")
    pc=PieChart(); pc.title="Where the money went"; pc.height=9; pc.width=13
    pc.add_data(Reference(dash,min_col=9,min_row=14,max_row=30),titles_from_data=True); pc.set_categories(Reference(dash,min_col=8,min_row=15,max_row=30))
    dash.add_chart(pc,"K13")

    # Start Here
    title(st,"Budget & Expense Planner","Plan your money, track every expense and see where it goes. Works in Excel and Google Sheets.",col="B")
    st.column_dimensions["A"].width=3; st.column_dimensions["B"].width=100
    lines=[("How to start",True),
      ("1.  Settings: choose your currency label, the year and the month to review. Edit the income and expense categories and type a monthly budget for each expense.",False),
      ("2.  Transactions: add every income and expense. Pick the type, category and account from the lists. The Check column warns you about missing or mismatched info.",False),
      ("3.  Monthly Budget: see budget vs actual for the month you picked, what's left over and your savings rate. Overspent categories turn red.",False),
      ("4.  Year Overview and Dashboard update by themselves: totals per month, spending by category and charts.",False),
      ("5.  Savings Goals and Debt Payoff: list your goals and debts to see the monthly amount needed and when each debt will be paid off.",False),
      ("",False),("Colour key",True),
      ("Yellow cells are inputs: type here.  Grey cells calculate themselves: please don't overwrite them.",False),
      ("",False),("Good to know",True),
      ("•  Room for 1,500 transactions, 10 income categories, 25 expense categories, 15 goals and 12 debts.",False),
      ("•  New year? Change the year on the Settings tab: all reports follow it. Keep adding transactions in the same list.",False),
      ("•  To use Google Sheets: open drive.google.com, upload this file, then choose Open with > Google Sheets.",False),
      ("•  Card payments: log them when you spend, not when you pay the card bill, so categories stay accurate.",False)]
    if sample: lines+=[("",False),("This copy contains sample data",True),("All amounts are made-up examples. Use the blank copy in your download for your own budget.",False)]
    for i,(t,b) in enumerate(lines,4):
        c=st.cell(row=i,column=2,value=t); c.alignment=Alignment(wrap_text=True,vertical="top"); c.font=Font(name=F,bold=b,size=12 if b else 10,color=ACC if b else INK)
    for ws in wb.worksheets:
        ws.page_setup.orientation="landscape"; ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    wb.active=1 if sample else 0
    wb.save(path)

META = {
    "name": "Budget & Expense Planner – Excel & Google Sheets Template",
    "file_base": "Budget-Expense-Planner",
    "price": 19,
    "color": "#4338CA",
    "title": "Budget & Expense Planner",
    "tagline": "Budget vs actual · savings goals · debt payoff — Excel & Google Sheets",
    "intro": "Know exactly where your money goes. Set a budget for every category, log your spending in seconds and watch the dashboard update by itself.",
    "features": [
        "Dashboard: income, spending, money left over, budget used and savings rate, with charts by month and by category",
        "Monthly budget vs actual for 25 expense categories – overspent categories turn red",
        "Transactions log with drop-down lists and a built-in check for typos and missing info",
        "Year overview: every category, every month, with totals and monthly averages",
        "Savings goals: progress bars and how much to save each month to hit your date",
        "Debt payoff: months to pay off, payoff date and estimated interest for each debt",
        "Any currency, room for 1,500 transactions",
    ],
    "tags": ["budget planner", "budget spreadsheet", "expense tracker", "google sheets budget", "excel budget", "monthly budget", "savings tracker",
             "debt payoff", "finance planner", "personal finance", "budget template", "money tracker", "paycheck budget"],
}
