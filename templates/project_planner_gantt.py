"""Project Planner with Gantt Chart (sample + blank versions).

Gantt = task table (phase, task, owner, start or 'depends on', duration, % done) on the left and a
60-column timeline on the right drawn with conditional formatting only (days or weeks, chosen on Settings).
A task without a start date starts the day after the task it depends on (or on the project start).
The Dashboard shows progress vs plan, late tasks, phases, owners and what is due next.
"""
import random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.utils import get_column_letter as L
from _common import Kit, F, INK, B, DATE, font, fill, style_range, input_cell, finish

MAIN = "1D4ED8"; SOFT = "E8EEFC"; BAR = "93C5FD"; DONE = "1D4ED8"; LATE = "FCA5A5"; MILE = "F59E0B"
K = Kit(MAIN, SOFT)
NT = 300
TF, TL = 5, 4 + NT
NCOL = 60
TL0 = 18                       # first timeline column (R)
G = "Gantt"
STATUSES = ["Done", "Late", "In progress", "Not started", "No date"]


def t_(col): return f"{G}!${col}${TF}:${col}${TL}"


def build(sample, path):
    wb = Workbook()
    st = wb.active; st.title = "Start Here"
    dash = wb.create_sheet("Dashboard"); g = wb.create_sheet(G); s = wb.create_sheet("Settings")
    today = dt.date.today()
    pstart = today - dt.timedelta(days=today.weekday()) - dt.timedelta(weeks=4)

    # ---------- Settings ----------
    K.title(s, "Settings", "Yellow cells are yours to edit.")
    rows = [("Project name", "Website Relaunch" if sample else None, None),
            ("Project manager", "Anna" if sample else None, None),
            ("Project start date", pstart if sample else "=TODAY()", DATE),
            ("Timeline scale", "Days", None),
            ("Timeline starts on (optional)", None, DATE)]
    for i, (k_, v, fmt) in enumerate(rows, 4):
        s.cell(row=i, column=1, value=k_).font = font(bold=True)
        c = s.cell(row=i, column=2); input_cell(c, v)
        if fmt: c.number_format = fmt
    s["C7"] = "Days = one column per day (about 2 months). Weeks = one column per week (about 14 months)."
    s["C8"] = "Leave empty to start the timeline on the Monday before the project start."
    for a in ("C7", "C8"): s[a].font = font(italic=True, size=9)
    calc = [("Days per column", '=IF(B7="Weeks",7,1)', "0"),
            ("Timeline first day", '=IF(ISNUMBER(B8),B8,B6-WEEKDAY(B6,2)+1)', DATE)]
    for i, (k_, f, fmt) in enumerate(calc, 10):
        s.cell(row=i, column=1, value=k_).font = font(italic=True)
        c = s.cell(row=i, column=2, value=f); c.number_format = fmt; c.border = B; c.fill = fill("F3F4F6"); c.font = font()
    lists = [(1, "Phases", ["Planning", "Design", "Build", "Content", "Testing", "Launch"] if sample else ["Phase 1", "Phase 2", "Phase 3"]),
             (3, "Team / owners", ["Anna", "Ben", "Chloe", "Dev team", "Client"] if sample else ["Me"]),
             (5, "Priority", ["High", "Medium", "Low"]), (7, "Scale", ["Days", "Weeks"])]
    for col, head, items in lists:
        K.hdr(s, 14, [head], col0=col)
        for i in range(12):
            c = s.cell(row=15 + i, column=col, value=items[i] if i < len(items) else None)
            if col in (1, 3): input_cell(c)
            elif i < len(items): c.border = B; c.font = font(); c.fill = fill("F3F4F6")
    for c, w in zip("ABCDEFG", [30, 16, 3, 18, 3, 12, 3]): s.column_dimensions[c].width = w
    s.column_dimensions["H"].width = 10
    dv = DataValidation(type="list", formula1="=Settings!$G$15:$G$16", allow_blank=False); s.add_data_validation(dv); dv.add("B7")
    PH, OW, PR = "Settings!$A$15:$A$26", "Settings!$C$15:$C$26", "Settings!$E$15:$E$17"
    STEP, T0, PS = "Settings!$B$10", "Settings!$B$11", "Settings!$B$6"

    # ---------- Gantt ----------
    K.title(g, "Project Plan & Gantt Chart", "Type tasks in the yellow columns. Leave Start empty and fill 'Depends on' to chain tasks. Duration 0 = milestone ◆.")
    g["A3"] = f'=IF(Settings!B4="","",Settings!B4)&IF(Settings!B5="",""," · PM: "&Settings!B5)'; g["A3"].font = font(bold=True)
    heads = ["ID", "Phase", "Task", "Owner", "Priority", "Start (optional)", "Depends on (ID)", "Duration (days)", "% done", "Notes",
             "Start", "End", "Work days", "Status", "Planned % today", "Due key", ""]
    K.hdr(g, 4, heads, [5, 12, 32, 10, 9, 11, 9, 9, 8, 20, 11, 11, 7, 11, 8, 9, 2])
    style_range(g, TF, TL, 1, 16, inputs=range(2, 11), fmt={6: DATE, 7: "0", 8: "0", 9: "0%", 11: "dd mmm yy", 12: "dd mmm yy", 13: "0", 15: "0%"})
    for r in range(TF, TL + 1):
        g[f"A{r}"] = f'=IF(C{r}="","",ROW()-{TF - 1})'
        g[f"K{r}"] = (f'=IF(C{r}="","",IF(ISNUMBER(F{r}),F{r},IF(N(G{r})>0,IFERROR(INDEX($L$4:L{r - 1},G{r}+1)+1,""),{PS})))')
        g[f"L{r}"] = f'=IF(ISNUMBER(K{r}),K{r}+MAX(N(H{r}),1)-1,"")'
        g[f"M{r}"] = f'=IF(ISNUMBER(K{r}),IF(N(H{r})=0,0,NETWORKDAYS(K{r},L{r})),"")'
        g[f"N{r}"] = (f'=IF(C{r}="","",IF(N(I{r})>=1,"Done",IF(NOT(ISNUMBER(K{r})),"No date",IF(L{r}<TODAY(),"Late",'
                      f'IF(K{r}<=TODAY(),"In progress","Not started")))))')
        g[f"O{r}"] = f'=IF(NOT(ISNUMBER(K{r})),"",MIN(1,MAX(0,(TODAY()-K{r}+1)/(L{r}-K{r}+1))))'
        g[f"P{r}"] = f'=IF(OR(N{r}="",N{r}="Done",N{r}="No date"),"",IF(L{r}<=TODAY()+14,L{r}+ROW()/10000000,""))'
    for ref, f1 in ((f"B{TF}:B{TL}", PH), (f"D{TF}:D{TL}", OW), (f"E{TF}:E{TL}", PR)):
        dv = DataValidation(type="list", formula1="=" + f1, allow_blank=True); g.add_data_validation(dv); dv.add(ref)
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                        error="Type a percentage between 0% and 100%", showErrorMessage=True); g.add_data_validation(dv); dv.add(f"I{TF}:I{TL}")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(NT), allow_blank=True,
                        error="Type the ID of an earlier task", showErrorMessage=True); g.add_data_validation(dv); dv.add(f"G{TF}:G{TL}")
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="3650", allow_blank=True,
                        error="Type a number of days (0 for a milestone)", showErrorMessage=True); g.add_data_validation(dv); dv.add(f"H{TF}:H{TL}")
    g.conditional_formatting.add(f"N{TF}:N{TL}", CellIsRule(operator="equal", formula=['"Late"'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    g.conditional_formatting.add(f"N{TF}:N{TL}", CellIsRule(operator="equal", formula=['"Done"'], fill=fill("D1FAE5"), font=Font(name=F, color="065F46")))
    g.conditional_formatting.add(f"N{TF}:N{TL}", CellIsRule(operator="equal", formula=['"In progress"'], fill=fill("DBEAFE"), font=Font(name=F, color="1E40AF")))
    g.conditional_formatting.add(f"A{TF}:J{TL}", FormulaRule(formula=[f'$N{TF}="Done"'], font=Font(name=F, color="9CA3AF")))

    # timeline
    first, lastc = L(TL0), L(TL0 + NCOL - 1)
    thin = Side(style="thin", color="E5E7EB")
    for i in range(NCOL):
        col = L(TL0 + i); g.column_dimensions[col].width = 3.2
        dc = g[f"{col}4"]
        dc.value = f"={T0}+{i}*{STEP}" if i == 0 else f"={L(TL0 + i - 1)}4+{STEP}"
        dc.number_format = "d"; dc.font = Font(name=F, size=8, bold=True, color="FFFFFF"); dc.fill = fill(MAIN)
        dc.alignment = Alignment(horizontal="center", vertical="center", text_rotation=0)
        mc = g[f"{col}3"]
        mc.value = (f'=TEXT({col}4,"mmm yy")' if i == 0 else f'=IF(MONTH({col}4)<>MONTH({L(TL0 + i - 1)}4),TEXT({col}4,"mmm yy"),"")')
        mc.font = Font(name=F, size=8, bold=True, color=MAIN)
        wc = g[f"{col}2"]
        wc.value = f'=IF({STEP}=1,LEFT(TEXT({col}4,"ddd"),1),"")'
        wc.font = Font(name=F, size=7, color="6B7280"); wc.alignment = Alignment(horizontal="center")
    g.column_dimensions["Q"].width = 1.5
    tl = f"{first}{TF}:{lastc}{TL}"
    for r in range(TF, TL + 1):
        for i in range(NCOL):
            g.cell(row=r, column=TL0 + i).border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hit = f"ISNUMBER($K{TF}),$L{TF}>={first}$4,$K{TF}<={first}$4+{STEP}-1"
    g.conditional_formatting.add(tl, FormulaRule(formula=[f'AND({hit},N($H{TF})=0,$H{TF}<>"")'], fill=fill(MILE), stopIfTrue=True))
    g.conditional_formatting.add(tl, FormulaRule(formula=[f"AND({hit},{first}$4<$K{TF}+ROUND(($L{TF}-$K{TF}+1)*N($I{TF}),0))"], fill=fill(DONE), stopIfTrue=True))
    g.conditional_formatting.add(tl, FormulaRule(formula=[f'AND({hit},$N{TF}="Late")'], fill=fill(LATE), stopIfTrue=True))
    g.conditional_formatting.add(tl, FormulaRule(formula=[f"AND({hit})"], fill=fill(BAR), stopIfTrue=True))
    g.conditional_formatting.add(tl, FormulaRule(formula=[f"AND(TODAY()>={first}$4,TODAY()<={first}$4+{STEP}-1)"], fill=fill("FEF9C3")))
    g.conditional_formatting.add(tl, FormulaRule(formula=[f"AND({STEP}=1,WEEKDAY({first}$4,2)>5)"], fill=fill("F3F4F6")))
    g.conditional_formatting.add(f"{first}4:{lastc}4", FormulaRule(formula=[f"AND(TODAY()>={first}$4,TODAY()<={first}$4+{STEP}-1)"],
                                                                  fill=fill("F59E0B"), font=Font(name=F, size=8, bold=True, color="FFFFFF")))
    g.freeze_panes = f"D{TF}"
    g.sheet_view.zoomScale = 90

    if sample:
        random.seed(3)
        plan = [("Planning", "Kick-off meeting", "Anna", "High", 1), ("Planning", "Goals, audience & success metrics", "Anna", "High", 3),
                ("Planning", "Sitemap & page list", "Ben", "Medium", 3), ("Planning", "Budget & plan approved", "Client", "High", 0),
                ("Design", "Moodboard & style direction", "Chloe", "Medium", 4), ("Design", "Wireframes (key pages)", "Chloe", "High", 6),
                ("Design", "Visual design – home page", "Chloe", "High", 5), ("Design", "Visual design – inner pages", "Chloe", "Medium", 6),
                ("Design", "Design sign-off", "Client", "High", 0), ("Build", "Set up hosting & CMS", "Dev team", "Medium", 2),
                ("Build", "Build page templates", "Dev team", "High", 10), ("Build", "Booking & contact forms", "Dev team", "High", 4),
                ("Build", "Blog & SEO settings", "Ben", "Medium", 3), ("Content", "Write page copy", "Ben", "High", 8),
                ("Content", "Product photos", "Chloe", "Medium", 5), ("Content", "Load content into CMS", "Ben", "Medium", 4),
                ("Testing", "Cross-browser & mobile testing", "Dev team", "High", 4), ("Testing", "Speed & SEO check", "Ben", "Medium", 2),
                ("Testing", "Client review & fixes", "Client", "High", 5), ("Launch", "Go-live", "Dev team", "High", 0),
                ("Launch", "Launch email & social posts", "Anna", "Medium", 2), ("Launch", "Analytics review (2 weeks)", "Anna", "Low", 10)]
        deps = {2: 1, 3: 2, 4: 3, 5: 4, 6: 5, 7: 6, 8: 7, 9: 8, 10: 4, 11: 9, 12: 11, 13: 10, 14: 6, 15: 9, 16: 14, 17: 12, 18: 13,
                19: 17, 20: 19, 21: 20, 22: 20}
        for i, (ph, task, ow, pr, du) in enumerate(plan):
            r = TF + i; tid = i + 1
            g[f"B{r}"] = ph; g[f"C{r}"] = task; g[f"D{r}"] = ow; g[f"E{r}"] = pr; g[f"H{r}"] = du
            if tid == 1: g[f"F{r}"] = pstart
            elif tid in deps: g[f"G{r}"] = deps[tid]
        # % done from a python replica of the schedule
        start, end = {}, {}
        for i, (ph, task, ow, pr, du) in enumerate(plan):
            tid = i + 1
            start[tid] = pstart if tid == 1 else end[deps[tid]] + dt.timedelta(days=1)
            end[tid] = start[tid] + dt.timedelta(days=max(du, 1) - 1)
        for tid in start:
            r = TF + tid - 1
            if end[tid] < today:
                done = 1 if tid not in (12, 8) else 0.8
            elif start[tid] <= today:
                el = (today - start[tid]).days + 1; tot = (end[tid] - start[tid]).days + 1
                done = round(min(0.9, el / tot * random.uniform(0.6, 1.1)), 1)
            else: done = 0
            if done: g[f"I{r}"] = done
        g[f"J{TF}"] = "Video call, 1 hour"; g[f"J{TF + 3}"] = "Milestone"; g[f"J{TF + 8}"] = "Milestone"; g[f"J{TF + 19}"] = "Milestone"
        g[f"J{TF + 14}"] = "Book photographer early"

    # ---------- Dashboard ----------
    K.title(dash, "Project Dashboard", "Updates by itself from the Gantt tab.")
    dash["A3"] = f'=IF(Settings!B4="","Project",Settings!B4)&"  ·  Today: "&TEXT(TODAY(),"d mmm yyyy")'
    dash["A3"].font = font(bold=True)
    kp = [("Tasks", f'=COUNTIF({t_("C")},"?*")', "0"),
          ("Done", f'=COUNTIF({t_("N")},"Done")', "0"),
          ("Late tasks", f'=COUNTIF({t_("N")},"Late")', "0"),
          ("Actual progress", f'=IF(SUMIF({t_("C")},"?*",{t_("M")})=0,IF(A6=0,"",C6/A6),SUMPRODUCT({t_("M")},{t_("I")})/SUMIF({t_("C")},"?*",{t_("M")}))', "0%"),
          ("Planned by today", f'=IF(SUMIF({t_("C")},"?*",{t_("M")})=0,"",SUMPRODUCT({t_("M")},{t_("O")})/SUMIF({t_("C")},"?*",{t_("M")}))', "0%"),
          ("Schedule", '=IF(OR(A9="",C9=""),"",IF(A9>=C9-0.05,"On track","Behind"))', "General"),
          ("Project start", f'=IF(COUNT({t_("K")})=0,"",MIN({t_("K")}))', "d mmm yyyy"),
          ("Project finish", f'=IF(COUNT({t_("L")})=0,"",MAX({t_("L")}))', "d mmm yyyy"),
          ("Days to finish", '=IF(C12="","",MAX(0,C12-TODAY()))', "0")]
    for i, (lab, f, fmt) in enumerate(kp):
        K.kpi(dash, 5 + (i // 3) * 3, 1 + (i % 3) * 2, lab, f, fmt)
    for c, w in zip("ABCDEFGHIJKLM", [24, 12, 24, 18, 13, 12, 3, 16, 10, 10, 10, 3, 6]): dash.column_dimensions[c].width = w
    dash.conditional_formatting.add("E6", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=F, size=18, bold=True, color="B91C1C")))
    dash.conditional_formatting.add("E9", CellIsRule(operator="equal", formula=['"Behind"'], font=Font(name=F, size=18, bold=True, color="B91C1C")))
    dash.conditional_formatting.add("E9", CellIsRule(operator="equal", formula=['"On track"'], font=Font(name=F, size=18, bold=True, color="047857")))

    K.section(dash, "A15", "Progress by phase")
    K.hdr(dash, 16, ["Phase", "Tasks", "Done", "Late", "Progress", "Planned"])
    for i in range(10):
        r = 17 + i
        dash[f"A{r}"] = f'=IF(Settings!A{15 + i}="","",Settings!A{15 + i})'
        dash[f"B{r}"] = f'=IF(A{r}="","",COUNTIFS({t_("B")},A{r},{t_("C")},"?*"))'
        dash[f"C{r}"] = f'=IF(A{r}="","",COUNTIFS({t_("B")},A{r},{t_("N")},"Done"))'
        dash[f"D{r}"] = f'=IF(A{r}="","",COUNTIFS({t_("B")},A{r},{t_("N")},"Late"))'
        dash[f"E{r}"] = (f'=IF(OR(A{r}="",N(B{r})=0),"",IF(SUMIF({t_("B")},A{r},{t_("M")})=0,C{r}/B{r},'
                         f'SUMPRODUCT(--({t_("B")}=A{r}),{t_("M")},{t_("I")})/SUMIF({t_("B")},A{r},{t_("M")})))')
        dash[f"F{r}"] = (f'=IF(OR(A{r}="",N(B{r})=0),"",IF(SUMIF({t_("B")},A{r},{t_("M")})=0,"",'
                         f'SUMPRODUCT(--({t_("B")}=A{r}),{t_("M")},{t_("O")})/SUMIF({t_("B")},A{r},{t_("M")})))')
    style_range(dash, 17, 26, 1, 6, fmt={2: "0", 3: "0", 4: "0", 5: "0%", 6: "0%"})
    from openpyxl.formatting.rule import DataBarRule
    dash.conditional_formatting.add("E17:E26", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="93C5FD"))

    K.section(dash, "H15", "Status")
    K.hdr(dash, 16, ["Status", "Tasks"], col0=8)
    for i, stt in enumerate(STATUSES[:4]):
        r = 17 + i
        dash[f"H{r}"] = stt; dash[f"I{r}"] = f'=COUNTIF({t_("N")},H{r})'
    style_range(dash, 17, 20, 8, 9, fmt={9: "0"})
    K.section(dash, "H23", "Team workload")
    K.hdr(dash, 24, ["Owner", "Open", "Late", "Done"], col0=8)
    for i in range(8):
        r = 25 + i
        dash[f"H{r}"] = f'=IF(Settings!C{15 + i}="","",Settings!C{15 + i})'
        dash[f"I{r}"] = f'=IF(H{r}="","",COUNTIFS({t_("D")},H{r},{t_("N")},"<>Done",{t_("C")},"?*"))'
        dash[f"J{r}"] = f'=IF(H{r}="","",COUNTIFS({t_("D")},H{r},{t_("N")},"Late"))'
        dash[f"K{r}"] = f'=IF(H{r}="","",COUNTIFS({t_("D")},H{r},{t_("N")},"Done"))'
    style_range(dash, 25, 32, 8, 11, fmt={9: "0", 10: "0", 11: "0"})

    K.section(dash, "A29", "Late and due in the next 14 days")
    K.hdr(dash, 30, ["Task", "End", "Owner", "Phase", "% done", "Status"])
    for k in range(12):
        r = 31 + k
        dash[f"M{r}"] = f'=IF({k + 1}<=COUNT({t_("P")}),MATCH(SMALL({t_("P")},{k + 1}),{t_("P")},0),"")'
        dash[f"M{r}"].font = Font(name=F, size=8, color="FFFFFF")
        for col, src in zip("ABCDEF", "CLDBIN"):
            dash[f"{col}{r}"] = f'=IF($M{r}="","",IF(INDEX({t_(src)},$M{r})="","",INDEX({t_(src)},$M{r})))'
    style_range(dash, 31, 42, 1, 6, fmt={2: "ddd d mmm", 5: "0%"})
    dash.conditional_formatting.add("A31:F42", FormulaRule(formula=['$F31="Late"'], fill=fill("FEE2E2")))

    ch = BarChart(); ch.type = "bar"; ch.title = "Progress vs plan by phase"; ch.height = 7.5; ch.width = 13
    ch.add_data(Reference(dash, min_col=5, min_row=16, max_col=6, max_row=26), titles_from_data=True)
    ch.set_categories(Reference(dash, min_col=1, min_row=17, max_row=26)); ch.y_axis.majorGridlines = None
    ch.x_axis.scaling.orientation = "maxMin"; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1; ch.y_axis.number_format = "0%"
    for ser, col in zip(ch.series, (MAIN, "C7D2FE")):
        ser.graphicalProperties = GraphicalProperties(solidFill=col); ser.graphicalProperties.line.solidFill = col
    dash.add_chart(ch, "H35")
    pc = PieChart(); pc.title = "Tasks by status"; pc.height = 7.5; pc.width = 10
    pc.add_data(Reference(dash, min_col=9, min_row=16, max_row=20), titles_from_data=True)
    pc.set_categories(Reference(dash, min_col=8, min_row=17, max_row=20))
    dash.add_chart(pc, "N15")

    # ---------- Start Here ----------
    lines = [("How to start", True),
             ("1.  Settings: type the project name, manager and start date. Rename the phases and add your team.", False),
             ("2.  Gantt: add one row per task – phase, task name, owner, duration in days. Type a start date, or leave it empty and put the ID of the task it depends on.", False),
             ("3.  The bar appears on the timeline by itself. Update '% done' as you go – the dark part of the bar shows progress.", False),
             ("4.  Dashboard: progress vs plan, late tasks, phases, team workload and everything due in the next 14 days.", False),
             ("", False), ("Gantt colours", True),
             ("Light blue = planned  ·  Dark blue = done  ·  Red = late (end date passed, not 100% done)  ·  Orange = milestone (duration 0)  ·  Yellow column = today  ·  Grey = weekend.", False),
             ("", False), ("Tips", True),
             ("•  Dependencies: a task with no start date starts the day after the task in 'Depends on' ends. Only earlier tasks (smaller ID) can be used. With neither, it starts on the project start date.", False),
             ("•  Durations are calendar days; 'Work days' shows Monday–Friday days in that period.", False),
             ("•  Long project? Switch the timeline scale to Weeks on Settings, or set 'Timeline starts on' to scroll the chart.", False),
             ("•  Progress is weighted by work days, so long tasks count more than short ones.", False),
             ("", False), ("Colour key", True),
             ("Yellow cells are inputs: type here.  Grey cells calculate themselves: please don't overwrite them.", False),
             ("", False), ("Good to know", True),
             ("•  Room for 300 tasks. Works in Excel 2010+ and Google Sheets.", False),
             ("•  To use Google Sheets: open drive.google.com, upload this file, then choose Open with > Google Sheets.", False)]
    if sample:
        lines += [("", False), ("This copy contains sample data", True),
                  ("The 'Website Relaunch' project is a made-up example. Use the blank copy in your download for your own project.", False)]
    K.start(st, "Project Planner with Gantt Chart", "Plan tasks, chain dependencies and track progress on an automatic Gantt chart. Works in Excel and Google Sheets.", lines)
    finish(wb, sample, path)


META = {
    "name": "Project Planner with Gantt Chart – Excel & Google Sheets",
    "file_base": "Project-Planner-Gantt-Chart",
    "price": 19,
    "color": "#1D4ED8",
    "title": "Project Planner with Gantt Chart",
    "tagline": "Automatic Gantt · dependencies · progress vs plan — Excel & Google Sheets",
    "intro": "Plan any project in minutes. Type your tasks, durations and dependencies and the Gantt chart draws itself – with progress, late tasks, milestones and a dashboard that tells you if you're on track.",
    "features": [
        "Automatic Gantt chart: 60-column timeline by days or weeks, today marker and weekend shading",
        "Task dependencies: tasks start the day after the task they depend on – change one date and the plan moves",
        "Progress shading on each bar, red bars for late tasks and orange milestones",
        "Dashboard: actual vs planned progress, on-track/behind status, finish date and days left",
        "Progress by phase, team workload and a 'late and due in the next 14 days' list",
        "Phases, owners, priorities, work-day counts and notes; room for 300 tasks",
        "For client projects, website launches, events, renovations, product launches and school projects",
    ],
    "tags": ["gantt chart", "project planner", "project management", "project tracker", "google sheets template", "excel template",
             "task tracker", "project timeline", "gantt chart template", "work planner", "small business", "project schedule", "team planner"],
}
