"""Simple CRM & Sales Pipeline (sample + blank versions).

Contacts = people/companies, Deals = one row per opportunity (stage, value, follow-up date),
Activities = calls/emails/meetings log. Stage probabilities live on Settings and drive the weighted
forecast. The Dashboard shows open pipeline, weighted forecast, won vs target for the chosen month,
the funnel by stage, the last 6 months, owners, lead sources, follow-ups due and the biggest open deals.
"""
import random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from _common import Kit, F, B, MONEY, DATE, font, fill, style_range, input_cell, finish

MAIN = "0F766E"; SOFT = "E6F4F1"
K = Kit(MAIN, SOFT)
NC, ND, NA = 500, 500, 1000
CF, CL = 5, 4 + NC
DF, DL = 5, 4 + ND
AF, AL = 5, 4 + NA
NUM = '#,##0;-#,##0;"-"'
CUR = '#,##0;(#,##0);"-"'
STAGES = [("Lead", 0.1), ("Contacted", 0.2), ("Qualified", 0.4), ("Proposal sent", 0.6), ("Negotiation", 0.8), ("Won", 1.0), ("Lost", 0.0)]
OPEN_STAGES = [s for s, _ in STAGES[:5]]


def c_(col): return f"Contacts!${col}${CF}:${col}${CL}"
def d_(col): return f"Deals!${col}${DF}:${col}${DL}"
def a_(col): return f"Activities!${col}${AF}:${col}${AL}"


def build(sample, path):
    wb = Workbook()
    st = wb.active; st.title = "Start Here"
    dash = wb.create_sheet("Dashboard"); de = wb.create_sheet("Deals"); co = wb.create_sheet("Contacts")
    ac = wb.create_sheet("Activities"); s = wb.create_sheet("Settings")
    today = dt.date.today(); m0 = dt.date(today.year, today.month, 1)

    # ---------- Settings ----------
    K.title(s, "Settings", "Yellow cells are yours to edit. Stage names are fixed; change the win probability of each stage if you like.")
    rows = [("Business name", "Northwind Studio" if sample else None, None),
            ("Month to view", m0 if sample else "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)", "mmmm yyyy"),
            ("Monthly sales target", 40000 if sample else 10000, CUR),
            ("Stale after (days without activity)", 14, "0")]
    for i, (k_, v, fmt) in enumerate(rows, 4):
        s.cell(row=i, column=1, value=k_).font = font(bold=True)
        c = s.cell(row=i, column=2); input_cell(c, v)
        if fmt: c.number_format = fmt
    s["C5"] = "Type any date in the month you want on the Dashboard."; s["C5"].font = font(italic=True, size=9)
    s["C6"] = "Amounts are in your own currency (no symbol is shown)."; s["C6"].font = font(italic=True, size=9)
    s["A9"] = "Month key"; s["A9"].font = font(italic=True)
    c = s["B9"]; c.value = "=YEAR(B5)*100+MONTH(B5)"; c.number_format = "0"; c.border = B; c.fill = fill("F3F4F6"); c.font = font()
    K.hdr(s, 12, ["Stage (fixed)", "Win probability"], col0=1)
    for i, (nm, p) in enumerate(STAGES):
        a = s.cell(row=13 + i, column=1, value=nm); a.border = B; a.font = font(); a.fill = fill("F3F4F6")
        b = s.cell(row=13 + i, column=2); input_cell(b, p); b.number_format = "0%"
    lists = [(4, "Owners / sales reps", ["Anna", "Ben", "Chloe"] if sample else ["Me"]),
             (6, "Lead sources", ["Website", "Referral", "LinkedIn", "Cold email", "Event", "Google Ads", "Partner"]),
             (8, "Activity types", ["Call", "Email", "Meeting", "Demo", "LinkedIn message", "Note"]),
             (10, "Lost reasons", ["Price", "Went with competitor", "No budget", "No response", "Timing", "Not a fit"])]
    for col, head, items in lists:
        K.hdr(s, 12, [head], col0=col)
        for i in range(12):
            input_cell(s.cell(row=13 + i, column=col, value=items[i] if i < len(items) else None))
    s["A21"] = "Won and Lost close a deal. Open stages count towards the pipeline and the weighted forecast (value × probability)."
    s["A21"].font = font(italic=True, size=9)
    for c, w in zip("ABCDEFGHIJ", [34, 16, 3, 20, 3, 18, 3, 18, 3, 22]): s.column_dimensions[c].width = w
    STG, PRB = "Settings!$A$13:$A$19", "Settings!$B$13:$B$19"
    OW, SRC, AT, LR = "Settings!$D$13:$D$24", "Settings!$F$13:$F$24", "Settings!$H$13:$H$24", "Settings!$J$13:$J$24"
    MK, MS, TG, STALE = "Settings!$B$9", "Settings!$B$5", "Settings!$B$6", "Settings!$B$7"

    # ---------- Contacts ----------
    K.title(co, "Contacts", "One row per person. Deals and activities link to the name in column A, so keep names unique.")
    heads = ["Name", "Company", "Job title", "Email", "Phone", "Country", "Lead source", "Owner", "Added", "Notes",
             "Open deals", "Open value", "Won value", "Last activity", "Days since contact", "Flag"]
    K.hdr(co, 4, heads, [20, 20, 18, 26, 15, 12, 13, 10, 11, 26, 9, 12, 12, 12, 10, 14])
    style_range(co, CF, CL, 1, 16, inputs=range(1, 11), fmt={9: DATE, 11: "0", 12: CUR, 13: CUR, 14: DATE, 15: "0"})
    for r in range(CF, CL + 1):
        co[f"K{r}"] = f'=IF(A{r}="","",COUNTIFS({d_("B")},A{r},{d_("O")},">0"))'
        co[f"L{r}"] = f'=IF(A{r}="","",SUMIFS({d_("P")},{d_("B")},A{r}))'
        co[f"M{r}"] = f'=IF(A{r}="","",SUMIFS({d_("D")},{d_("B")},A{r},{d_("E")},"Won"))'
        co[f"N{r}"] = f'=IF(A{r}="","",IF(COUNTIF({a_("B")},A{r})=0,"",SUMPRODUCT(MAX(({a_("B")}=A{r})*{a_("A")}))))'
        co[f"O{r}"] = f'=IF(N{r}="","",TODAY()-N{r})'
        co[f"P{r}"] = f'=IF(OR(A{r}="",N(K{r})=0),"",IF(N{r}="","Never contacted",IF(O{r}>{STALE},"Going cold","")))'
    for ref, f1 in ((f"G{CF}:G{CL}", SRC), (f"H{CF}:H{CL}", OW)):
        dv = DataValidation(type="list", formula1="=" + f1, allow_blank=True); co.add_data_validation(dv); dv.add(ref)
    co.conditional_formatting.add(f"P{CF}:P{CL}", CellIsRule(operator="notEqual", formula=['""'], fill=fill("FEF3C7"), font=Font(name=F, bold=True, color="92400E")))
    co.freeze_panes = "B5"

    # ---------- Deals ----------
    K.title(de, "Deals (pipeline)", "One row per opportunity. Pick the contact and stage from the lists, and always set a follow-up date for open deals.")
    heads = ["Deal / opportunity", "Contact", "Company", "Deal value", "Stage", "Owner", "Lead source", "Created", "Expected close",
             "Next step", "Follow-up date", "Closed date", "Lost reason", "Notes",
             "Probability", "Weighted value", "Age (days)", "Flag", "Month key", "Open key", "Follow-up key"]
    K.hdr(de, 4, heads, [28, 18, 18, 12, 13, 10, 12, 11, 12, 26, 12, 11, 16, 22, 10, 12, 8, 15, 9, 10, 11])
    style_range(de, DF, DL, 1, 21, inputs=[1, 2] + list(range(4, 15)),
                fmt={4: CUR, 8: DATE, 9: DATE, 11: DATE, 12: DATE, 15: "0%", 16: CUR, 17: "0"})
    for r in range(DF, DL + 1):
        de[f"C{r}"] = f'=IF(B{r}="","",IFERROR(INDEX({c_("B")},MATCH(B{r},{c_("A")},0))&"",""))'
        de[f"O{r}"] = f'=IF(E{r}="","",IF(OR(E{r}="Won",E{r}="Lost"),0,IFERROR(INDEX({PRB},MATCH(E{r},{STG},0)),0)))'
        de[f"P{r}"] = f'=IF(O{r}="","",N(D{r})*O{r})'
        de[f"Q{r}"] = f'=IF(NOT(ISNUMBER(H{r})),"",IF(AND(ISNUMBER(L{r}),OR(E{r}="Won",E{r}="Lost")),L{r},TODAY())-H{r})'
        de[f"R{r}"] = (f'=IF(OR(E{r}="",E{r}="Won",E{r}="Lost"),"",IF(NOT(ISNUMBER(K{r})),"Set follow-up",IF(K{r}<TODAY(),"Follow up now",'
                       f'IF(AND(ISNUMBER(I{r}),I{r}<TODAY()),"Past close date",IF(K{r}<=TODAY()+2,"Due soon","")))))')
        de[f"S{r}"] = (f'=IF(AND(ISNUMBER(L{r}),OR(E{r}="Won",E{r}="Lost")),YEAR(L{r})*100+MONTH(L{r}),'
                       f'IF(AND(ISNUMBER(I{r}),E{r}<>"Won",E{r}<>"Lost"),YEAR(I{r})*100+MONTH(I{r}),""))')
        de[f"T{r}"] = f'=IF(N(O{r})=0,"",P{r}+ROW()/10000000)'
        de[f"U{r}"] = f'=IF(OR(E{r}="",E{r}="Won",E{r}="Lost",NOT(ISNUMBER(K{r}))),"",K{r}+ROW()/10000000)'
    for ref, f1 in ((f"B{DF}:B{DL}", f"Contacts!$A${CF}:$A${CL}"), (f"E{DF}:E{DL}", STG), (f"F{DF}:F{DL}", OW),
                    (f"G{DF}:G{DL}", SRC), (f"M{DF}:M{DL}", LR)):
        dv = DataValidation(type="list", formula1="=" + f1, allow_blank=True); de.add_data_validation(dv); dv.add(ref)
    dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, error="Type a date, for example 2026-11-03", showErrorMessage=True)
    de.add_data_validation(dv); dv.add(f"H{DF}:I{DL}"); dv.add(f"K{DF}:L{DL}")
    de.conditional_formatting.add(f"R{DF}:R{DL}", CellIsRule(operator="equal", formula=['"Follow up now"'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    de.conditional_formatting.add(f"R{DF}:R{DL}", CellIsRule(operator="notEqual", formula=['""'], fill=fill("FEF3C7"), font=Font(name=F, bold=True, color="92400E")))
    de.conditional_formatting.add(f"E{DF}:E{DL}", CellIsRule(operator="equal", formula=['"Won"'], fill=fill("D1FAE5"), font=Font(name=F, bold=True, color="065F46")))
    de.conditional_formatting.add(f"A{DF}:N{DL}", FormulaRule(formula=[f'$E{DF}="Lost"'], font=Font(name=F, color="9CA3AF")))
    de.freeze_panes = "B5"

    # ---------- Activities ----------
    K.title(ac, "Activities", "Log every call, email and meeting. The Contacts tab uses this to show the last contact date.")
    K.hdr(ac, 4, ["Date", "Contact", "Deal (optional)", "Type", "Summary", "Outcome / next step", "Owner", "Month key"], [11, 20, 28, 15, 40, 30, 10, 9])
    style_range(ac, AF, AL, 1, 8, inputs=range(1, 8), fmt={1: DATE})
    for r in range(AF, AL + 1):
        ac[f"H{r}"] = f'=IF(ISNUMBER(A{r}),YEAR(A{r})*100+MONTH(A{r}),"")'
    for ref, f1 in ((f"B{AF}:B{AL}", f"Contacts!$A${CF}:$A${CL}"), (f"C{AF}:C{AL}", f"Deals!$A${DF}:$A${DL}"), (f"D{AF}:D{AL}", AT), (f"G{AF}:G{AL}", OW)):
        dv = DataValidation(type="list", formula1="=" + f1, allow_blank=True); ac.add_data_validation(dv); dv.add(ref)
    ac.freeze_panes = "A5"

    # ---------- sample data ----------
    if sample:
        random.seed(11)
        first = ["Olivia", "Liam", "Emma", "Noah", "Ava", "James", "Sophia", "Lucas", "Mia", "Ethan", "Isla", "Mason", "Grace", "Leo", "Zoe",
                 "Jack", "Chloe", "Henry", "Ella", "Oscar", "Ruby", "Daniel", "Lily", "Sam", "Nora", "Max", "Hannah", "Ryan"]
        last = ["Smith", "Brown", "Miller", "Davis", "Wilson", "Taylor", "Clark", "Lewis", "Walker", "Hall", "Young", "King", "Wright", "Green",
                "Baker", "Adams", "Nelson", "Hill", "Campbell", "Murphy", "Cooper", "Reed", "Ward", "Fischer", "Novak", "Rossi", "Dubois", "Jensen"]
        comps = ["Brightline Dental", "Oakridge Realty", "Bluefin Logistics", "Maple & Co", "Summit Fitness", "Harbor Coffee", "Peak Legal",
                 "Greenleaf Organics", "Atlas Tutoring", "Copperpot Kitchen", "Nova Clinics", "Pinecrest Hotels", "Riverstone Builders",
                 "Swift Couriers", "Lumen Optics", "Cedar Accounting", "Urban Nest Interiors", "Silverline Auto", "Bloom Florists",
                 "Kestrel Software", "Willow Spa", "Ironworks Gym", "Coastal Vets", "Redwood Schools", "Echo Media", "Starling Travel",
                 "Fable Books", "Granite Insurance", "Vista Eyecare", "Orchid Weddings"]
        countries = ["USA", "USA", "USA", "UK", "UK", "Canada", "Australia", "Germany", "Netherlands", "Ireland"]
        titles = ["Owner", "Founder", "Marketing Manager", "Operations Director", "Office Manager", "CEO", "Head of Sales", "Practice Manager"]
        srcs = ["Website", "Website", "Referral", "Referral", "LinkedIn", "Cold email", "Event", "Google Ads", "Partner"]
        people = []
        for i, comp in enumerate(comps):
            nm = f"{first[i % len(first)]} {last[(i * 7) % len(last)]}"
            dom = comp.lower().replace(" & ", "").replace(" ", "") + ".com"
            people.append((nm, comp, random.choice(titles), nm.split()[0].lower() + "@" + dom,
                           f"+1 555 {random.randint(100, 999)} {random.randint(1000, 9999)}", random.choice(countries),
                           random.choice(srcs), random.choice(["Anna", "Anna", "Ben", "Chloe"]), today - dt.timedelta(days=random.randint(20, 200))))
        for i, p in enumerate(people):
            for j, v in enumerate(p, 1): co.cell(row=CF + i, column=j, value=v)
        co[f"J{CF}"] = "Prefers email, decision in Q1"; co[f"J{CF + 3}"] = "Met at the trade show"
        services = ["Website redesign", "SEO retainer (6 months)", "Brand identity package", "Google Ads management", "Booking system setup",
                    "Social media package", "E-commerce store build", "Monthly support plan", "Logo & signage", "Email marketing setup"]
        values = {"Website redesign": (4000, 12000), "SEO retainer (6 months)": (3000, 9000), "Brand identity package": (2500, 7000),
                  "Google Ads management": (1800, 6000), "Booking system setup": (1500, 4500), "Social media package": (1200, 4800),
                  "E-commerce store build": (6000, 18000), "Monthly support plan": (900, 3600), "Logo & signage": (800, 2500),
                  "Email marketing setup": (1000, 3000)}
        steps = ["Send proposal", "Book demo call", "Follow-up call", "Send revised quote", "Get sign-off from owner", "Share case studies",
                 "Intro call", "Send contract"]
        r = DF; deals = []
        for i in range(46):
            p = people[i % len(people)] if i < len(people) else random.choice(people)
            svc = random.choice(services); lo, hi = values[svc]; val = round(random.randint(lo, hi), -2)
            created = today - dt.timedelta(days=random.randint(3, 170))
            age = (today - created).days
            roll = random.random()
            if age > 40 and roll < 0.55: stage = "Won" if random.random() < 0.6 else "Lost"
            else: stage = random.choice(["Lead", "Contacted", "Qualified", "Qualified", "Proposal sent", "Proposal sent", "Negotiation"])
            de[f"A{r}"] = f"{p[1]} – {svc}"; de[f"B{r}"] = p[0]; de[f"D{r}"] = val; de[f"E{r}"] = stage; de[f"F{r}"] = p[7]
            de[f"G{r}"] = p[6]; de[f"H{r}"] = created
            if stage in ("Won", "Lost"):
                closed = created + dt.timedelta(days=random.randint(14, min(60, age)))
                de[f"I{r}"] = closed; de[f"L{r}"] = closed
                if stage == "Lost": de[f"M{r}"] = random.choice(["Price", "Went with competitor", "No budget", "No response", "Timing"])
            else:
                de[f"I{r}"] = today + dt.timedelta(days=random.randint(-5, 60))
                de[f"J{r}"] = random.choice(steps)
                de[f"K{r}"] = today + dt.timedelta(days=random.choice([-3, -1, 0, 1, 2, 3, 5, 7, 9, 12, 14]))
            deals.append((p[0], f"{p[1]} – {svc}", p[7], created))
            r += 1
        de[f"N{DF}"] = "Wants launch before spring"
        notes = {"Call": ["Discovery call – needs and budget", "Checked in on the proposal", "Answered questions about timeline"],
                 "Email": ["Sent proposal PDF", "Sent case studies", "Sent revised quote", "Follow-up email"],
                 "Meeting": ["On-site meeting with owner", "Kick-off meeting"], "Demo": ["Live demo of booking system"],
                 "LinkedIn message": ["Connected and sent intro"], "Note": ["Budget approved next month"]}
        outs = ["Positive – wants proposal", "Waiting for answer", "Asked for discount", "Follow up next week", "Sent to decision maker", ""]
        acts = []
        for nm, dl, ow, created in deals:
            n = random.randint(1, 4)
            for _ in range(n):
                d = created + dt.timedelta(days=random.randint(0, max(1, (today - created).days)))
                if d > today: d = today
                t = random.choice(list(notes)); acts.append((d, nm, dl, t, random.choice(notes[t]), random.choice(outs), ow))
        acts.sort()
        for i, a in enumerate(acts[:NA]):
            for j, v in enumerate(a, 1): ac.cell(row=AF + i, column=j, value=v or None)

    # ---------- Dashboard ----------
    K.title(dash, "Sales Dashboard", "Updates by itself from Deals, Contacts and Activities. Change the month and target on the Settings tab.")
    dash["A3"] = f'=IF(Settings!B4="","",Settings!B4&"  ·  ")&TEXT({MS},"mmmm yyyy")&"  ·  Target: "&TEXT({TG},"#,##0")'
    dash["A3"].font = font(bold=True)
    won_m = f'SUMIFS({d_("D")},{d_("E")},"Won",{d_("S")},{MK})'
    kp = [("Open pipeline value", f'=SUMIFS({d_("D")},{d_("O")},">0")', CUR),
          ("Weighted forecast", f'=SUM({d_("P")})', CUR),
          ("Open deals", f'=COUNTIF({d_("O")},">0")', "0"),
          ("Won this month", f"={won_m}", CUR),
          ("Target reached", f'=IF(N({TG})=0,"",A9/{TG})', "0%"),
          ("Follow-ups overdue", f'=COUNTIF({d_("R")},"Follow up now")', "0"),
          ("Win rate (all time)", f'=IF(COUNTIF({d_("E")},"Won")+COUNTIF({d_("E")},"Lost")=0,"",COUNTIF({d_("E")},"Won")/(COUNTIF({d_("E")},"Won")+COUNTIF({d_("E")},"Lost")))', "0%"),
          ("Average won deal", f'=IF(COUNTIF({d_("E")},"Won")=0,"",SUMIFS({d_("D")},{d_("E")},"Won")/COUNTIF({d_("E")},"Won"))', CUR),
          ("Avg. days to win", f'=IF(COUNTIF({d_("E")},"Won")=0,"",SUMIFS({d_("Q")},{d_("E")},"Won")/COUNTIF({d_("E")},"Won"))', "0")]
    for i, (lab, f, fmt) in enumerate(kp):
        K.kpi(dash, 5 + (i // 3) * 3, 1 + (i % 3) * 2, lab, f, fmt)
    for c, w in zip("ABCDEFGHIJKLM", [24, 12, 13, 26, 13, 12, 3, 16, 10, 12, 10, 3, 6]): dash.column_dimensions[c].width = w
    dash.conditional_formatting.add("E9", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=F, size=18, bold=True, color="B91C1C")))

    K.section(dash, "A15", "Pipeline by stage")
    K.hdr(dash, 16, ["Stage", "Deals", "Value", "Probability", "Weighted", "Share of open"])
    for i in range(7):
        r = 17 + i
        dash[f"A{r}"] = f"=Settings!A{13 + i}"
        dash[f"B{r}"] = f'=COUNTIF({d_("E")},A{r})'
        dash[f"C{r}"] = f'=SUMIFS({d_("D")},{d_("E")},A{r})'
        dash[f"D{r}"] = f"=Settings!B{13 + i}"
        dash[f"E{r}"] = f"=C{r}*D{r}" if i < 5 else '=""'
        dash[f"F{r}"] = f'=IF(OR({i}>4,N($C$24)=0),"",C{r}/$C$24)'
    style_range(dash, 17, 23, 1, 6, fmt={2: "0", 3: CUR, 4: "0%", 5: CUR, 6: "0%"})
    dash["A24"] = "Open total"; dash["B24"] = "=SUM(B17:B21)"; dash["C24"] = "=SUM(C17:C21)"; dash["E24"] = "=SUM(E17:E21)"
    for c, fmt in zip("ABCDEF", ["General", "0", CUR, "0%", CUR, "0%"]):
        x = dash[f"{c}24"]; x.font = font(bold=True); x.border = B; x.fill = fill(SOFT); x.number_format = fmt

    K.section(dash, "H15", "Lead sources (all time)")
    K.hdr(dash, 16, ["Source", "Deals", "Won value", "Win rate"], col0=8)
    for i in range(8):
        r = 17 + i
        dash[f"H{r}"] = f'=IF(Settings!F{13 + i}="","",Settings!F{13 + i})'
        dash[f"I{r}"] = f'=IF(H{r}="","",COUNTIF({d_("G")},H{r}))'
        dash[f"J{r}"] = f'=IF(H{r}="","",SUMIFS({d_("D")},{d_("G")},H{r},{d_("E")},"Won"))'
        dash[f"K{r}"] = (f'=IF(H{r}="","",IF(COUNTIFS({d_("G")},H{r},{d_("E")},"Won")+COUNTIFS({d_("G")},H{r},{d_("E")},"Lost")=0,"",'
                         f'COUNTIFS({d_("G")},H{r},{d_("E")},"Won")/(COUNTIFS({d_("G")},H{r},{d_("E")},"Won")+COUNTIFS({d_("G")},H{r},{d_("E")},"Lost"))))')
    style_range(dash, 17, 24, 8, 11, fmt={9: "0", 10: CUR, 11: "0%"})

    K.section(dash, "A27", "Last 6 months")
    K.hdr(dash, 28, ["Month", "Won value", "Target", "Deals won", "New deals", "Activities"])
    for i in range(6):
        r = 29 + i
        dash[f"A{r}"] = f"=DATE(YEAR({MS}),MONTH({MS})-{5 - i},1)"
        key = f"(YEAR(A{r})*100+MONTH(A{r}))"
        dash[f"B{r}"] = f'=SUMIFS({d_("D")},{d_("E")},"Won",{d_("S")},{key})'
        dash[f"C{r}"] = f"=N({TG})"
        dash[f"D{r}"] = f'=COUNTIFS({d_("E")},"Won",{d_("S")},{key})'
        dash[f"E{r}"] = f'=COUNTIFS({d_("H")},">="&A{r},{d_("H")},"<"&DATE(YEAR(A{r}),MONTH(A{r})+1,1))'
        dash[f"F{r}"] = f'=COUNTIF({a_("H")},{key})'
    style_range(dash, 29, 34, 1, 6, fmt={1: "mmm yyyy", 2: CUR, 3: CUR, 4: "0", 5: "0", 6: "0"})

    K.section(dash, "H27", "Sales reps")
    K.hdr(dash, 28, ["Owner", "Open deals", "Weighted", "Won (month)"], col0=8)
    for i in range(6):
        r = 29 + i
        dash[f"H{r}"] = f'=IF(Settings!D{13 + i}="","",Settings!D{13 + i})'
        dash[f"I{r}"] = f'=IF(H{r}="","",COUNTIFS({d_("F")},H{r},{d_("O")},">0"))'
        dash[f"J{r}"] = f'=IF(H{r}="","",SUMIFS({d_("P")},{d_("F")},H{r}))'
        dash[f"K{r}"] = f'=IF(H{r}="","",SUMIFS({d_("D")},{d_("F")},H{r},{d_("E")},"Won",{d_("S")},{MK}))'
    style_range(dash, 29, 34, 8, 11, fmt={9: "0", 10: CUR, 11: CUR})

    K.section(dash, "A37", "Follow-ups (overdue first)")
    K.hdr(dash, 38, ["Deal", "Follow-up", "Contact", "Next step", "Stage", "Owner"])
    for k in range(10):
        r = 39 + k
        dash[f"M{r}"] = f'=IF({k + 1}<=COUNT({d_("U")}),MATCH(SMALL({d_("U")},{k + 1}),{d_("U")},0),"")'
        dash[f"M{r}"].font = Font(name=F, size=8, color="FFFFFF")
        for col, src in zip("ABCDEF", "AKBJEF"):
            dash[f"{col}{r}"] = f'=IF($M{r}="","",IF(INDEX({d_(src)},$M{r})="","",INDEX({d_(src)},$M{r})))'
    style_range(dash, 39, 48, 1, 6, fmt={2: "ddd d mmm"})
    dash.conditional_formatting.add("A39:F48", FormulaRule(formula=["AND(ISNUMBER($B39),$B39<TODAY())"], fill=fill("FEE2E2")))
    dash.conditional_formatting.add("A39:F48", FormulaRule(formula=["AND(ISNUMBER($B39),$B39<=TODAY()+2)"], fill=fill("FEF3C7")))

    K.section(dash, "A51", "Biggest open deals (weighted)")
    K.hdr(dash, 52, ["Deal", "Value", "Stage", "Expected close", "Weighted", "Owner"])
    for k in range(5):
        r = 53 + k
        dash[f"M{r}"] = f'=IF({k + 1}<=COUNT({d_("T")}),MATCH(LARGE({d_("T")},{k + 1}),{d_("T")},0),"")'
        dash[f"M{r}"].font = Font(name=F, size=8, color="FFFFFF")
        for col, src in zip("ABCDEF", "ADEIPF"):
            dash[f"{col}{r}"] = f'=IF($M{r}="","",IF(INDEX({d_(src)},$M{r})="","",INDEX({d_(src)},$M{r})))'
    style_range(dash, 53, 57, 1, 6, fmt={2: CUR, 4: "d mmm yyyy", 5: CUR})

    ch = BarChart(); ch.type = "bar"; ch.title = "Open pipeline by stage"; ch.height = 7.5; ch.width = 12; ch.legend = None
    ch.add_data(Reference(dash, min_col=3, min_row=16, max_row=21), titles_from_data=True)
    ch.set_categories(Reference(dash, min_col=1, min_row=17, max_row=21)); ch.y_axis.majorGridlines = None
    ch.x_axis.scaling.orientation = "maxMin"
    ch.series[0].graphicalProperties = GraphicalProperties(solidFill=MAIN); ch.series[0].graphicalProperties.line.solidFill = MAIN
    dash.add_chart(ch, "H37")
    bc = BarChart(); bc.title = "Won vs target"; bc.height = 7.5; bc.width = 12
    bc.add_data(Reference(dash, min_col=2, min_row=28, max_col=3, max_row=34), titles_from_data=True)
    bc.set_categories(Reference(dash, min_col=1, min_row=29, max_row=34)); bc.y_axis.majorGridlines = None
    for ser, col in zip(bc.series, (MAIN, "C7D2CF")):
        ser.graphicalProperties = GraphicalProperties(solidFill=col); ser.graphicalProperties.line.solidFill = col
    dash.add_chart(bc, "H52")

    # ---------- Start Here ----------
    lines = [("How to start", True),
             ("1.  Settings: type your business name, monthly sales target, your sales reps and lead sources. Adjust the win probability of each stage if you like.", False),
             ("2.  Contacts: add one row per person (name, company, email, phone, source). Names must be unique – deals and activities link to them.", False),
             ("3.  Deals: add one row per opportunity – pick the contact, type the value, pick the stage and set a follow-up date. Company is filled in for you.", False),
             ("4.  Move deals through the stages. When a deal closes, choose Won or Lost and type the closed date (and the lost reason).", False),
             ("5.  Activities: log calls, emails and meetings. Contacts then show the last contact date and a 'Going cold' flag.", False),
             ("6.  Dashboard: open pipeline, weighted forecast, won vs target, funnel, follow-ups due and your biggest deals.", False),
             ("", False), ("Stages", True),
             ("Lead → Contacted → Qualified → Proposal sent → Negotiation → Won / Lost. Weighted forecast = deal value × the stage's win probability.", False),
             ("", False), ("Flags", True),
             ("Follow up now = follow-up date has passed.  Due soon = within 2 days.  Past close date = expected close has passed.  Set follow-up = an open deal has no follow-up date.", False),
             ("", False), ("Colour key", True),
             ("Yellow cells are inputs: type here.  Grey cells calculate themselves: please don't overwrite them.", False),
             ("", False), ("Good to know", True),
             ("•  Room for 500 contacts, 500 deals and 1,000 activities. Use any currency – amounts have no symbol.", False),
             ("•  Change 'Month to view' on Settings to look back at any month.", False),
             ("•  To use Google Sheets: open drive.google.com, upload this file, then choose Open with > Google Sheets.", False)]
    if sample:
        lines += [("", False), ("This copy contains sample data", True),
                  ("Northwind Studio and all contacts, deals and activities are made-up examples. Use the blank copy in your download for your own sales.", False)]
    K.start(st, "Simple CRM & Sales Pipeline", "Track leads, deals and follow-ups – and see your forecast at a glance. Works in Excel and Google Sheets.", lines)
    finish(wb, sample, path)


META = {
    "name": "Simple CRM & Sales Pipeline Tracker – Excel & Google Sheets",
    "file_base": "Simple-CRM-Sales-Pipeline",
    "price": 24,
    "color": "#0F766E",
    "title": "Simple CRM & Sales Pipeline Tracker",
    "tagline": "Contacts · deals · follow-ups · weighted forecast — Excel & Google Sheets",
    "intro": "A simple CRM without the monthly subscription. Keep your contacts, deals and follow-ups in one spreadsheet and see your pipeline, weighted forecast and sales vs target on an automatic dashboard.",
    "features": [
        "Deal pipeline with 7 stages (Lead to Won/Lost) and editable win probabilities",
        "Weighted sales forecast, open pipeline value, win rate, average deal size and days to win",
        "Follow-up tracker: overdue, due soon and past-close-date flags, plus a follow-up list on the dashboard",
        "Contacts list with open deals, won value, last contact date and 'going cold' alerts",
        "Activity log for calls, emails, meetings and demos",
        "Dashboard: won vs monthly target (last 6 months), funnel by stage, lead sources, sales reps and biggest open deals",
        "Room for 500 contacts, 500 deals and 1,000 activities; any currency",
    ],
    "tags": ["crm template", "sales pipeline", "sales tracker", "client tracker", "lead tracker", "google sheets crm", "excel crm",
             "small business", "customer database", "sales forecast", "freelancer crm", "deal tracker", "follow up tracker"],
}
