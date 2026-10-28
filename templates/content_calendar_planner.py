"""Content Calendar & Social Media Planner (sample + blank versions).

Content Plan = one row per post (inputs + helper columns). The Calendar tab shows any month as a
6-week grid with up to 3 posts per day (looked up with a date*100+sequence key, no array formulas),
colour-coded by platform. The Dashboard shows the month's plan vs goal, results per platform and
pillar, the next 14 days and the top posts by engagement.
"""
import random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.utils import get_column_letter as L
from _common import Kit, F, INK, B, MONEY, DATE, font, fill, style_range, input_cell, finish

MAIN = "7C3AED"; SOFT = "F3EEFF"
K = Kit(MAIN, SOFT)
CR = 600                      # content rows
FIRST, LAST = 5, 4 + CR
CP = "'Content Plan'"
NUM = '#,##0;-#,##0;"-"'
PLATFORMS = [("Instagram", "FCE7F3", "9D174D"), ("TikTok", "CCFBF1", "115E59"), ("Facebook", "DBEAFE", "1E40AF"),
             ("LinkedIn", "E0E7FF", "3730A3"), ("YouTube", "FEE2E2", "991B1B"), ("Pinterest", "FFE4E6", "9F1239"),
             ("X (Twitter)", "E5E7EB", "111827"), ("Threads", "EDE9FE", "5B21B6"), ("Blog", "FEF3C7", "92400E"),
             ("Newsletter", "DCFCE7", "166534")]
TYPES = ["Post", "Carousel", "Reel / Short", "Story", "Video", "Live", "Article", "Email", "Thread", "Pin"]
PILLARS = ["Educational", "Behind the scenes", "Promotional", "Community / UGC", "Entertainment", "Announcement", "Inspiration"]
STATUSES = ["Idea", "Drafting", "Designing", "Ready", "Scheduled", "Published", "Cancelled"]


def rng(c):
    return f"{CP}!${c}${FIRST}:${c}${LAST}"


def build(sample, path):
    wb = Workbook()
    st = wb.active; st.title = "Start Here"
    dash = wb.create_sheet("Dashboard"); cp = wb.create_sheet("Content Plan"); cal = wb.create_sheet("Calendar")
    ideas = wb.create_sheet("Ideas Bank"); tags = wb.create_sheet("Hashtags"); s = wb.create_sheet("Settings")
    today = dt.date.today()
    m0 = dt.date(today.year, today.month, 1)

    # ---------- Settings ----------
    K.title(s, "Settings", "Yellow cells are yours to edit. Rename platforms, types, pillars and team members to fit your brand.")
    rows = [("Brand / account name", "Bloom Studio" if sample else None, None),
            ("Month to view", m0 if sample else "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)", "mmmm yyyy"),
            ("Week starts on", "Sunday", None), ("Monthly post goal", 50 if sample else 30, "0")]
    for i, (k_, v, fmt) in enumerate(rows, 4):
        s.cell(row=i, column=1, value=k_).font = font(bold=True)
        c = s.cell(row=i, column=2); input_cell(c, v)
        if fmt: c.number_format = fmt
    s["C5"] = "Type any date in the month you want on the Calendar and Dashboard."; s["C5"].font = font(italic=True, size=9)
    calc = [("Month key", "=YEAR(B5)*100+MONTH(B5)", "0"), ("Month start", "=DATE(YEAR(B5),MONTH(B5),1)", DATE),
            ("Calendar grid start", '=B10-WEEKDAY(B10,IF(B6="Monday",2,1))+1', DATE)]
    for i, (k_, f, fmt) in enumerate(calc, 9):
        s.cell(row=i, column=1, value=k_).font = font(italic=True)
        c = s.cell(row=i, column=2, value=f); c.number_format = fmt; c.border = B; c.fill = fill("F3F4F6"); c.font = font()
    lists = [(1, "Platforms", [p[0] for p in PLATFORMS], True), (3, "Content types", TYPES, True), (5, "Content pillars", PILLARS, True),
             (7, "Status (fixed)", STATUSES, False), (9, "Team / owners", ["Maya", "Leo", "Sam"] if sample else ["Me"], True),
             (11, "Week start", ["Sunday", "Monday"], False), (13, "Priority", ["High", "Medium", "Low"], False)]
    for col, head, items, editable in lists:
        K.hdr(s, 14, [head], col0=col)
        for i in range(12):
            c = s.cell(row=15 + i, column=col, value=items[i] if i < len(items) else None)
            if editable: input_cell(c)
            else:
                c.border = B; c.font = font(); c.fill = fill("F3F4F6")
    s["A28"] = "Status names are used in formulas, so they can't be changed. Platform colours on the Calendar follow the default platform names."
    s["A28"].font = font(italic=True, size=9)
    for c, w in zip("ABCDEFGHIJKLM", [22, 16, 16, 3, 20, 3, 14, 3, 16, 3, 12, 3, 10]): s.column_dimensions[c].width = w
    for col, ref in ((2, "$K$15:$K$16"),):
        dv = DataValidation(type="list", formula1=f"=Settings!{ref}", allow_blank=False); s.add_data_validation(dv); dv.add("B6")
    PL, TY, PI, SS, TM, PR = ("Settings!$A$15:$A$26", "Settings!$C$15:$C$26", "Settings!$E$15:$E$26", "Settings!$G$15:$G$21",
                              "Settings!$I$15:$I$26", "Settings!$M$15:$M$17")
    MK, MS, GS = "Settings!$B$9", "Settings!$B$10", "Settings!$B$11"

    # ---------- Content Plan ----------
    K.title(cp, "Content Plan", "One row per post. Pick platform, type, pillar, owner and status from the lists. Add results after publishing.")
    heads = ["Publish date", "Time", "Platform", "Content type", "Pillar", "Title / hook", "Caption / notes", "Hashtags / keywords",
             "Link / asset", "Owner", "Status", "Campaign", "Reach", "Likes", "Comments", "Shares", "Saves", "Clicks",
             "Engagement", "Eng. rate", "Month key", "Day #", "Calendar key", "Calendar label", "Sort key", "Flag", "Top key", "Upcoming key"]
    K.hdr(cp, 4, heads, [12, 8, 13, 13, 16, 36, 30, 22, 18, 10, 11, 14, 9, 8, 9, 8, 8, 8, 11, 9, 9, 6, 11, 30, 11, 10, 10, 11])
    style_range(cp, FIRST, LAST, 1, 28, inputs=range(1, 19),
                fmt={1: DATE, 2: "h:mm AM/PM", 13: NUM, 14: NUM, 15: NUM, 16: NUM, 17: NUM, 18: NUM, 19: NUM, 20: "0.0%"})
    for r in range(FIRST, LAST + 1):
        cp[f"S{r}"] = f'=IF(COUNT(N{r}:Q{r})=0,"",SUM(N{r}:Q{r}))'
        cp[f"T{r}"] = f'=IF(OR(S{r}="",N(M{r})=0),"",S{r}/M{r})'
        cp[f"U{r}"] = f'=IF(A{r}="","",YEAR(A{r})*100+MONTH(A{r}))'
        cp[f"V{r}"] = f'=IF(OR(A{r}="",K{r}="Cancelled"),"",COUNTIFS($A${FIRST}:A{r},A{r},$K${FIRST}:K{r},"<>Cancelled"))'
        cp[f"W{r}"] = f'=IF(V{r}="","",A{r}*100+V{r})'
        cp[f"X{r}"] = f'=IF(V{r}="","",IF(K{r}="Published","✓ ","")&IF(C{r}="","",C{r}&": ")&IF(F{r}="",D{r},F{r}))'
        cp[f"Y{r}"] = f'=IF(A{r}="","",A{r}+IF(ISNUMBER(B{r}),MOD(B{r},1),0.5)+ROW()/10000000)'
        cp[f"Z{r}"] = (f'=IF(A{r}="","",IF(OR(K{r}="Published",K{r}="Cancelled"),"",IF(A{r}<TODAY(),"Overdue",'
                       f'IF(AND(A{r}<=TODAY()+2,K{r}<>"Scheduled",K{r}<>"Ready"),"Due soon",""))))')
        cp[f"AA{r}"] = f'=IF(AND(U{r}={MK},ISNUMBER(S{r}),K{r}="Published"),S{r}+ROW()/10000000,"")'
        cp[f"AB{r}"] = f'=IF(OR(A{r}="",K{r}="Cancelled",K{r}="Published"),"",IF(AND(A{r}>=TODAY(),A{r}<=TODAY()+13),Y{r},""))'
    for ref, f1 in ((f"C{FIRST}:C{LAST}", PL), (f"D{FIRST}:D{LAST}", TY), (f"E{FIRST}:E{LAST}", PI), (f"J{FIRST}:J{LAST}", TM), (f"K{FIRST}:K{LAST}", SS)):
        dv = DataValidation(type="list", formula1="=" + f1, allow_blank=True); cp.add_data_validation(dv); dv.add(ref)
    dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, error="Type a date, for example 2026-11-03", showErrorMessage=True)
    cp.add_data_validation(dv); dv.add(f"A{FIRST}:A{LAST}")
    cp.conditional_formatting.add(f"Z{FIRST}:Z{LAST}", CellIsRule(operator="equal", formula=['"Overdue"'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    cp.conditional_formatting.add(f"Z{FIRST}:Z{LAST}", CellIsRule(operator="equal", formula=['"Due soon"'], fill=fill("FEF3C7"), font=Font(name=F, bold=True, color="92400E")))
    cp.conditional_formatting.add(f"K{FIRST}:K{LAST}", CellIsRule(operator="equal", formula=['"Published"'], fill=fill("D1FAE5"), font=Font(name=F, color="065F46")))
    cp.conditional_formatting.add(f"A{FIRST}:R{LAST}", FormulaRule(formula=[f'$K{FIRST}="Cancelled"'], font=Font(name=F, color="9CA3AF", strike=True)))
    cp.freeze_panes = "C5"

    if sample:
        random.seed(5)
        weights = [("Instagram", 30), ("TikTok", 20), ("Facebook", 9), ("LinkedIn", 10), ("Pinterest", 9), ("YouTube", 5),
                   ("Blog", 5), ("Newsletter", 5), ("X (Twitter)", 5), ("Threads", 2)]
        plist = [p for p, w in weights for _ in range(w)]
        types = {"Instagram": ["Carousel", "Reel / Short", "Post", "Story"], "TikTok": ["Reel / Short"], "Facebook": ["Post", "Video"],
                 "LinkedIn": ["Post", "Article", "Carousel"], "Pinterest": ["Pin"], "YouTube": ["Video", "Reel / Short"], "Blog": ["Article"],
                 "Newsletter": ["Email"], "X (Twitter)": ["Thread", "Post"], "Threads": ["Thread"]}
        hooks = ["5 mistakes everyone makes with {t}", "How we pack an order (behind the scenes)", "New drop: the {t} collection is here",
                 "Customer of the week: Jenna's setup", "3 ways to style {t} for winter", "Q&A: your top questions about {t}",
                 "Before & after: {t} makeover", "Our 60-second morning routine", "Weekend sale – 20% off {t}", "Meet the team: Leo in the studio",
                 "Myth vs fact: {t}", "The story behind our logo", "Tutorial: {t} in 4 easy steps", "Poll: which colour should we make next?",
                 "Unboxing our newest {t}", "Monthly favourites: {t} edition", "Free checklist: planning your {t}", "Day in the life of a small business"]
        topics = ["linen bags", "desk organizers", "ceramic mugs", "planners", "candles", "gift boxes"]
        pill = {"5 mistakes": "Educational", "How we pack": "Behind the scenes", "New drop": "Announcement", "Customer": "Community / UGC",
                "3 ways": "Inspiration", "Q&A": "Educational", "Before & after": "Inspiration", "Our 60-second": "Behind the scenes",
                "Weekend sale": "Promotional", "Meet the team": "Behind the scenes", "Myth vs fact": "Educational", "The story": "Behind the scenes",
                "Tutorial": "Educational", "Poll": "Community / UGC", "Unboxing": "Entertainment", "Monthly favourites": "Promotional",
                "Free checklist": "Educational", "Day in the life": "Entertainment"}
        reach = {"Instagram": (900, 4200), "TikTok": (1500, 12000), "Facebook": (400, 1800), "LinkedIn": (500, 2600), "Pinterest": (700, 5200),
                 "YouTube": (300, 2400), "Blog": (200, 900), "Newsletter": (1100, 1400), "X (Twitter)": (300, 1500), "Threads": (200, 900)}
        tags_by = {"Instagram": "#smallbusiness #shopsmall #madewithlove", "TikTok": "#smallbusinesscheck #packwithme", "Pinterest": "gift ideas, desk setup",
                   "LinkedIn": "#entrepreneurship #ecommerce", "YouTube": "studio vlog, small business"}
        nxt = dt.date(m0.year + (m0.month == 12), m0.month % 12 + 1, 1)
        end = dt.date(nxt.year + (nxt.month == 12), nxt.month % 12 + 1, 1)
        posts = []
        d = m0
        while d < end:
            n = random.choice([1, 2, 2, 2, 3]) if d < nxt else random.choice([0, 1, 1, 2])
            if d.weekday() >= 5: n = max(0, n - 1)
            for _ in range(n):
                p = random.choice(plist); h = random.choice(hooks)
                posts.append((d, random.choice([dt.time(9, 0), dt.time(12, 30), dt.time(17, 0), dt.time(19, 30)]), p, random.choice(types[p]),
                              next(v for k_, v in pill.items() if h.startswith(k_)), h.format(t=random.choice(topics))))
            d += dt.timedelta(days=1)
        r = FIRST
        for (d, tm, p, ty, pi, title) in posts:
            cp[f"A{r}"] = d; cp[f"B{r}"] = tm; cp[f"C{r}"] = p; cp[f"D{r}"] = ty; cp[f"E{r}"] = pi; cp[f"F{r}"] = title
            if p in tags_by: cp[f"H{r}"] = tags_by[p]
            cp[f"J{r}"] = random.choice(["Maya", "Maya", "Leo", "Sam"])
            if "sale" in title.lower() or "drop" in title.lower(): cp[f"L{r}"] = "Autumn launch"
            if d < today:
                roll = random.random()
                status = "Published" if roll < 0.9 else ("Cancelled" if roll < 0.95 else "Drafting")
            elif d <= today + dt.timedelta(days=6):
                status = random.choice(["Scheduled", "Scheduled", "Ready", "Designing"])
            else:
                status = random.choice(["Drafting", "Idea", "Idea", "Designing"])
            cp[f"K{r}"] = status
            if status == "Published":
                lo, hi = reach[p]; rc = random.randint(lo, hi)
                cp[f"M{r}"] = rc; cp[f"N{r}"] = int(rc * random.uniform(0.02, 0.07)); cp[f"O{r}"] = int(rc * random.uniform(0.002, 0.01))
                cp[f"P{r}"] = int(rc * random.uniform(0.001, 0.01)); cp[f"Q{r}"] = int(rc * random.uniform(0.003, 0.02))
                cp[f"R{r}"] = int(rc * random.uniform(0.002, 0.015))
            r += 1

    # ---------- Calendar ----------
    K.title(cal, "Content Calendar", "Up to 3 posts per day. Change the month on the Settings tab. ✓ = published.")
    cal["A3"] = f'=TEXT({MS},"mmmm yyyy")'; cal["A3"].font = Font(name=F, size=16, bold=True, color=INK)
    cal["J1"] = f"={MS}"; cal["J2"] = f"={GS}"
    for a in ("J1", "J2"): cal[a].number_format = DATE; cal[a].font = Font(name=F, size=8, color="FFFFFF")
    for d in range(7):
        c = cal.cell(row=5, column=1 + d, value=f'=TEXT($J$2+{d},"dddd")')
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = fill(MAIN); c.alignment = Alignment(horizontal="center"); c.border = B
        cal.column_dimensions[L(1 + d)].width = 26
    thin = Side(style="thin", color="C4B5FD"); none_ = Side(style=None)
    day_rows, post_rows, more_rows = [], [], []
    for w in range(6):
        r0 = 6 + w * 5
        day_rows.append(r0); post_rows.append((r0 + 1, r0 + 3)); more_rows.append(r0 + 4)
        for d in range(7):
            col = L(1 + d)
            c = cal[f"{col}{r0}"]; c.value = f"=$J$2+{w * 7 + d}"; c.number_format = "d"
            c.font = Font(name=F, bold=True, size=11, color=INK); c.alignment = Alignment(horizontal="left")
            for k in range(1, 4):
                x = cal[f"{col}{r0 + k}"]
                x.value = f'=IFERROR(INDEX({rng("X")},MATCH({col}{r0}*100+{k},{rng("W")},0)),"")'
                x.font = font(size=9); x.alignment = Alignment(shrink_to_fit=False, wrap_text=False)
            m = cal[f"{col}{r0 + 4}"]
            m.value = (f'=IF(COUNTIFS({rng("A")},{col}{r0},{rng("K")},"<>Cancelled")>3,'
                       f'"+"&(COUNTIFS({rng("A")},{col}{r0},{rng("K")},"<>Cancelled")-3)&" more","")')
            m.font = Font(name=F, italic=True, size=8, color="6B7280")
            for k in range(5):
                cal[f"{col}{r0 + k}"].border = Border(left=thin, right=thin, top=thin if k == 0 else none_, bottom=thin if k == 4 else none_)
        for k in range(5): cal.row_dimensions[r0 + k].height = 16 if k else 18
    day_ref = " ".join(f"A{r}:G{r}" for r in day_rows)
    cal.conditional_formatting.add(day_ref, FormulaRule(formula=["A6=TODAY()"], fill=fill("EDE9FE"), font=Font(name=F, bold=True, color=MAIN)))
    cal.conditional_formatting.add(day_ref, FormulaRule(formula=["MONTH(A6)<>MONTH($J$1)"], font=Font(name=F, color="C0C4CC")))
    for k in range(1, 5):   # rows k below each day row: compare with the day cell k rows up
        ref = " ".join(f"A{r + k}:G{r + k}" for r in day_rows)
        cal.conditional_formatting.add(ref, FormulaRule(formula=["MONTH(A6)<>MONTH($J$1)"], font=Font(name=F, color="C0C4CC")))
    post_ref = " ".join(f"A{a}:G{b}" for a, b in post_rows)
    for name, bg, fg in PLATFORMS:
        cal.conditional_formatting.add(post_ref, FormulaRule(formula=[f'ISNUMBER(SEARCH("{name}:",A7))'], fill=fill(bg), font=Font(name=F, size=9, color=fg)))
    cal.sheet_view.zoomScale = 90

    # ---------- Ideas Bank ----------
    K.title(ideas, "Ideas Bank", "Park every idea here. When you schedule one, add it to the Content Plan and mark it Used.")
    K.hdr(ideas, 4, ["Idea", "Platform", "Pillar", "Format", "Angle / notes", "Priority", "Added", "Used?"], [40, 13, 18, 13, 36, 10, 12, 8])
    style_range(ideas, 5, 204, 1, 8, inputs=range(1, 9), fmt={7: DATE})
    for ref, f1 in (("B5:B204", PL), ("C5:C204", PI), ("D5:D204", TY), ("F5:F204", PR)):
        dv = DataValidation(type="list", formula1="=" + f1, allow_blank=True); ideas.add_data_validation(dv); dv.add(ref)
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True); ideas.add_data_validation(dv); dv.add("H5:H204")
    ideas.conditional_formatting.add("F5:F204", CellIsRule(operator="equal", formula=['"High"'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    ideas.conditional_formatting.add("A5:G204", FormulaRule(formula=['$H5="Yes"'], font=Font(name=F, color="9CA3AF", strike=True)))
    if sample:
        idea_rows = [("Holiday gift guide under $30", "Pinterest", "Promotional", "Pin", "Pin 6 product photos as a board", "High", "No"),
                     ("Studio tour in 30 seconds", "TikTok", "Behind the scenes", "Reel / Short", "Trending audio, fast cuts", "High", "No"),
                     ("How to choose the right planner", "Blog", "Educational", "Article", "SEO: 'best planner for 2027'", "Medium", "No"),
                     ("Customer photo contest", "Instagram", "Community / UGC", "Post", "Prize: gift box; tag us to enter", "High", "No"),
                     ("Year in review newsletter", "Newsletter", "Announcement", "Email", "Top 5 products + thank-you code", "Medium", "No"),
                     ("Packaging sustainability story", "LinkedIn", "Behind the scenes", "Article", "Numbers: plastic saved this year", "Low", "No"),
                     ("Desk setup inspiration", "Pinterest", "Inspiration", "Pin", "", "Medium", "Yes")]
        for i, row in enumerate(idea_rows):
            for j, v in enumerate(row[:5], 1): ideas.cell(row=5 + i, column=j, value=v or None)
            ideas.cell(row=5 + i, column=6, value=row[5]); ideas.cell(row=5 + i, column=7, value=today - dt.timedelta(days=3 * i + 1))
            ideas.cell(row=5 + i, column=8, value=row[6])
    ideas.freeze_panes = "A5"

    # ---------- Hashtags ----------
    K.title(tags, "Hashtag & Keyword Bank", "Save tested hashtag sets and keywords so you can copy them into posts.")
    K.hdr(tags, 4, ["Group / topic", "Hashtags or keywords", "Best for", "Notes"], [22, 60, 14, 30])
    style_range(tags, 5, 64, 1, 4, inputs=range(1, 5))
    dv = DataValidation(type="list", formula1="=" + PL, allow_blank=True); tags.add_data_validation(dv); dv.add("C5:C64")
    for c in range(5, 65): tags[f"B{c}"].alignment = Alignment(wrap_text=True, vertical="top")
    if sample:
        tag_rows = [("Brand", "#bloomstudio #bloomstudiomade", "Instagram", "Use on every post"),
                    ("Small business", "#smallbusiness #shopsmall #supportsmallbusiness #handmade", "Instagram", ""),
                    ("Packing videos", "#packwithme #smallbusinesscheck #asmrpacking", "TikTok", "Best reach in the evening"),
                    ("Gifts", "gift ideas, gifts for her, desk gifts, christmas gifts", "Pinterest", "Keywords, not hashtags"),
                    ("Founder", "#entrepreneurship #ecommerce #founderstory", "LinkedIn", "Max 3 on LinkedIn")]
        for i, row in enumerate(tag_rows):
            for j, v in enumerate(row, 1): tags.cell(row=5 + i, column=j, value=v or None)

    # ---------- Dashboard ----------
    K.title(dash, "Content Dashboard", "Updates by itself from the Content Plan. Change the month on the Settings tab.")
    dash["A3"] = f'=IF(Settings!B4="","",Settings!B4&"  ·  ")&TEXT({MS},"mmmm yyyy")&"  ·  Goal: "&Settings!B7&" posts"'
    dash["A3"].font = font(bold=True)
    kp = [("Posts planned (month)", f'=COUNTIFS({rng("U")},{MK},{rng("K")},"<>Cancelled")', "0"),
          ("Published (month)", f'=COUNTIFS({rng("U")},{MK},{rng("K")},"Published")', "0"),
          ("Goal progress (published)", '=IF(N(Settings!B7)=0,"",C6/Settings!B7)', "0%"),
          ("Scheduled or ready", f'=COUNTIFS({rng("U")},{MK},{rng("K")},"Scheduled")+COUNTIFS({rng("U")},{MK},{rng("K")},"Ready")', "0"),
          ("Overdue posts (all)", f'=COUNTIF({rng("Z")},"Overdue")', "0"),
          ("Engagement rate (month)", f'=IF(SUMIFS({rng("M")},{rng("U")},{MK},{rng("K")},"Published")=0,"",'
                                      f'SUMIFS({rng("S")},{rng("U")},{MK},{rng("K")},"Published")/SUMIFS({rng("M")},{rng("U")},{MK},{rng("K")},"Published"))', "0.0%")]
    for i, (lab, f, fmt) in enumerate(kp):
        K.kpi(dash, 5 + (i // 3) * 3, 1 + (i % 3) * 2, lab, f, fmt)
    for c, w in zip("ABCDEFGHIJKL", [16, 12, 14, 38, 13, 12, 3, 22, 10, 10, 3, 6]): dash.column_dimensions[c].width = w
    dash.conditional_formatting.add("C9", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=F, size=18, bold=True, color="B91C1C")))

    K.section(dash, "A12", "By platform (month)")
    K.hdr(dash, 13, ["Platform", "Planned", "Published", "Reach", "Engagement", "Eng. rate"])
    for i in range(10):
        r = 14 + i
        dash[f"A{r}"] = f'=IF(Settings!A{15+i}="","",Settings!A{15+i})'
        dash[f"B{r}"] = f'=IF(A{r}="","",COUNTIFS({rng("C")},A{r},{rng("U")},{MK},{rng("K")},"<>Cancelled"))'
        dash[f"C{r}"] = f'=IF(A{r}="","",COUNTIFS({rng("C")},A{r},{rng("U")},{MK},{rng("K")},"Published"))'
        dash[f"D{r}"] = f'=IF(A{r}="","",SUMIFS({rng("M")},{rng("C")},A{r},{rng("U")},{MK},{rng("K")},"Published"))'
        dash[f"E{r}"] = f'=IF(A{r}="","",SUMIFS({rng("S")},{rng("C")},A{r},{rng("U")},{MK},{rng("K")},"Published"))'
        dash[f"F{r}"] = f'=IF(OR(A{r}="",N(D{r})=0),"",E{r}/D{r})'
    style_range(dash, 14, 23, 1, 6, fmt={2: "0", 3: "0", 4: NUM, 5: NUM, 6: "0.0%"})
    dash["A24"] = "Total"
    for c in "BCDE": dash[f"{c}24"] = f"=SUM({c}14:{c}23)"
    dash["F24"] = '=IF(N(D24)=0,"",E24/D24)'
    for c, fmt in zip("ABCDEF", ["General", "0", "0", NUM, NUM, "0.0%"]):
        x = dash[f"{c}24"]; x.font = font(bold=True); x.border = B; x.fill = fill(SOFT); x.number_format = fmt

    K.section(dash, "H12", "Content mix by pillar (month)")
    K.hdr(dash, 13, ["Pillar", "Posts", "Share"], col0=8)
    for i in range(10):
        r = 14 + i
        dash[f"H{r}"] = f'=IF(Settings!E{15+i}="","",Settings!E{15+i})'
        dash[f"I{r}"] = f'=IF(H{r}="","",COUNTIFS({rng("E")},H{r},{rng("U")},{MK},{rng("K")},"<>Cancelled"))'
        dash[f"J{r}"] = f'=IF(OR(H{r}="",N($B$24)=0),"",I{r}/$B$24)'
    style_range(dash, 14, 23, 8, 10, fmt={9: "0", 10: "0%"})

    K.section(dash, "A27", "Coming up (next 14 days)")
    K.hdr(dash, 28, ["Date", "Time", "Platform", "Title / hook", "Owner", "Status"])
    for k in range(10):
        r = 29 + k
        dash[f"L{r}"] = f'=IF({k+1}<=COUNT({rng("AB")}),MATCH(SMALL({rng("AB")},{k+1}),{rng("AB")},0),"")'
        dash[f"L{r}"].font = Font(name=F, size=8, color="FFFFFF")
        for col, src in zip("ABCDEF", "ABCFJK"):
            dash[f"{col}{r}"] = f'=IF($L{r}="","",IF(INDEX({rng(src)},$L{r})="","",INDEX({rng(src)},$L{r})))'
    style_range(dash, 29, 38, 1, 6, fmt={1: "ddd d mmm", 2: "h:mm AM/PM"})
    dash.conditional_formatting.add("A29:F38", FormulaRule(formula=["AND(ISNUMBER($A29),$A29<=TODAY()+1)"], fill=fill("FEF3C7")))

    K.section(dash, "A41", "Top posts by engagement (month)")
    K.hdr(dash, 42, ["Date", "Platform", "Type", "Title / hook", "Engagement", "Eng. rate"])
    for k in range(5):
        r = 43 + k
        dash[f"L{r}"] = f'=IF({k+1}<=COUNT({rng("AA")}),MATCH(LARGE({rng("AA")},{k+1}),{rng("AA")},0),"")'
        dash[f"L{r}"].font = Font(name=F, size=8, color="FFFFFF")
        for col, src in zip("ABCDEF", "ACDFST"):
            dash[f"{col}{r}"] = f'=IF($L{r}="","",INDEX({rng(src)},$L{r}))'
    style_range(dash, 43, 47, 1, 6, fmt={1: "d mmm", 5: NUM, 6: "0.0%"})

    ch = BarChart(); ch.type = "bar"; ch.title = "Planned vs published by platform"; ch.height = 7.5; ch.width = 13
    ch.add_data(Reference(dash, min_col=2, min_row=13, max_col=3, max_row=23), titles_from_data=True)
    ch.set_categories(Reference(dash, min_col=1, min_row=14, max_row=23)); ch.y_axis.majorGridlines = None
    for ser, col in zip(ch.series, ("C4B5FD", MAIN)):
        ser.graphicalProperties = GraphicalProperties(solidFill=col); ser.graphicalProperties.line.solidFill = col
    dash.add_chart(ch, "H26")
    pc = PieChart(); pc.title = "Content mix"; pc.height = 7.5; pc.width = 10
    pc.add_data(Reference(dash, min_col=9, min_row=13, max_row=23), titles_from_data=True)
    pc.set_categories(Reference(dash, min_col=8, min_row=14, max_row=23))
    dash.add_chart(pc, "N12")

    # ---------- Start Here ----------
    lines = [("How to start", True),
             ("1.  Settings: type your brand name, the month you want to see, your monthly post goal, and rename platforms, content types, pillars and team members if you like.", False),
             ("2.  Content Plan: add one row per post – date, time, platform, type, pillar, title and status. Everything else is optional.", False),
             ("3.  Calendar: see the month as a calendar, up to 3 posts per day, colour-coded by platform. ✓ marks published posts.", False),
             ("4.  After posting: add reach, likes, comments, shares, saves and clicks. The Dashboard shows engagement and your top posts.", False),
             ("5.  Ideas Bank and Hashtags: park ideas and tested hashtag sets so you never start from a blank page.", False),
             ("", False), ("Status flow", True),
             ("Idea → Drafting → Designing → Ready → Scheduled → Published (or Cancelled). Posts past their date that aren't published are flagged Overdue; posts due within 2 days that aren't ready are flagged Due soon.", False),
             ("", False), ("Colour key", True),
             ("Yellow cells are inputs: type here.  Grey cells calculate themselves: please don't overwrite them.", False),
             ("", False), ("Good to know", True),
             ("•  Room for 600 posts, 200 ideas and 60 hashtag sets. Change the month on Settings to plan ahead or look back.", False),
             ("•  Calendar colours follow the default platform names (Instagram, TikTok, Facebook, LinkedIn, YouTube, Pinterest, X (Twitter), Threads, Blog, Newsletter).", False),
             ("•  To use Google Sheets: open drive.google.com, upload this file, then choose Open with > Google Sheets.", False)]
    if sample:
        lines += [("", False), ("This copy contains sample data", True),
                  ("Bloom Studio and all posts and results are made-up examples. Use the blank copy in your download for your own plan.", False)]
    K.start(st, "Content Calendar & Social Media Planner", "Plan, schedule and track every post across all your channels. Works in Excel and Google Sheets.", lines)
    finish(wb, sample, path)


META = {
    "name": "Content Calendar & Social Media Planner – Excel & Google Sheets",
    "file_base": "Content-Calendar-Planner",
    "price": 19,
    "color": "#7C3AED",
    "title": "Content Calendar & Social Media Planner",
    "tagline": "Monthly calendar · post planner · results per platform — Excel & Google Sheets",
    "intro": "Plan a month of content in one place. Add your posts once and get a colour-coded monthly calendar, an overdue/due-soon tracker and a dashboard that shows what's working on every platform.",
    "features": [
        "Monthly calendar view: up to 3 posts per day, colour-coded by platform, any month you choose",
        "Content plan for 600 posts: date, time, platform, format, pillar, hook, caption, hashtags, owner, status and campaign",
        "Status workflow from Idea to Published, with Overdue and Due soon flags",
        "Dashboard: posts planned vs goal, published, engagement rate, next 14 days and top posts by engagement",
        "Results per platform and content mix by pillar, with charts",
        "Ideas bank and hashtag & keyword bank",
        "Works for Instagram, TikTok, Facebook, LinkedIn, YouTube, Pinterest, X, Threads, blogs and newsletters",
    ],
    "tags": ["content calendar", "social media planner", "instagram planner", "content planner", "social media calendar", "google sheets",
             "excel template", "tiktok planner", "marketing calendar", "small business", "content creator", "post scheduler", "editorial calendar"],
}
