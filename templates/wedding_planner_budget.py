"""Builds the Wedding Planner & Budget spreadsheet (sample + blank versions)."""
import random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter as L
from _common import Kit, F, INK, B, MONEY, DATE, font, fill, style_range, input_cell, finish

MAIN = "8B4D6B"; SOFT = "F7EEF3"
K = Kit(MAIN, SOFT)
BR = 120; GR = 400; VR = 60; TR = 120; TABLES = 40
K0 = '#,##0;(#,##0);"-"'

CATS = [("Venue & rentals", 0.30), ("Catering & cake", 0.25), ("Photography & video", 0.12), ("Attire & beauty", 0.08), ("Flowers & decor", 0.08),
        ("Music & entertainment", 0.05), ("Stationery & invitations", 0.02), ("Rings", 0.03), ("Officiant & ceremony", 0.01), ("Transportation", 0.02),
        ("Favors & gifts", 0.02), ("Other / buffer", 0.02)]

TASKS = [(12, "Set the budget and decide who is paying"), (12, "Draft the guest list"), (12, "Choose a date and book the venue"), (11, "Hire a planner (optional)"),
         (10, "Book photographer and videographer"), (10, "Book caterer"), (9, "Choose wedding party"), (9, "Send save-the-dates"), (9, "Shop for wedding dress / suit"),
         (8, "Book band or DJ"), (8, "Book florist"), (7, "Register for gifts"), (7, "Build a wedding website"), (6, "Book officiant"), (6, "Plan honeymoon and check passports"),
         (5, "Order cake / book tasting"), (5, "Book hair & makeup trial"), (4, "Order invitations"), (4, "Book transportation"), (3, "Buy wedding rings"),
         (3, "Plan rehearsal dinner"), (2, "Mail invitations"), (2, "Final dress fitting"), (2, "Apply for marriage license"), (1, "Chase missing RSVPs"),
         (1, "Give final headcount to caterer"), (1, "Make the seating chart"), (1, "Write vows"), (0.5, "Confirm all vendors and send day-of timeline"),
         (0.5, "Pay final vendor balances"), (0.25, "Pack for honeymoon"), (0, "Enjoy the day!")]

def build(sample, path):
    wb = Workbook()
    st = wb.active; st.title = "Start Here"
    dash = wb.create_sheet("Dashboard"); bud = wb.create_sheet("Budget"); gl = wb.create_sheet("Guest List")
    ven = wb.create_sheet("Vendors"); chk = wb.create_sheet("Checklist"); seat = wb.create_sheet("Seating"); s = wb.create_sheet("Settings")

    # ---------- Settings ----------
    K.title(s, "Settings", "Set these once. Yellow cells are yours to edit.")
    vals = [("Couple names", "Emma & Daniel" if sample else ""), ("Wedding date", dt.date(2027, 6, 12) if sample else None),
            ("Total budget", 30000 if sample else None), ("Currency label", "USD"), ("Venue", "Rosewood Gardens" if sample else "")]
    for i, (k, v) in enumerate(vals, 4):
        s.cell(row=i, column=1, value=k).font = font(bold=True); input_cell(s.cell(row=i, column=2), v if v != "" else None)
    s["B5"].number_format = DATE; s["B6"].number_format = K0
    K.section(s, "A11", "Budget categories & suggested split")
    K.hdr(s, 12, ["Category", "Suggested %", "Planned amount"])
    for i in range(15):
        r = 13 + i
        input_cell(s.cell(row=r, column=1, value=CATS[i][0] if i < len(CATS) else None))
        input_cell(s.cell(row=r, column=2, value=CATS[i][1] if i < len(CATS) else None)); s.cell(row=r, column=2).number_format = "0%"
        c = s.cell(row=r, column=3, value=f'=IF(A{r}="","",B{r}*$B$6)'); c.number_format = MONEY; c.border = B; c.fill = fill("F3F4F6"); c.font = font()
    s["A28"] = "Total"; s["B28"] = "=SUM(B13:B27)"; s["C28"] = "=SUM(C13:C27)"
    s["B28"].number_format = "0%"; s["C28"].number_format = MONEY
    for a in ("A28", "B28", "C28"): s[a].font = font(bold=True)
    s["D28"] = '=IF(ROUND(B28,4)<>1,"Split adds up to "&TEXT(B28,"0%")&" - aim for 100%","")'; s["D28"].font = Font(name=F, bold=True, color="B91C1C")
    s["A30"] = "Split based on typical wedding budget shares - adjust to your priorities."; s["A30"].font = font(italic=True, size=9)
    K.section(s, "F11", "Guest groups"); K.hdr(s, 12, ["Group"], col0=6)
    groups = ["Bride's family", "Groom's family", "Bride's friends", "Groom's friends", "Mutual friends", "Work", "Plus-ones", "", "", ""]
    for i, g in enumerate(groups): input_cell(s.cell(row=13 + i, column=6, value=g or None))
    K.section(s, "H11", "Meal choices"); K.hdr(s, 12, ["Meal"], col0=8)
    for i, m in enumerate(["Beef", "Chicken", "Fish", "Vegetarian", "Vegan", "Kids meal", "", ""]): input_cell(s.cell(row=13 + i, column=8, value=m or None))
    s["J12"] = "RSVP"; s["K12"] = "Vendor status"; s["L12"] = "Yes / No"
    for c in ("J12", "K12", "L12"): s[c].font = font(bold=True)
    for i, v in enumerate(["Pending", "Attending", "Declined"]): s.cell(row=13 + i, column=10, value=v)
    for i, v in enumerate(["Researching", "Quote received", "Booked", "Paid in full", "Not used"]): s.cell(row=13 + i, column=11, value=v)
    for i, v in enumerate(["Yes", "No"]): s.cell(row=13 + i, column=12, value=v)
    for col, w in zip("ABCDEFGHIJKL", [26, 16, 16, 3, 3, 18, 3, 14, 3, 12, 16, 9]): s.column_dimensions[col].width = w
    WD = "Settings!$B$5"; BUDGET = "Settings!$B$6"; CATR = "Settings!$A$13:$A$27"

    # ---------- Budget ----------
    K.title(bud, "Budget & Payments", "One row per item. Estimated = what you expect; Actual = the final price; Paid = what you've paid so far.")
    K.hdr(bud, 4, ["Category", "Item", "Vendor", "Estimated cost", "Actual cost", "Paid so far", "Next payment due", "Notes", "Cost used", "Balance due", "Status"],
          [24, 26, 20, 13, 13, 13, 14, 24, 12, 12, 12])
    style_range(bud, 5, 4 + BR, 1, 11, inputs=range(1, 9), fmt={4: MONEY, 5: MONEY, 6: MONEY, 7: DATE, 9: MONEY, 10: MONEY})
    items = []
    if sample:
        items = [("Venue & rentals", "Venue hire", "Rosewood Gardens", 8000, 8500, 4000, dt.date(2027, 5, 1)),
                 ("Venue & rentals", "Chairs & tables", "PartyRent", 900, 850, 850, None),
                 ("Catering & cake", "Dinner (110 guests)", "Golden Fork Catering", 6600, 7150, 2000, dt.date(2027, 5, 29)),
                 ("Catering & cake", "Wedding cake", "Sugar & Bloom", 650, None, 150, dt.date(2027, 6, 1)),
                 ("Catering & cake", "Bar package", "Golden Fork Catering", 1500, None, 0, dt.date(2027, 5, 29)),
                 ("Photography & video", "Photographer", "Lens & Light", 2800, 2950, 1000, dt.date(2027, 6, 1)),
                 ("Photography & video", "Videographer", "MotionStory", 1200, None, 300, dt.date(2027, 6, 1)),
                 ("Attire & beauty", "Wedding dress", "Bella Bridal", 1600, 1750, 1750, None),
                 ("Attire & beauty", "Suit", "Tailor & Co", 700, 640, 320, dt.date(2027, 4, 15)),
                 ("Attire & beauty", "Hair & makeup", "Glow Studio", 450, None, 100, dt.date(2027, 6, 12)),
                 ("Flowers & decor", "Bouquets & centrepieces", "Petal Lane", 1900, 2100, 500, dt.date(2027, 5, 15)),
                 ("Music & entertainment", "DJ", "DJ Sam", 1100, 1100, 300, dt.date(2027, 6, 1)),
                 ("Stationery & invitations", "Save-the-dates & invitations", "Paper Moon", 500, 460, 460, None),
                 ("Rings", "Wedding bands", "Gold Street", 1100, 1240, 1240, None),
                 ("Officiant & ceremony", "Officiant fee", "Rev. Miller", 350, 350, 0, dt.date(2027, 6, 12)),
                 ("Transportation", "Vintage car", "Classic Rides", 450, None, 100, dt.date(2027, 6, 5)),
                 ("Favors & gifts", "Guest favors", "Etsy", 300, None, 0, None)]
    for i, it in enumerate(items):
        for j, v in enumerate(it):
            bud.cell(row=5 + i, column=[1, 2, 3, 4, 5, 6, 7][j], value=v)
    for r in range(5, 5 + BR):
        bud[f"I{r}"] = f'=IF(B{r}="","",IF(E{r}<>"",E{r},D{r}))'
        bud[f"J{r}"] = f'=IF(B{r}="","",MAX(0,I{r}-F{r}))'
        bud[f"K{r}"] = f'=IF(B{r}="","",IF(J{r}=0,"Paid",IF(AND(G{r}<>"",G{r}<TODAY()),"Overdue",IF(AND(G{r}<>"",G{r}-TODAY()<=30),"Due soon","Open"))))'
    dv = DataValidation(type="list", formula1=f"={CATR}", allow_blank=True); bud.add_data_validation(dv); dv.add(f"A5:A{4+BR}")
    rg = f"K5:K{4+BR}"
    bud.conditional_formatting.add(rg, CellIsRule(operator="equal", formula=['"Paid"'], fill=fill("D1FAE5")))
    bud.conditional_formatting.add(rg, CellIsRule(operator="equal", formula=['"Overdue"'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    bud.conditional_formatting.add(rg, CellIsRule(operator="equal", formula=['"Due soon"'], fill=fill("FEF3C7")))
    bud.freeze_panes = "C5"
    BRg = lambda c: f"Budget!${c}$5:${c}${4+BR}"

    # ---------- Guest List ----------
    K.title(gl, "Guest List & RSVPs", "One row per invitation (household). Count adults and children; set RSVP when they reply.")
    K.hdr(gl, 4, ["Guest / household name", "Group", "Email or phone", "Address", "Adults invited", "Children invited", "Save-the-date sent", "Invitation sent",
                  "RSVP", "Adults attending", "Children attending", "Meal choice", "Dietary notes", "Table #", "Gift received", "Thank-you sent",
                  "Total invited", "Total attending"],
          [24, 16, 22, 26, 9, 9, 10, 10, 11, 10, 10, 13, 16, 8, 18, 10, 9, 9])
    style_range(gl, 5, 4 + GR, 1, 18, inputs=range(1, 17))
    YN = "Settings!$L$13:$L$14"
    for ref, f1 in ((f"B5:B{4+GR}", "=Settings!$F$13:$F$22"), (f"I5:I{4+GR}", "=Settings!$J$13:$J$15"), (f"L5:L{4+GR}", "=Settings!$H$13:$H$20"),
                    (f"G5:H{4+GR}", f"={YN}"), (f"P5:P{4+GR}", f"={YN}")):
        dv = DataValidation(type="list", formula1=f1, allow_blank=True); gl.add_data_validation(dv); dv.add(ref)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(TABLES), allow_blank=True); gl.add_data_validation(dv); dv.add(f"N5:N{4+GR}")
    for r in range(5, 5 + GR):
        gl[f"Q{r}"] = f'=IF(A{r}="","",N(E{r})+N(F{r}))'
        gl[f"R{r}"] = f'=IF(A{r}="","",IF(I{r}="Attending",IF(AND(J{r}="",K{r}=""),Q{r},N(J{r})+N(K{r})),0))'
    gl.conditional_formatting.add(f"I5:I{4+GR}", CellIsRule(operator="equal", formula=['"Attending"'], fill=fill("D1FAE5")))
    gl.conditional_formatting.add(f"I5:I{4+GR}", CellIsRule(operator="equal", formula=['"Declined"'], fill=fill("FEE2E2")))
    gl.conditional_formatting.add(f"I5:I{4+GR}", CellIsRule(operator="equal", formula=['"Pending"'], fill=fill("FEF3C7")))
    gl.freeze_panes = "B5"
    if sample:
        random.seed(5)
        first = ["Anna", "Ben", "Clara", "David", "Ella", "Frank", "Grace", "Henry", "Isla", "Jack", "Kate", "Leo", "Maya", "Nate", "Olive", "Paul", "Quinn",
                 "Rosa", "Sam", "Tara", "Uma", "Victor", "Wendy", "Xavier", "Yara", "Zoe"]
        last = ["Smith", "Johnson", "Clark", "Lewis", "Walker", "Hall", "Young", "King", "Wright", "Scott", "Green", "Baker", "Adams", "Nelson", "Hill", "Carter"]
        meals = ["Beef", "Chicken", "Fish", "Vegetarian", "Chicken", "Beef"]
        for i in range(58):
            r = 5 + i; adults = random.choice([1, 2, 2, 2]); kids = random.choice([0, 0, 0, 1, 2]) if adults == 2 else 0
            gl[f"A{r}"] = f"{random.choice(first)} {random.choice(last)}" + (" & family" if kids else (" +1" if adults == 2 else ""))
            gl[f"B{r}"] = random.choice(groups[:7]); gl[f"C{r}"] = f"guest{i+1}@example.com"
            gl[f"E{r}"] = adults; gl[f"F{r}"] = kids or None; gl[f"G{r}"] = "Yes"; gl[f"H{r}"] = "Yes"
            rs = random.choices(["Attending", "Declined", "Pending"], [0.7, 0.12, 0.18])[0]; gl[f"I{r}"] = rs
            if rs == "Attending":
                gl[f"L{r}"] = random.choice(meals); gl[f"N{r}"] = 1 + i // 5
                if random.random() < 0.15: gl[f"J{r}"] = max(1, adults - 1)
                if random.random() < 0.25: gl[f"O{r}"] = random.choice(["Gift card", "Kitchen set", "Cash", "Vase"]); gl[f"P{r}"] = random.choice(["Yes", "No"])
                if random.random() < 0.1: gl[f"M{r}"] = random.choice(["Gluten-free", "Nut allergy", "No pork"])

    # ---------- Vendors ----------
    K.title(ven, "Vendors", "Everyone you're talking to or have booked. Contract and payments in one place.")
    K.hdr(ven, 4, ["Category", "Vendor name", "Contact person", "Phone", "Email / website", "Quote", "Deposit paid", "Deposit date", "Balance due date",
                   "Status", "Contract signed", "Notes", "Balance", "Days to balance due"],
          [22, 22, 16, 14, 24, 12, 12, 12, 13, 14, 10, 26, 12, 10])
    style_range(ven, 5, 4 + VR, 1, 14, inputs=range(1, 13), fmt={6: MONEY, 7: MONEY, 8: DATE, 9: DATE, 13: MONEY})
    for ref, f1 in ((f"A5:A{4+VR}", f"={CATR}"), (f"J5:J{4+VR}", "=Settings!$K$13:$K$17"), (f"K5:K{4+VR}", f"={YN}")):
        dv = DataValidation(type="list", formula1=f1, allow_blank=True); ven.add_data_validation(dv); dv.add(ref)
    for r in range(5, 5 + VR):
        ven[f"M{r}"] = f'=IF(OR(B{r}="",F{r}=""),"",MAX(0,F{r}-G{r}))'
        ven[f"N{r}"] = f'=IF(OR(B{r}="",I{r}="",J{r}="Paid in full",J{r}="Not used"),"",I{r}-TODAY())'
    ven.conditional_formatting.add(f"N5:N{4+VR}", FormulaRule(formula=['AND(ISNUMBER(N5),N5<=30)'], fill=fill("FEF3C7"), font=Font(name=F, bold=True, color="92400E")))
    if sample:
        vs = [("Venue & rentals", "Rosewood Gardens", "Linda", "555-0110", "rosewood.example.com", 8500, 4000, dt.date(2026, 7, 3), dt.date(2027, 5, 1), "Booked", "Yes"),
              ("Catering & cake", "Golden Fork Catering", "Marco", "555-0182", "goldenfork.example.com", 8650, 2000, dt.date(2026, 8, 20), dt.date(2027, 5, 29), "Booked", "Yes"),
              ("Photography & video", "Lens & Light", "Jess", "555-0144", "lenslight.example.com", 2950, 1000, dt.date(2026, 9, 2), dt.date(2027, 6, 1), "Booked", "Yes"),
              ("Photography & video", "MotionStory", "Ravi", "555-0151", "motionstory.example.com", 1200, 300, dt.date(2026, 10, 1), dt.date(2027, 6, 1), "Booked", "Yes"),
              ("Flowers & decor", "Petal Lane", "Ava", "555-0129", "petallane.example.com", 2100, 500, dt.date(2026, 10, 10), dt.date(2027, 5, 15), "Booked", "Yes"),
              ("Music & entertainment", "DJ Sam", "Sam", "555-0177", "djsam.example.com", 1100, 300, dt.date(2026, 10, 15), dt.date(2027, 6, 1), "Booked", "No"),
              ("Catering & cake", "Sugar & Bloom", "Nina", "555-0165", "sugarbloom.example.com", 650, 150, dt.date(2026, 10, 20), dt.date(2027, 6, 1), "Booked", "Yes"),
              ("Transportation", "Classic Rides", "Tom", "555-0102", "classicrides.example.com", 450, 100, dt.date(2026, 10, 22), dt.date(2027, 6, 5), "Quote received", "No"),
              ("Music & entertainment", "String Trio Aria", "Elena", "555-0193", "aria.example.com", 900, None, None, None, "Researching", "No")]
        for i, v in enumerate(vs):
            for j, x in enumerate(v, 1): ven.cell(row=5 + i, column=j, value=x)

    # ---------- Checklist ----------
    K.title(chk, "Planning Checklist", "Due dates count back from your wedding date. Tick Done with 'Yes'. Add your own tasks at the bottom.")
    K.hdr(chk, 4, ["Months before", "Task", "Owner", "Done", "Notes", "Due date", "Status"], [10, 46, 14, 9, 30, 13, 12])
    style_range(chk, 5, 4 + TR, 1, 7, inputs=range(1, 6), fmt={6: DATE})
    for i, (mb, t) in enumerate(TASKS):
        chk[f"A{5+i}"] = mb; chk[f"B{5+i}"] = t
        if sample and mb >= 9: chk[f"D{5+i}"] = "Yes"
        if sample and mb in (8, 7): chk[f"D{5+i}"] = "Yes" if i % 2 else None
    for r in range(5, 5 + TR):
        chk[f"F{r}"] = f'=IF(OR(B{r}="",{WD}=""),"",{WD}-ROUND(N(A{r})*30.4,0))'
        chk[f"G{r}"] = f'=IF(B{r}="","",IF(D{r}="Yes","Done",IF(F{r}="","",IF(F{r}<TODAY(),"Overdue",IF(F{r}-TODAY()<=30,"This month","Upcoming")))))'
    dv = DataValidation(type="list", formula1=f"={YN}", allow_blank=True); chk.add_data_validation(dv); dv.add(f"D5:D{4+TR}")
    rg = f"G5:G{4+TR}"
    chk.conditional_formatting.add(rg, CellIsRule(operator="equal", formula=['"Done"'], fill=fill("D1FAE5")))
    chk.conditional_formatting.add(rg, CellIsRule(operator="equal", formula=['"Overdue"'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    chk.conditional_formatting.add(rg, CellIsRule(operator="equal", formula=['"This month"'], fill=fill("FEF3C7")))
    chk.conditional_formatting.add(f"A5:G{4+TR}", FormulaRule(formula=['$D5="Yes"'], font=Font(name=F, color="9CA3AF", strike=True)))
    chk.freeze_panes = "A5"

    # ---------- Seating ----------
    K.title(seat, "Seating Plan", "Set seats per table. Guests are placed with the 'Table #' column on the Guest List; this shows who sits where.")
    K.hdr(seat, 4, ["Table #", "Table name (optional)", "Seats", "Guests seated", "Free seats", "Status"], [9, 22, 8, 10, 10, 12])
    style_range(seat, 5, 4 + TABLES, 1, 6, inputs=(2, 3))
    for i in range(TABLES):
        r = 5 + i; seat[f"A{r}"] = i + 1
        seat[f"C{r}"] = 10 if (sample and i < 14) else None
        if sample and i == 0: seat[f"B{r}"] = "Head table"
        seat[f"D{r}"] = f"=SUMIFS('Guest List'!$R$5:$R${4+GR},'Guest List'!$N$5:$N${4+GR},A{r})"
        seat[f"E{r}"] = f'=IF(C{r}="","",C{r}-D{r})'
        seat[f"F{r}"] = f'=IF(AND(C{r}="",D{r}=0),"",IF(C{r}="","Set seats",IF(D{r}>C{r},"Over capacity",IF(D{r}=C{r},"Full","Space"))))'
    seat.conditional_formatting.add(f"F5:F{4+TABLES}", CellIsRule(operator="equal", formula=['"Over capacity"'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    seat.conditional_formatting.add(f"F5:F{4+TABLES}", CellIsRule(operator="equal", formula=['"Full"'], fill=fill("D1FAE5")))
    seat["H4"] = "Attending, no table yet"; seat["H4"].font = font(bold=True)
    seat["H5"] = f"=SUMIFS('Guest List'!$R$5:$R${4+GR},'Guest List'!$N$5:$N${4+GR},\"\")"; seat["H5"].font = Font(name=F, size=16, bold=True, color=MAIN)
    seat["H7"] = "Total seats set"; seat["H7"].font = font(bold=True); seat["H8"] = f"=SUM(C5:C{4+TABLES})"; seat["H8"].font = Font(name=F, size=16, bold=True, color=MAIN)
    seat.column_dimensions["G"].width = 3; seat.column_dimensions["H"].width = 24
    seat.freeze_panes = "A5"

    # ---------- Dashboard ----------
    K.title(dash, "Wedding Dashboard", "Everything here updates from your other tabs.")
    dash["A3"] = '=IF(Settings!B4="","",Settings!B4)&IF(Settings!B5="",""," · "&TEXT(Settings!B5,"dddd, mmmm d, yyyy"))&IF(Settings!B8="",""," · "&Settings!B8)'
    dash["A3"].font = Font(name=F, bold=True, size=12, color=INK)
    for c, w in zip("ABCDEFGHIJ", [24, 14, 15, 14, 15, 14, 15, 14, 3, 14]): dash.column_dimensions[c].width = w
    used = f"SUM({BRg('I')})"; paid = f"SUM({BRg('F')})"
    kp = [("Days to go", f'=IF({WD}="","",MAX(0,{WD}-TODAY()))', "0"), ("Total budget", f"={BUDGET}", K0), ("Committed cost (actual or estimate)", f"={used}", K0), ("Left to allocate", f"=N({BUDGET})-{used}", K0),
          ("Paid so far", f"={paid}", K0), ("Still to pay", f"=SUM({BRg('J')})", K0), ("Payments overdue / due in 30 days", f'=COUNTIF({BRg("K")},"Overdue")&" / "&COUNTIF({BRg("K")},"Due soon")', "@"),
          ("Checklist done", f"=IFERROR(COUNTIF(Checklist!$G$5:$G${4+TR},\"Done\")/COUNTIF(Checklist!$B$5:$B${4+TR},\"?*\"),0)", "0%")]
    pos = [(5, 1), (5, 3), (5, 5), (5, 7), (8, 1), (8, 3), (8, 5), (8, 7)]
    for (r, c), (lab, f, fmt) in zip(pos, kp): K.kpi(dash, r, c, lab, f, fmt)
    dash["D9"].number_format = K0
    dash.conditional_formatting.add("G6", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, size=18, bold=True, color="B91C1C")))
    # guests block
    K.section(dash, "A11", "Guests")
    GQ = f"'Guest List'!$Q$5:$Q${4+GR}"; GI = f"'Guest List'!$I$5:$I${4+GR}"
    gstats = [("Invited (people)", f"=SUM({GQ})"), ("Attending", f"=SUM('Guest List'!$R$5:$R${4+GR})"),
              ("Declined", f'=SUMIFS({GQ},{GI},"Declined")'), ("Awaiting reply", f'=SUMIFS({GQ},{GI},"Pending")+SUMIFS({GQ},{GI},"")'),
              ("Invitations (households)", f'=COUNTIF(\'Guest List\'!$A$5:$A${4+GR},"?*")'), ("Thank-you notes to write", f'=COUNTIFS(\'Guest List\'!$O$5:$O${4+GR},"?*",\'Guest List\'!$P$5:$P${4+GR},"<>Yes")')]
    K.hdr(dash, 12, ["Guest count", "People"])
    for i, (lab, f) in enumerate(gstats):
        r = 13 + i; dash[f"A{r}"] = lab; dash[f"B{r}"] = f
        for c in "AB": dash[f"{c}{r}"].border = B; dash[f"{c}{r}"].font = font(bold=(i == 1))
    K.hdr(dash, 12, ["Meal choice", "Count"], col0=4)
    for i in range(8):
        r = 13 + i
        dash[f"D{r}"] = f'=IF(Settings!H{13+i}="","",Settings!H{13+i})'
        dash[f"E{r}"] = f"=IF(D{r}=\"\",\"\",SUMIFS('Guest List'!$R$5:$R${4+GR},'Guest List'!$L$5:$L${4+GR},D{r}))"
        for c in "DE": dash[f"{c}{r}"].border = B; dash[f"{c}{r}"].font = font()
    dash["D21"] = "No meal picked yet"; dash["E21"] = f"=SUMIFS('Guest List'!$R$5:$R${4+GR},'Guest List'!$L$5:$L${4+GR},\"\")"
    for c in "DE": dash[f"{c}21"].border = B; dash[f"{c}21"].font = font(italic=True)
    K.hdr(dash, 12, ["By group", "Invited", "Attending"], col0=7)
    for i in range(10):
        r = 13 + i
        dash[f"G{r}"] = f'=IF(Settings!F{13+i}="","",Settings!F{13+i})'
        dash[f"H{r}"] = f"=IF(G{r}=\"\",\"\",SUMIFS({GQ},'Guest List'!$B$5:$B${4+GR},G{r}))"
        dash[f"I{r}"] = f"=IF(G{r}=\"\",\"\",SUMIFS('Guest List'!$R$5:$R${4+GR},'Guest List'!$B$5:$B${4+GR},G{r}))"
        for c in "GHI": dash[f"{c}{r}"].border = B; dash[f"{c}{r}"].font = font()
    dash.column_dimensions["I"].width = 10
    # budget by category
    K.section(dash, "A24", "Budget by category")
    K.hdr(dash, 25, ["Category", "Planned", "Estimated", "Committed", "Paid", "Still to pay", "vs plan", "% of plan used"])
    for i in range(15):
        r = 26 + i; sr = 13 + i
        dash[f"A{r}"] = f'=IF(Settings!A{sr}="","",Settings!A{sr})'
        dash[f"B{r}"] = f'=IF(A{r}="","",Settings!C{sr})'
        dash[f"C{r}"] = f'=IF(A{r}="","",SUMIFS({BRg("D")},{BRg("A")},A{r}))'
        dash[f"D{r}"] = f'=IF(A{r}="","",SUMIFS({BRg("I")},{BRg("A")},A{r}))'
        dash[f"E{r}"] = f'=IF(A{r}="","",SUMIFS({BRg("F")},{BRg("A")},A{r}))'
        dash[f"F{r}"] = f'=IF(A{r}="","",SUMIFS({BRg("J")},{BRg("A")},A{r}))'
        dash[f"G{r}"] = f'=IF(A{r}="","",B{r}-D{r})'
        dash[f"H{r}"] = f'=IF(A{r}="","",IFERROR(D{r}/B{r},0))'
        for j in range(1, 9):
            c = dash.cell(row=r, column=j); c.border = B; c.font = font(); c.number_format = "0%" if j == 8 else (MONEY if j > 1 else "General")
    dash["A41"] = "Total"
    for c in "BCDEFG": dash[f"{c}41"] = f"=SUM({c}26:{c}40)"
    dash["H41"] = "=IFERROR(D41/B41,0)"
    for j in range(1, 9):
        c = dash.cell(row=41, column=j); c.border = B; c.font = font(bold=True); c.fill = fill(SOFT); c.number_format = "0%" if j == 8 else (MONEY if j > 1 else "General")
    dash.conditional_formatting.add("G26:G41", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, bold=True, color="B91C1C")))
    dash.conditional_formatting.add("H26:H40", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1.2, color="D8A7C0"))
    ch = BarChart(); ch.type = "bar"; ch.title = "Planned vs committed"; ch.height = 9; ch.width = 16
    ch.add_data(Reference(dash, min_col=2, min_row=25, max_row=37), titles_from_data=True)
    ch.add_data(Reference(dash, min_col=4, min_row=25, max_row=37), titles_from_data=True)
    ch.set_categories(Reference(dash, min_col=1, min_row=26, max_row=37)); ch.y_axis.majorGridlines = None; ch.x_axis.scaling.orientation = "maxMin"
    for ser, col in zip(ch.series, ("D8A7C0", MAIN)):
        ser.graphicalProperties = GraphicalProperties(solidFill=col); ser.graphicalProperties.line.solidFill = col
    dash.add_chart(ch, "A44")
    pc = PieChart(); pc.title = "RSVP status"; pc.height = 9; pc.width = 10
    dash["K12"] = "RSVP"; dash["L12"] = "People"
    for i, (lab, src) in enumerate([("Attending", "B14"), ("Declined", "B15"), ("Awaiting", "B16")]):
        dash[f"K{13+i}"] = lab; dash[f"L{13+i}"] = f"={src}"
    for a in ("K12", "L12", "K13", "L13", "K14", "L14", "K15", "L15"): dash[a].font = font(size=8, color="FFFFFF")
    pc.add_data(Reference(dash, min_col=12, min_row=12, max_row=15), titles_from_data=True)
    pc.set_categories(Reference(dash, min_col=11, min_row=13, max_row=15))
    pc.dataLabels = DataLabelList(); pc.dataLabels.showVal = True
    from openpyxl.chart.series import DataPoint
    for idx, col in enumerate(("6BA368", "D9534F", "F0AD4E")):
        pt = DataPoint(idx=idx); pt.graphicalProperties = GraphicalProperties(solidFill=col); pc.series[0].dPt.append(pt)
    dash.add_chart(pc, "F44")

    # ---------- Start Here ----------
    lines = [("How to start", True),
             ("1.  Settings tab: type your names, wedding date, total budget and currency. Adjust the budget split by category if you like (it should add up to 100%).", False),
             ("2.  Budget tab: add each item with an estimate. When you book, type the actual cost, what you've paid and the next payment date.", False),
             ("3.  Guest List tab: one row per household with adults and children invited. Set RSVP, meal and table number as replies come in.", False),
             ("4.  Vendors tab: keep contacts, quotes, deposits and balance due dates together.", False),
             ("5.  Checklist tab: 32 tasks with due dates counted back from your wedding date. Mark Done with 'Yes'.", False),
             ("6.  Seating tab: set the seats per table; it counts guests per table from the Guest List and warns when a table is over capacity.", False),
             ("", False),
             ("Colour key", True),
             ("Yellow cells are inputs: type here.  Grey cells calculate themselves: please don't overwrite them.", False),
             ("Status colours: green = done / paid / attending, amber = due soon / pending, red = overdue / declined / over capacity.", False),
             ("", False),
             ("Good to know", True),
             ("•  Committed cost uses the actual cost when you've typed one, otherwise the estimate.", False),
             ("•  Attending people = adults + children attending. If you leave those blank for an 'Attending' household, everyone invited is counted.", False),
             ("•  Google Sheets: open drive.google.com, upload this file, then choose Open with > Google Sheets.", False),
             ("•  Room for 120 budget items, 400 households, 60 vendors, 120 tasks and 40 tables.", False)]
    if sample:
        lines += [("", False), ("This copy contains sample data", True),
                  ("The names, guests, vendors and amounts are made-up examples. Clear the yellow cells before entering your own, or use the blank copy in your download.", False)]
    K.start(st, "Wedding Planner & Budget", "Budget, guest list & RSVPs, vendors, checklist and seating in one file. Works in Excel and Google Sheets.", lines)
    finish(wb, sample, path)

META = {
    "name": "Wedding Planner & Budget Spreadsheet – Excel & Google Sheets Template",
    "file_base": "Wedding-Planner-Budget",
    "price": 19,
    "color": "#8B4D6B",
    "title": "Wedding Planner & Budget",
    "tagline": "Budget · guest list & RSVP · vendors · checklist · seating — Excel & Google Sheets",
    "intro": "Plan your whole wedding in one calm, organised file. Keep the budget on track, see RSVPs and meal counts at a glance, and never miss a vendor payment or planning task.",
    "features": [
        "Dashboard: countdown, budget vs committed cost, paid and still to pay, overdue payments, checklist progress, guest and meal counts",
        "Budget: 12 categories with a suggested split of your total budget, estimated vs actual cost, payments and due dates",
        "Guest list & RSVPs: households, adults and children, meal choice, dietary notes, table number, gifts and thank-you notes",
        "Vendors: contacts, quotes, deposits, balance due dates and contract status, with reminders 30 days ahead",
        "12-month planning checklist with 32 tasks, due dates counted back from your wedding date",
        "Seating plan: seats per table, guests seated and over-capacity warnings",
        "Any currency, no macros; blank copy plus a filled-in example",
    ],
    "tags": ["wedding planner", "wedding budget", "wedding spreadsheet", "guest list template", "rsvp tracker", "wedding checklist", "google sheets",
             "excel template", "seating chart", "bride to be", "wedding organizer", "vendor tracker", "engagement gift"],
}
