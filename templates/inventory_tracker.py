"""Inventory & Stock Tracker for small shops, with reorder alerts (sample + blank versions).

Products (SKU, cost, price, opening stock, reorder point/qty, supplier) + Stock Movements log
(purchase, sale, customer return, return to supplier, adjustment, damaged/lost)
-> current stock, status (OK / Reorder / Out of stock), stock value, days of stock left, sales and profit
-> Reorder List (automatic, with suggested quantities and supplier), Suppliers, Dashboard.
"""
import random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as L
from _common import Kit, F, INK, B, MONEY, DATE, font, fill, style_range, input_cell, finish

MAIN = "1E3A8A"; SOFT = "EAF0FB"; ACC = "60A5FA"
K = Kit(MAIN, SOFT)
NP = 300; P0 = 6; P1 = P0 + NP - 1
NM = 3000; M0 = 6; M1 = M0 + NM - 1
NS = 30; S0 = 6; S1 = S0 + NS - 1
NRL = 100
TYPES = [("Purchase (stock in)", 1), ("Sale", -1), ("Customer return", 1), ("Return to supplier", -1),
         ("Adjustment (+ or -)", 1), ("Damaged / lost", -1)]
SM = "'Stock Movements'"
def mv(c): return f"{SM}!${c}${M0}:${c}${M1}"
def pr(c): return f"Products!${c}${P0}:${c}${P1}"
CUR = '#,##0.00;-#,##0.00;"-"'; QTY = '#,##0;-#,##0;0'

SUPPLIERS = [("Wick & Wax Supply Co.", "Laura Benson", "orders@wickwax.example", "+1 555 0101", "", 7),
             ("Pure Botanicals Ltd", "Sam Ortiz", "sales@purebotanicals.example", "+1 555 0144", "", 10),
             ("Glass & Jar Wholesale", "Mia Chen", "mia@glassjar.example", "+1 555 0190", "", 5),
             ("PackRight Packaging", "Tom Reyes", "hello@packright.example", "+1 555 0172", "", 4)]
PRODUCTS = [  # sku, name, category, supplier idx, cost, price, opening, reorder point, reorder qty, location, daily demand
    ("CN-LAV-08", "Lavender Soy Candle 8oz", "Candles", 0, 6.2, 24, 60, 20, 60, "Shelf A1", 0.9),
    ("CN-VAN-08", "Vanilla Bean Soy Candle 8oz", "Candles", 0, 6.2, 24, 55, 20, 60, "Shelf A1", 0.8),
    ("CN-CED-12", "Cedarwood Candle 12oz", "Candles", 0, 8.9, 32, 30, 12, 36, "Shelf A2", 0.45),
    ("CN-CIT-04", "Citrus Travel Tin 4oz", "Candles", 0, 3.1, 12, 80, 25, 80, "Shelf A3", 1.1),
    ("CN-PUM-08", "Pumpkin Spice Candle 8oz", "Candles", 0, 6.4, 26, 40, 15, 48, "Shelf A2", 0.75),
    ("SP-OAT-01", "Oat & Honey Soap Bar", "Soap", 1, 2.1, 9, 90, 30, 100, "Shelf B1", 1.3),
    ("SP-CHA-01", "Charcoal Detox Soap Bar", "Soap", 1, 2.3, 10, 70, 25, 80, "Shelf B1", 0.9),
    ("SP-ROS-01", "Rose Clay Soap Bar", "Soap", 1, 2.4, 10, 60, 25, 80, "Shelf B1", 0.7),
    ("SP-LIQ-250", "Lemongrass Hand Soap 250ml", "Soap", 1, 3.6, 14, 40, 15, 48, "Shelf B2", 0.55),
    ("HF-DIF-100", "Reed Diffuser 100ml – Fig", "Home fragrance", 2, 7.5, 28, 25, 10, 30, "Shelf C1", 0.35),
    ("HF-DIF-200", "Reed Diffuser 200ml – Amber", "Home fragrance", 2, 11.0, 42, 18, 8, 24, "Shelf C1", 0.22),
    ("HF-SPR-100", "Room Spray 100ml – Linen", "Home fragrance", 2, 4.2, 18, 35, 12, 36, "Shelf C2", 0.5),
    ("HF-MLT-06", "Wax Melts 6-pack", "Home fragrance", 0, 2.6, 11, 70, 20, 60, "Shelf C3", 0.85),
    ("BT-SLT-300", "Bath Salts 300g – Eucalyptus", "Bath", 1, 3.9, 16, 40, 15, 40, "Shelf D1", 0.45),
    ("BT-BOM-01", "Bath Bomb – Lavender", "Bath", 1, 1.6, 7, 120, 40, 120, "Shelf D2", 1.6),
    ("BT-OIL-50", "Body Oil 50ml – Sweet Almond", "Bath", 1, 5.2, 22, 22, 10, 30, "Shelf D1", 0.3),
    ("BT-SCR-200", "Sugar Scrub 200g", "Bath", 1, 4.4, 19, 30, 12, 36, "Shelf D1", 0.4),
    ("AC-WIC-TR", "Wick Trimmer", "Accessories", 3, 3.0, 14, 20, 8, 24, "Drawer E1", 0.15),
    ("AC-SNF-01", "Candle Snuffer – Brass", "Accessories", 3, 4.1, 18, 15, 6, 20, "Drawer E1", 0.1),
    ("AC-MAT-01", "Long Matches Jar", "Accessories", 3, 2.2, 9, 40, 15, 40, "Drawer E2", 0.4),
    ("GS-SPA-01", "Spa Gift Box", "Gift sets", 3, 14.5, 49, 15, 6, 20, "Shelf F1", 0.25),
    ("GS-CAN-03", "Candle Trio Gift Set", "Gift sets", 3, 17.0, 58, 12, 5, 15, "Shelf F1", 0.2),
    ("GS-MIN-01", "Mini Self-Care Kit", "Gift sets", 3, 8.0, 29, 20, 8, 24, "Shelf F2", 0.3),
    ("PK-BOX-SM", "Shipping Box Small (packaging)", "Packaging", 3, 0.45, 0, 300, 100, 500, "Back room", 3.2)]
CATS = ["Candles", "Soap", "Home fragrance", "Bath", "Accessories", "Gift sets", "Packaging"]


def simulate():
    random.seed(21)
    rows = []; stock = {p[0]: p[6] for p in PRODUCTS}; pending = []
    d = dt.date(2026, 8, 1); end = dt.date(2026, 10, 30); order_no = 1040; po = 300
    while d <= end:
        season = 1 + max(0, (d - dt.date(2026, 9, 15)).days) / 60          # autumn pick-up
        for (sku, eta, q, cost) in [x for x in pending if x[1] == d]:
            rows.append((d, sku, "Purchase (stock in)", q, cost, f"PO-{po}", "")); stock[sku] += q; po += 1
            pending.remove((sku, eta, q, cost))
        for p in PRODUCTS:
            sku = p[0]; dem = p[10] * season * (0.6 if d.weekday() < 4 else 1.5)
            if sku == "PK-BOX-SM": continue
            n = sum(1 for _ in range(4) if random.random() < dem / 4)
            n = min(n, stock[sku])
            if n:
                price = p[5] if random.random() > 0.12 else round(p[5] * 0.85, 2)
                rows.append((d, sku, "Sale", n, price if price != p[5] else None, f"#{order_no}", "15% promo" if price != p[5] else ""))
                stock[sku] -= n; order_no += 1
        # reorders (some deliberately late at the end so the sample shows alerts)
        for p in PRODUCTS:
            sku = p[0]
            if stock[sku] <= p[7] and not any(x[0] == sku for x in pending) and d < (dt.date(2026, 9, 20) if sku == "GS-CAN-03" else dt.date(2026, 10, 18)):
                lead = SUPPLIERS[p[3]][5]
                pending.append((sku, d + dt.timedelta(days=lead), p[8], None))
        if d == dt.date(2026, 9, 3): rows.append((d, "BT-BOM-01", "Damaged / lost", 4, None, "", "Crushed in delivery"))
        if d == dt.date(2026, 9, 21): rows.append((d, "CN-VAN-08", "Customer return", 1, None, "#1102", "Changed mind")); stock["CN-VAN-08"] += 1
        if d == dt.date(2026, 10, 1): rows.append((d, "SP-ROS-01", "Adjustment (+ or -)", -2, None, "Stock count", "Count was 2 lower"))
        if d == dt.date(2026, 10, 9): rows.append((d, "HF-SPR-100", "Return to supplier", 3, None, "RMA-17", "Leaking caps"))
        d += dt.timedelta(days=1)
    rows.sort(key=lambda r: r[0])
    return rows


def build(sample, path):
    wb = Workbook()
    st = wb.active; st.title = "Start Here"
    dash = wb.create_sheet("Dashboard"); ps = wb.create_sheet("Products"); sm = wb.create_sheet("Stock Movements")
    rl = wb.create_sheet("Reorder List"); sp = wb.create_sheet("Suppliers"); cc = wb.create_sheet("Calc"); s = wb.create_sheet("Settings")

    # ---------- Settings ----------
    K.title(s, "Settings", "Yellow cells are yours to edit.")
    rows = [("Currency label", "USD", None), ("Today / as-of date", dt.date(2026, 10, 30) if sample else "=TODAY()", DATE),
            ("Sales period for 'days of stock left' (days)", 30, "0"), ("Shop name", "Ember & Bloom Home" if sample else None, None)]
    for i, (k_, v, fmt) in enumerate(rows, 4):
        s.cell(row=i, column=1, value=k_).font = font(bold=True)
        c = s.cell(row=i, column=2); input_cell(c, v)
        if fmt: c.number_format = fmt
    s["C5"] = "Keep =TODAY() so 'this month' and 'last 30 days' stay current." + (" (The sample uses a fixed date.)" if sample else "")
    s["C5"].font = font(italic=True, size=9)
    calc = [("Sales period starts", "=B5-B6+1", DATE), ("This month starts", "=DATE(YEAR(B5),MONTH(B5),1)", DATE)]
    for i, (k_, f, fmt) in enumerate(calc, 9):
        s.cell(row=i, column=1, value=k_).font = font(italic=True)
        c = s.cell(row=i, column=2, value=f); c.number_format = fmt; c.border = B; c.fill = fill("F3F4F6"); c.font = font()
    K.hdr(s, 4, ["Categories"], col0=8)
    for i in range(15):
        c = s.cell(row=5 + i, column=8, value=CATS[i] if (sample and i < len(CATS)) else (["General", "Other"][i] if (not sample and i < 2) else None)); input_cell(c)
    K.hdr(s, 4, ["Movement type", "Effect on stock"], col0=10)
    for i, (t, sg) in enumerate(TYPES):
        s.cell(row=5 + i, column=10, value=t).font = font(size=9); s.cell(row=5 + i, column=11, value=sg).font = font(size=9)
        s.cell(row=5 + i, column=10).border = B; s.cell(row=5 + i, column=11).border = B
    for c, w in zip("ABCDEFGHIJK", [40, 18, 3, 3, 3, 3, 3, 18, 3, 22, 10]): s.column_dimensions[c].width = w
    s["J12"] = "Adjustment: type a negative quantity to remove stock (e.g. -2 after a stock count)."; s["J12"].font = font(italic=True, size=9)
    AS, FROM, MS = "Settings!$B$5", "Settings!$B$9", "Settings!$B$10"

    # ---------- Suppliers ----------
    K.title(sp, "Suppliers", "Who you buy from. Grey columns calculate themselves.")
    K.hdr(sp, 5, ["Supplier", "Contact person", "Email", "Phone", "Website", "Lead time (days)", "Notes", "Products", "Items to reorder", "Reorder cost"],
          [28, 16, 30, 14, 20, 9, 20, 9, 9, 12])
    style_range(sp, S0, S1, 1, 10, inputs=range(1, 8), fmt={10: CUR})
    if sample:
        for i, row in enumerate(SUPPLIERS):
            for j, v in enumerate(row, 1):
                if v != "": sp.cell(row=S0 + i, column=j, value=v)
    for r in range(S0, S1 + 1):
        sp[f"H{r}"] = f'=IF(A{r}="","",COUNTIF({pr("D")},A{r}))'
        sp[f"I{r}"] = f'=IF(A{r}="","",COUNTIFS({pr("D")},A{r},{pr("W")},">0"))'
        sp[f"J{r}"] = f'=IF(A{r}="","",SUMIFS({pr("Y")},{pr("D")},A{r}))'
    sp.freeze_panes = "B6"

    # ---------- Products ----------
    K.title(ps, "Products", "One row per product (SKU). Yellow = type here. Opening stock = what you had when you started using this file.")
    ps["A3"] = '="Prices in "&Settings!B4&" · status: Out of stock / Reorder (at or below reorder point) / OK"'; ps["A3"].font = font(bold=True)
    heads = ["SKU", "Product", "Category", "Supplier", "Unit cost", "Sale price", "Opening stock", "Reorder point", "Reorder qty", "Location", "Notes",
             "Stock in", "Stock out", "In stock", "Status", "Stock value (cost)", "Retail value", "Sold (period)", "Days of stock left", "Margin %",
             "Units sold (all)", "Revenue (all)", "Reorder key", "Suggested order", "Order cost", "Top-seller key"]
    K.hdr(ps, 5, heads, [12, 28, 14, 22, 9, 9, 9, 9, 9, 11, 14, 8, 8, 8, 12, 11, 11, 8, 9, 8, 9, 11, 7, 9, 10, 9])
    style_range(ps, P0, P1, 1, 26, inputs=range(1, 12), fmt={5: CUR, 6: CUR, 7: QTY, 8: QTY, 9: QTY, 12: QTY, 13: QTY, 14: QTY, 16: CUR, 17: CUR,
                                                            18: QTY, 19: '0;-0;""', 20: '0%;-0%;""', 21: QTY, 22: CUR, 25: CUR, 26: "0.000"})
    if sample:
        for i, p in enumerate(PRODUCTS):
            r = P0 + i
            for j, v in enumerate((p[0], p[1], p[2], SUPPLIERS[p[3]][0], p[4], p[5] or None, p[6], p[7], p[8], p[9]), 1):
                if v is not None: ps.cell(row=r, column=j, value=v)
        ps[f"K{P0 + len(PRODUCTS) - 1}"] = "Not for sale – used per order"
    sku, sgn = mv("B"), mv("I")
    for r in range(P0, P1 + 1):
        on = f'A{r}=""'
        ps[f"L{r}"] = f'=IF({on},0,SUMIFS({sgn},{sku},A{r},{sgn},">0"))'
        ps[f"M{r}"] = f'=IF({on},0,-SUMIFS({sgn},{sku},A{r},{sgn},"<0"))'
        ps[f"N{r}"] = f"=IF({on},0,N(G{r})+L{r}-M{r})"
        ps[f"O{r}"] = f'=IF({on},"",IF(N{r}<=0,"Out of stock",IF(N{r}<=N(H{r}),"Reorder","OK")))'
        ps[f"P{r}"] = f"=MAX(0,N{r})*N(E{r})"
        ps[f"Q{r}"] = f"=MAX(0,N{r})*N(F{r})"
        ps[f"R{r}"] = f'=IF({on},0,SUMIFS({mv("D")},{sku},A{r},{mv("C")},"Sale",{mv("A")},">="&{FROM},{mv("A")},"<="&{AS}))'
        ps[f"S{r}"] = f'=IF(OR({on},R{r}=0),"",MAX(0,N{r})/(R{r}/Settings!$B$6))'
        ps[f"T{r}"] = f'=IF(N(F{r})=0,"",(F{r}-N(E{r}))/F{r})'
        ps[f"U{r}"] = f'=IF({on},0,SUMIFS({mv("D")},{sku},A{r},{mv("C")},"Sale")-SUMIFS({mv("D")},{sku},A{r},{mv("C")},"Customer return"))'
        ps[f"V{r}"] = f'=IF({on},0,SUMIFS({mv("J")},{sku},A{r}))'
        ps[f"W{r}"] = f'=IF(OR(O{r}="Reorder",O{r}="Out of stock"),ROW(),0)'
        ps[f"X{r}"] = f"=IF(W{r}=0,0,IF(N(I{r})>0,I{r},MAX(1,2*N(H{r})-MAX(0,N{r}))))"
        ps[f"Y{r}"] = f"=X{r}*N(E{r})"
        ps[f"Z{r}"] = f"=IF(R{r}>0,R{r}+ROW()/100000,0)"
        ps[f"O{r}"].alignment = Alignment(horizontal="center")
    for ref, f1 in ((f"C{P0}:C{P1}", "=Settings!$H$5:$H$19"), (f"D{P0}:D{P1}", f"=Suppliers!$A${S0}:$A${S1}")):
        dv = DataValidation(type="list", formula1=f1, allow_blank=True); ps.add_data_validation(dv); dv.add(ref)
    ps.conditional_formatting.add(f"O{P0}:O{P1}", CellIsRule(operator="equal", formula=['"Out of stock"'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    ps.conditional_formatting.add(f"O{P0}:O{P1}", CellIsRule(operator="equal", formula=['"Reorder"'], fill=fill("FEF3C7"), font=Font(name=F, bold=True, color="92400E")))
    ps.conditional_formatting.add(f"O{P0}:O{P1}", CellIsRule(operator="equal", formula=['"OK"'], fill=fill("D1FAE5"), font=Font(name=F, color="166534")))
    ps.conditional_formatting.add(f"A{P0}:A{P1}", FormulaRule(formula=[f'AND(A{P0}<>"",COUNTIF($A${P0}:$A${P1},A{P0})>1)'], fill=fill("FCA5A5")))
    ps.freeze_panes = "C6"

    # ---------- Stock Movements ----------
    K.title(sm, "Stock Movements", "Every time stock comes in or goes out, add a row. Quantity is always positive (except adjustments).")
    sm["A3"] = '="Unit price is optional: leave it empty to use the price on the Products tab (cost for purchases, sale price for sales)."'
    sm["A3"].font = font(bold=True, size=9)
    heads = ["Date", "SKU", "Type", "Quantity", "Unit price (optional)", "Order / reference", "Note", "Product", "Stock change", "Sales revenue",
             "Cost of goods sold", "Month", "Purchase cost", "Product #", "Sign"]
    K.hdr(sm, 5, heads, [11, 13, 20, 9, 10, 13, 20, 28, 9, 11, 11, 9, 11, 8, 6])
    style_range(sm, M0, M1, 1, 15, inputs=range(1, 8), fmt={1: DATE, 4: QTY, 5: CUR, 9: '+#,##0;-#,##0;0', 10: CUR, 11: CUR, 12: "mmm yy", 13: CUR})
    if sample:
        for i, row in enumerate(simulate()):
            for j, v in enumerate(row, 1):
                if v not in (None, ""): sm.cell(row=M0 + i, column=j, value=v)
    for r in range(M0, M1 + 1):
        sm[f"N{r}"] = f'=IF(B{r}="",0,IFERROR(MATCH(B{r},{pr("A")},0),0))'
        sm[f"O{r}"] = f'=IF(C{r}="",0,IFERROR(INDEX(Settings!$K$5:$K$10,MATCH(C{r},Settings!$J$5:$J$10,0)),0))'
        sm[f"H{r}"] = f'=IF(B{r}="","",IF(N{r}=0,"SKU not found",INDEX({pr("B")},N{r})))'
        sm[f"I{r}"] = f"=N(D{r})*O{r}"
        price = f"IF(N(E{r})>0,E{r},INDEX({pr('F')},N{r}))"; cost = f"INDEX({pr('E')},N{r})"
        sm[f"J{r}"] = f'=IF(N{r}=0,0,IF(C{r}="Sale",N(D{r})*{price},IF(C{r}="Customer return",-N(D{r})*{price},0)))'
        sm[f"K{r}"] = f'=IF(N{r}=0,0,IF(C{r}="Sale",N(D{r})*N({cost}),IF(C{r}="Customer return",-N(D{r})*N({cost}),0)))'
        sm[f"L{r}"] = f'=IF(N(A{r})=0,0,DATE(YEAR(A{r}),MONTH(A{r}),1))'
        sm[f"M{r}"] = f'=IF(OR(N{r}=0,C{r}<>"Purchase (stock in)"),0,N(D{r})*IF(N(E{r})>0,E{r},N({cost})))'
    dv = DataValidation(type="list", formula1=f"={pr('A')}", allow_blank=True); sm.add_data_validation(dv); dv.add(f"B{M0}:B{M1}")
    dv = DataValidation(type="list", formula1="=Settings!$J$5:$J$10", allow_blank=True); sm.add_data_validation(dv); dv.add(f"C{M0}:C{M1}")
    sm.conditional_formatting.add(f"H{M0}:H{M1}", FormulaRule(formula=[f'H{M0}="SKU not found"'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    sm.freeze_panes = "C6"

    # ---------- Reorder List ----------
    K.title(rl, "Reorder List", "Products at or below their reorder point. Updates by itself – send it to your suppliers.")
    rl["A3"] = f'="As of "&TEXT({AS},"dd mmm yyyy")&"  ·  "&COUNTIF({pr("W")},">0")&" items  ·  estimated order cost "&Settings!B4&" "&TEXT(SUM({pr("Y")}),"#,##0.00")'
    rl["A3"].font = font(bold=True)
    K.hdr(rl, 5, ["SKU", "Product", "Status", "In stock", "Reorder point", "Sold (period)", "Suggested order qty", "Unit cost", "Order cost",
                  "Supplier", "Email", "Lead time (days)"], [12, 28, 12, 8, 9, 9, 10, 9, 10, 24, 28, 9])
    for k in range(NRL):
        r = 6 + k
        rw = f"SMALL({pr('W')},COUNTIF({pr('W')},0)+{k + 1})"
        rl[f"A{r}"] = f'=IFERROR(INDEX(Products!$A:$A,{rw}),"")'
        for c, src in zip("BCDEFGHI", "BONHRXEY"):
            rl[f"{c}{r}"] = f'=IF($A{r}="","",INDEX(Products!${src}:${src},{rw}))'
        rl[f"J{r}"] = f'=IF($A{r}="","",INDEX(Products!$D:$D,{rw}))'
        rl[f"K{r}"] = f'=IF(OR($A{r}="",J{r}=""),"",IFERROR(INDEX(Suppliers!$C${S0}:$C${S1},MATCH(J{r},Suppliers!$A${S0}:$A${S1},0)),""))'
        rl[f"L{r}"] = f'=IF(OR($A{r}="",J{r}=""),"",IFERROR(INDEX(Suppliers!$F${S0}:$F${S1},MATCH(J{r},Suppliers!$A${S0}:$A${S1},0)),""))'
    style_range(rl, 6, 5 + NRL, 1, 12, fmt={4: QTY, 5: QTY, 6: QTY, 7: QTY, 8: CUR, 9: CUR})
    rl.conditional_formatting.add(f"C6:C{5 + NRL}", CellIsRule(operator="equal", formula=['"Out of stock"'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    rl.conditional_formatting.add(f"C6:C{5 + NRL}", CellIsRule(operator="equal", formula=['"Reorder"'], fill=fill("FEF3C7"), font=Font(name=F, bold=True, color="92400E")))
    rl.freeze_panes = "C6"

    # ---------- Calc ----------
    K.title(cc, "Calc", "Numbers for the dashboard charts. Calculated – nothing to type.")
    K.hdr(cc, 4, ["Month", "Label", "Sales", "Gross profit", "Units sold", "Purchases"], [11, 9, 12, 12, 9, 12])
    for i in range(12):
        r = 5 + i
        cc[f"A{r}"] = f"=DATE(YEAR({AS}),MONTH({AS})-{11 - i},1)"
        cc[f"B{r}"] = f'=TEXT(A{r},"mmm yy")'
        cc[f"C{r}"] = f'=SUMIFS({mv("J")},{mv("L")},A{r})'
        cc[f"D{r}"] = f'=C{r}-SUMIFS({mv("K")},{mv("L")},A{r})'
        cc[f"E{r}"] = f'=SUMIFS({mv("D")},{mv("L")},A{r},{mv("C")},"Sale")-SUMIFS({mv("D")},{mv("L")},A{r},{mv("C")},"Customer return")'
        cc[f"F{r}"] = f'=SUMIFS({mv("M")},{mv("L")},A{r})'
    style_range(cc, 5, 16, 1, 6, fmt={1: "mmm yyyy", 3: CUR, 4: CUR, 5: QTY, 6: CUR})
    K.hdr(cc, 19, ["Category", "Products", "Units in stock", "Stock value", "Retail value", "Sold (period)"], [16, 9, 10, 12, 12, 10])
    for i in range(15):
        r = 20 + i
        cc[f"A{r}"] = f'=IF(Settings!$H${5 + i}="","",Settings!$H${5 + i})'
        cc[f"B{r}"] = f'=IF(A{r}="",0,COUNTIF({pr("C")},A{r}))'
        for c, src in zip("CDEF", "NPQR"):
            cc[f"{c}{r}"] = f'=IF(A{r}="",0,SUMIFS({pr(src)},{pr("C")},A{r}))'
    style_range(cc, 20, 34, 1, 6, fmt={3: QTY, 4: CUR, 5: CUR, 6: QTY})

    # ---------- Dashboard ----------
    K.title(dash, "Inventory Dashboard", "Updates by itself from the Products and Stock Movements tabs.")
    dash["A3"] = f'=IF(Settings!B7="","",Settings!B7&"  ·  ")&"As of "&TEXT({AS},"dd mmm yyyy")&"  ·  amounts in "&Settings!B4'; dash["A3"].font = font(bold=True)
    inm = f'{mv("A")},">="&{MS},{mv("A")},"<="&{AS}'
    kp = [("Products", f'=COUNTA({pr("A")})', "0"), ("Units in stock", f'=SUMIF({pr("N")},">0")', "#,##0"),
          ("Stock value (at cost)", f'=SUM({pr("P")})', "#,##0.00"),
          ("Items to reorder", f'=COUNTIF({pr("O")},"Reorder")', "0"), ("Out of stock", f'=COUNTIF({pr("O")},"Out of stock")', "0"),
          ("Retail value of stock", f'=SUM({pr("Q")})', "#,##0.00"),
          ("Sales this month", f'=SUMIFS({mv("J")},{inm})', "#,##0.00"), ("Gross profit this month", f'=SUMIFS({mv("J")},{inm})-SUMIFS({mv("K")},{inm})', "#,##0.00"),
          ("Units sold this month", f'=SUMIFS({mv("D")},{inm},{mv("C")},"Sale")-SUMIFS({mv("D")},{inm},{mv("C")},"Customer return")', "#,##0")]
    for i, (lab, f, fmt) in enumerate(kp):
        K.kpi(dash, 5 + (i // 3) * 3, 1 + (i % 3) * 2, lab, f, fmt)
        dash.cell(row=6 + (i // 3) * 3, column=1 + (i % 3) * 2).alignment = Alignment(horizontal="left")
    for c, w in zip("ABCDEFGHIJKL", [30, 12, 16, 12, 18, 12, 3, 30, 13, 9, 9, 10]): dash.column_dimensions[c].width = w
    K.section(dash, "A15", "Best sellers (sales period)")
    K.hdr(dash, 16, ["Product", "Sold", "In stock", "Days left", "Status", "Margin %"])
    for n in range(10):
        r = 17 + n
        m = f"MATCH(LARGE({pr('Z')},{n + 1}),{pr('Z')},0)"
        dash[f"A{r}"] = f'=IF(LARGE({pr("Z")},{n + 1})=0,"",INDEX({pr("B")},{m}))'
        for c, src in zip("BCDEF", "RNSOT"):
            dash[f"{c}{r}"] = f'=IF($A{r}="","",INDEX({pr(src)},{m}))'
    style_range(dash, 17, 26, 1, 6, fmt={2: QTY, 3: QTY, 4: '0;-0;""', 6: '0%;-0%;""'})
    K.section(dash, "H15", "Needs attention (reorder list)")
    K.hdr(dash, 16, ["Product", "Status", "In stock", "Order qty", "Order cost"], col0=8)
    for n in range(10):
        r = 17 + n; src_r = 6 + n
        for c, src in zip("HIJKL", "BCDGI"): dash[f"{c}{r}"] = f"='Reorder List'!{src}{src_r}"
    style_range(dash, 17, 26, 8, 12, fmt={10: QTY, 11: QTY, 12: CUR})
    for sh_rng in ("E17:E26", "I17:I26"):
        dash.conditional_formatting.add(sh_rng, CellIsRule(operator="equal", formula=['"Out of stock"'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
        dash.conditional_formatting.add(sh_rng, CellIsRule(operator="equal", formula=['"Reorder"'], fill=fill("FEF3C7"), font=Font(name=F, bold=True, color="92400E")))
    for r in range(17, 27):
        for c in "BCDEIJK": dash[f"{c}{r}"].alignment = Alignment(horizontal="center")
    c1 = BarChart(); c1.type = "col"; c1.title = "Sales and gross profit by month"; c1.height = 7.5; c1.width = 17
    c1.add_data(Reference(cc, min_col=3, max_col=4, min_row=4, max_row=16), titles_from_data=True)
    c1.set_categories(Reference(cc, min_col=2, min_row=5, max_row=16))
    for ser, col in zip(c1.series, (MAIN, ACC)): ser.graphicalProperties.solidFill = col
    c1.y_axis.majorGridlines = None; c1.y_axis.scaling.min = 0; c1.y_axis.number_format = "#,##0"
    dash.add_chart(c1, "A29")
    c2 = BarChart(); c2.type = "bar"; c2.title = "Stock value by category"; c2.height = 7.5; c2.width = 14; c2.legend = None
    c2.add_data(Reference(cc, min_col=4, min_row=19, max_row=29), titles_from_data=True)
    c2.set_categories(Reference(cc, min_col=1, min_row=20, max_row=29))
    c2.y_axis.number_format = "#,##0"; c2.series[0].graphicalProperties.solidFill = MAIN; c2.x_axis.scaling.orientation = "maxMin"; c2.y_axis.majorGridlines = None
    dash.add_chart(c2, "H29")

    # ---------- Start Here ----------
    lines = [("How to start", True),
             ("1.  Settings tab: currency, shop name and your categories.", False),
             ("2.  Suppliers tab: who you buy from, with email and lead time.", False),
             ("3.  Products tab: one row per SKU – cost, sale price, opening stock (what you have today), reorder point and reorder quantity.", False),
             ("4.  Stock Movements tab: add a row for every purchase, sale, return, adjustment or damaged item. Stock updates by itself.", False),
             ("5.  Reorder List: everything at or below its reorder point, with a suggested order quantity, cost and the supplier's email.", False),
             ("6.  Dashboard: stock value, items to reorder, out-of-stock items, sales and gross profit this month, best sellers and charts.", False),
             ("", False), ("How it works", True),
             ("In stock = opening stock + stock in − stock out. Status turns 'Reorder' at or below the reorder point and 'Out of stock' at zero.", False),
             ("Days of stock left = units in stock ÷ average daily sales over the sales period (Settings, 30 days by default).", False),
             ("Suggested order = your reorder qty, or (if empty) enough to reach twice the reorder point.", False),
             ("", False), ("Colour key", True),
             ("Yellow cells are inputs: type here.  Grey cells calculate themselves: please don't overwrite them.", False),
             ("", False), ("Good to know", True),
             ("•  Room for 300 products and 3,000 stock movements. Each SKU must be unique (duplicates turn red).", False),
             ("•  Selling on Etsy, Shopify or at markets? Add one 'Sale' row per order (or one per day per product) – the order number goes in 'Order / reference'.", False),
             ("•  Google Sheets: open drive.google.com, upload this file, then choose Open with > Google Sheets.", False)]
    if sample:
        lines += [("", False), ("This copy contains sample data", True),
                  ("A made-up candle & home fragrance shop with 24 products and 3 months of stock movements, shown as of 30 Oct 2026. Use the blank copy in your download for your own shop.", False)]
    K.start(st, "Inventory & Stock Tracker", "Know what you have, what's selling and what to reorder – before you run out. Works in Excel and Google Sheets.", lines)
    finish(wb, sample, path)


META = {
    "name": "Inventory & Stock Tracker with Reorder Alerts – Excel & Google Sheets",
    "file_base": "Inventory-Stock-Tracker",
    "price": 19,
    "color": "#1E3A8A",
    "title": "Inventory & Stock Tracker",
    "tagline": "Stock levels · reorder alerts · sales & profit — Excel & Google Sheets",
    "intro": "Never run out of your best sellers again. Log purchases and sales and the tracker keeps your stock up to date, flags what to reorder, tells you how many days of stock you have left and shows your sales and profit.",
    "features": [
        "Live stock levels for up to 300 products, with status: OK, Reorder or Out of stock",
        "Automatic Reorder List with suggested quantities, order cost, supplier email and lead time",
        "Stock movements log: purchases, sales, customer returns, returns to supplier, adjustments and damaged items",
        "Days of stock left, units sold, margin % and stock value at cost and retail for every SKU",
        "Dashboard: stock value, items to reorder, sales and gross profit this month, best sellers and charts",
        "Suppliers tab with contacts and lead times; categories and locations",
        "Made for Etsy sellers, online shops, market stalls, makers and small retailers",
    ],
    "tags": ["inventory tracker", "inventory spreadsheet", "stock tracker", "small business", "google sheets", "excel template",
             "reorder alert", "inventory management", "etsy seller", "stock control", "product tracker", "shop inventory", "reseller"],
}
