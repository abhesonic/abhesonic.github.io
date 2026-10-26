"""Small Business Bookkeeping & Invoice Tracker (sample + blank)."""
import random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter as L

F="Arial"; INK="1F2A37"; ACC="0E7490"; SOFT="ECFEFF"; INPUT="FFF8DB"; GREY="F3F4F6"; LINE="D7DEE3"
thin=Side(style="thin",color=LINE); B=Border(left=thin,right=thin,top=thin,bottom=thin)
MONEY='#,##0.00;(#,##0.00);"-"'; DATE='yyyy-mm-dd'; IR=1000; ER=1500; VR=300
def font(**k): return Font(name=F, **{"color":INK, **k})
def fill(c): return PatternFill("solid", fgColor=c)
def hdr(ws,row,cols,widths=None,c1=1):
    for i,t in enumerate(cols):
        c=ws.cell(row=row,column=c1+i,value=t); c.font=Font(name=F,bold=True,color="FFFFFF"); c.fill=fill(ACC)
        c.alignment=Alignment(vertical="center",wrap_text=True); c.border=B
    ws.row_dimensions[row].height=30
    if widths:
        for i,w in enumerate(widths): ws.column_dimensions[L(c1+i)].width=w
def title(ws,t,s,col="A"):
    ws[f"{col}1"]=t; ws[f"{col}1"].font=Font(name=F,size=18,bold=True,color=ACC)
    ws[f"{col}2"]=s; ws[f"{col}2"].font=Font(name=F,size=10,italic=True,color="5B6770"); ws.sheet_view.showGridLines=False
def box(ws,r1,r2,c1,c2,inputs=(),fmt=None):
    for r in range(r1,r2+1):
        for c in range(c1,c2+1):
            x=ws.cell(row=r,column=c); x.border=B; x.font=font(); x.fill=fill(INPUT if c in inputs else GREY)
            if fmt and c in fmt: x.number_format=fmt[c]
def dv(ws,f,rng,**k):
    v=DataValidation(type="list",formula1=f,allow_blank=True,**k); ws.add_data_validation(v); v.add(rng)

INC=["Product sales","Services","Consulting","Subscriptions","Shipping charged","Other income"]
EXP=["Cost of goods","Advertising & marketing","Software & subscriptions","Office supplies","Rent","Utilities & phone","Travel","Meals",
     "Contractors","Payroll","Bank & payment fees","Insurance","Professional fees","Shipping & postage","Equipment","Training","Taxes & licenses","Other"]

def build(sample,path):
    wb=Workbook(); st=wb.active; st.title="Start Here"
    dash=wb.create_sheet("Dashboard"); inc=wb.create_sheet("Income"); exp=wb.create_sheet("Expenses"); inv=wb.create_sheet("Invoices")
    pl=wb.create_sheet("Profit & Loss"); tax=wb.create_sheet("Sales Tax"); s=wb.create_sheet("Settings")

    title(s,"Settings","Set these once. Yellow cells are yours to edit.")
    for i,(k,v) in enumerate([("Business name","Bright Studio LLC" if sample else ""),("Currency label","USD"),("Year",2026),("Default payment terms (days)",30)],4):
        s.cell(row=i,column=1,value=k).font=font(bold=True)
        c=s.cell(row=i,column=2,value=v or None); c.font=Font(name=F,color="0000FF"); c.fill=fill(INPUT); c.border=B
    s["A9"]="Income categories"; s["A9"].font=Font(name=F,bold=True,color=ACC,size=12); hdr(s,10,["Income category"])
    for i in range(10):
        c=s.cell(row=11+i,column=1,value=INC[i] if i<len(INC) else None); c.fill=fill(INPUT); c.border=B; c.font=font()
    s["C9"]="Expense categories"; s["C9"].font=Font(name=F,bold=True,color=ACC,size=12); hdr(s,10,["Expense category"],c1=3)
    for i in range(25):
        c=s.cell(row=11+i,column=3,value=EXP[i] if i<len(EXP) else None); c.fill=fill(INPUT); c.border=B; c.font=font()
    s["E9"]="Payment methods"; s["E9"].font=Font(name=F,bold=True,color=ACC,size=12); hdr(s,10,["Method"],c1=5)
    for i,m in enumerate(["Bank transfer","Card","Cash","PayPal","Stripe","Check","",""]):
        c=s.cell(row=11+i,column=5,value=m or None); c.fill=fill(INPUT); c.border=B; c.font=font()
    for c,w in zip("ABCDE",[28,22,26,3,16]): s.column_dimensions[c].width=w
    YR="Settings!$B$6"

    title(inc,"Income","Log every payment you receive. Add the invoice number when the payment is for an invoice.")
    hdr(inc,4,["Date received","Customer","Category","Description","Invoice #","Amount (before tax)","Sales tax","Total received","Method","Year","Month","Quarter"],
        [12,22,18,26,11,14,11,13,14,7,7,8])
    box(inc,5,4+IR,1,12,inputs=(1,2,3,4,5,6,7,9),fmt={1:DATE,6:MONEY,7:MONEY,8:MONEY})
    for r in range(5,5+IR):
        inc[f"H{r}"]=f'=IF(A{r}="","",N(F{r})+N(G{r}))'; inc[f"J{r}"]=f'=IF(A{r}="","",YEAR(A{r}))'
        inc[f"K{r}"]=f'=IF(A{r}="","",MONTH(A{r}))'; inc[f"L{r}"]=f'=IF(A{r}="","",ROUNDUP(MONTH(A{r})/3,0))'
    dv(inc,"=Settings!$A$11:$A$20",f"C5:C{4+IR}"); dv(inc,"=Settings!$E$11:$E$18",f"I5:I{4+IR}"); inc.freeze_panes="A5"

    title(exp,"Expenses","Log every business cost. Mark 'Tax paid' if the bill included sales tax/VAT you can claim back.")
    hdr(exp,4,["Date","Vendor","Category","Description","Amount (before tax)","Tax paid","Total paid","Method","Receipt kept?","Year","Month","Quarter"],
        [12,22,24,26,14,11,13,14,10,7,7,8])
    box(exp,5,4+ER,1,12,inputs=(1,2,3,4,5,6,8,9),fmt={1:DATE,5:MONEY,6:MONEY,7:MONEY})
    for r in range(5,5+ER):
        exp[f"G{r}"]=f'=IF(A{r}="","",N(E{r})+N(F{r}))'; exp[f"J{r}"]=f'=IF(A{r}="","",YEAR(A{r}))'
        exp[f"K{r}"]=f'=IF(A{r}="","",MONTH(A{r}))'; exp[f"L{r}"]=f'=IF(A{r}="","",ROUNDUP(MONTH(A{r})/3,0))'
    dv(exp,"=Settings!$C$11:$C$35",f"C5:C{4+ER}"); dv(exp,"=Settings!$E$11:$E$18",f"H5:H{4+ER}")
    v=DataValidation(type="list",formula1='"Yes,No"',allow_blank=True); exp.add_data_validation(v); v.add(f"I5:I{4+ER}")
    exp.conditional_formatting.add(f"I5:I{4+ER}",CellIsRule(operator="equal",formula=['"No"'],fill=fill("FEF3C7"))); exp.freeze_panes="A5"

    title(inv,"Invoices","List every invoice you send. Payments logged on the Income tab with the same invoice # mark it paid.")
    hdr(inv,4,["Invoice #","Customer","Issue date","Due date","Invoice total","Paid","Balance","Status","Days overdue"],[11,22,12,12,13,13,13,12,10])
    box(inv,5,4+VR,1,9,inputs=(1,2,3,4,5),fmt={3:DATE,4:DATE,5:MONEY,6:MONEY,7:MONEY})
    for r in range(5,5+VR):
        inv[f"F{r}"]=f'=IF(A{r}="","",SUMIFS(Income!$H$5:$H${4+IR},Income!$E$5:$E${4+IR},A{r}))'
        inv[f"G{r}"]=f'=IF(A{r}="","",MAX(0,N(E{r})-F{r}))'
        inv[f"H{r}"]=f'=IF(A{r}="","",IF(G{r}<=0,"Paid",IF(TODAY()>IF(D{r}="",C{r}+Settings!$B$7,D{r}),"Overdue",IF(F{r}>0,"Part paid","Open"))))'
        inv[f"I{r}"]=f'=IF(H{r}="Overdue",TODAY()-IF(D{r}="",C{r}+Settings!$B$7,D{r}),"")'
    inv.conditional_formatting.add(f"H5:H{4+VR}",CellIsRule(operator="equal",formula=['"Overdue"'],fill=fill("FEE2E2"),font=Font(name=F,bold=True,color="B91C1C")))
    inv.conditional_formatting.add(f"H5:H{4+VR}",CellIsRule(operator="equal",formula=['"Paid"'],fill=fill("D1FAE5")))
    inv["K4"]="Tip: leave Due date empty to use the default payment terms from Settings."; inv["K4"].font=font(italic=True,size=9); inv.freeze_panes="A5"

    if sample:
        random.seed(11); custs=["Northwind Co.","Acme Retail","Blue Fern Cafe","Harbor Dental","Kite & Co.","Summit Gym"]
        r=5; ir=5; n=1001
        for m in range(1,11):
            for _ in range(3):
                c=random.choice(custs); amt=random.choice([850,1200,1500,2400,3200]); d=dt.date(2026,m,random.randint(1,20))
                inv[f"A{ir}"]=f"INV-{n}"; inv[f"B{ir}"]=c; inv[f"C{ir}"]=d; inv[f"D{ir}"]=d+dt.timedelta(days=30); inv[f"E{ir}"]=round(amt*1.08,2)
                if not (m>=9 and random.random()<0.6):
                    pd_=d+dt.timedelta(days=random.randint(5,40))
                    if pd_<=dt.date(2026,10,26):
                        for j,vv in enumerate([pd_,c,"Services","Project work",f"INV-{n}",amt,round(amt*0.08,2),None,random.choice(["Bank transfer","Stripe"])],1):
                            if vv is not None: inc.cell(row=r,column=j,value=vv)
                        r+=1
                ir+=1; n+=1
            for _ in range(4):
                amt=random.choice([29,49,79,120]); inc.cell(row=r,column=1,value=dt.date(2026,m,random.randint(1,28))); inc.cell(row=r,column=2,value="Online store")
                inc.cell(row=r,column=3,value="Product sales"); inc.cell(row=r,column=4,value="Template sales"); inc.cell(row=r,column=6,value=amt*4)
                inc.cell(row=r,column=7,value=0); inc.cell(row=r,column=9,value="Stripe"); r+=1
        ex=[("Adobe","Software & subscriptions","Creative Cloud",55,4.4),("Google","Advertising & marketing","Search ads",300,0),("WeWork","Rent","Desk membership",450,36),
            ("Comcast","Utilities & phone","Internet",80,6.4),("Stripe","Bank & payment fees","Processing fees",62,0),("Upwork","Contractors","Freelance developer",900,0),
            ("Staples","Office supplies","Paper & ink",45,3.6),("Delta","Travel","Client visit",380,0),("Hiscox","Insurance","Liability insurance",95,0)]
        er=5
        for m in range(1,11):
            for e in ex[:5]+random.sample(ex[5:],2):
                for j,vv in enumerate([dt.date(2026,m,random.randint(1,28)),e[0],e[1],e[2],e[3],e[4],None,"Card",random.choice(["Yes","Yes","Yes","No"])],1):
                    if vv is not None: exp.cell(row=er,column=j,value=vv)
                er+=1

    I=lambda c: f"Income!${c}$5:${c}${4+IR}"; E=lambda c: f"Expenses!${c}$5:${c}${4+ER}"
    months=["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
    title(pl,"Profit & Loss","Monthly profit and loss for the year on the Settings tab. Amounts are before sales tax.")
    hdr(pl,4,["Line"]+months+["Total"],[26]+[10]*12+[12])
    pl["A5"]="INCOME"; pl["A5"].font=Font(name=F,bold=True,color=ACC)
    for i in range(10):
        r=6+i; pl[f"A{r}"]=f'=IF(Settings!A{11+i}="","",Settings!A{11+i})'
        for m in range(1,13): pl.cell(row=r,column=1+m,value=f'=IF($A{r}="","",SUMIFS({I("F")},{I("C")},$A{r},{I("J")},{YR},{I("K")},{m}))')
    pl["A16"]="Total income"; pl["A17"]="EXPENSES"; pl["A17"].font=Font(name=F,bold=True,color=ACC)
    for i in range(25):
        r=18+i; pl[f"A{r}"]=f'=IF(Settings!C{11+i}="","",Settings!C{11+i})'
        for m in range(1,13): pl.cell(row=r,column=1+m,value=f'=IF($A{r}="","",SUMIFS({E("E")},{E("C")},$A{r},{E("J")},{YR},{E("K")},{m}))')
    pl["A43"]="Total expenses"; pl["A44"]="Net profit"; pl["A45"]="Profit margin"
    for m in range(1,13):
        c=L(1+m); pl[f"{c}16"]=f"=SUM({c}6:{c}15)"; pl[f"{c}43"]=f"=SUM({c}18:{c}42)"; pl[f"{c}44"]=f"={c}16-{c}43"; pl[f"{c}45"]=f'=IF({c}16=0,"",{c}44/{c}16)'
    for r in list(range(6,17))+list(range(18,45)): pl[f"N{r}"]=f'=IF(A{r}="","",SUM(B{r}:M{r}))'
    pl["N45"]='=IF(N16=0,"",N44/N16)'
    for r in list(range(6,17))+list(range(18,46)):
        for c in range(1,15):
            x=pl.cell(row=r,column=c); x.border=B; x.font=font(bold=r in (16,43,44,45)); x.number_format='0%' if r==45 else MONEY
            if r in (16,43,44,45): x.fill=fill(SOFT)
    pl.conditional_formatting.add("B44:N44",CellIsRule(operator="lessThan",formula=["0"],font=Font(name=F,bold=True,color="B91C1C")))
    pl.freeze_panes="B5"

    title(tax,"Sales Tax / VAT by Quarter","Tax you collected from customers minus tax you paid on business costs.")
    hdr(tax,4,["Quarter","Tax collected","Tax paid on costs","Net tax (to pay)"],[12,16,18,16])
    for q in range(1,5):
        r=4+q; tax[f"A{r}"]=f"Q{q}"
        tax[f"B{r}"]=f"=SUMIFS({I('G')},{I('J')},{YR},{I('L')},{q})"; tax[f"C{r}"]=f"=SUMIFS({E('F')},{E('J')},{YR},{E('L')},{q})"; tax[f"D{r}"]=f"=B{r}-C{r}"
    tax["A9"]="Year"; 
    for c in "BCD": tax[f"{c}9"]=f"=SUM({c}5:{c}8)"
    for r in range(5,10):
        for c in "ABCD": x=tax[f"{c}{r}"]; x.border=B; x.font=font(bold=r==9); x.number_format=MONEY if c!="A" else "General"
    tax["A11"]="This is a summary to help you prepare, not tax advice. Check your local rules or ask an accountant."; tax["A11"].font=font(italic=True,size=9)

    title(dash,"Business Dashboard","Updates by itself from your Income, Expenses and Invoices. Change the year on the Settings tab.")
    dash["A3"]='=IF(Settings!B4="","",Settings!B4&"  ·  ")&"Year "&Settings!B6&"  ·  Amounts in "&Settings!B5'; dash["A3"].font=font(bold=True)
    kp=[("Income (year)","='Profit & Loss'!N16",MONEY),("Expenses (year)","='Profit & Loss'!N43",MONEY),("Net profit (year)","='Profit & Loss'!N44",MONEY),
        ("Profit margin","='Profit & Loss'!N45",'0%'),("Unpaid invoices",f"=SUM(Invoices!$G$5:$G${4+VR})",MONEY),("Overdue invoices",f'=COUNTIF(Invoices!$H$5:$H${4+VR},"Overdue")','0')]
    for i,(lab,f,fmt) in enumerate(kp):
        col=1+(i%3)*2; row=5+(i//3)*3
        a=dash.cell(row=row,column=col,value=lab); a.font=Font(name=F,size=9,bold=True,color="5B6770")
        v=dash.cell(row=row+1,column=col,value=f); v.font=Font(name=F,size=18,bold=True,color=ACC); v.number_format=fmt
        for rr in (row,row+1):
            for cc in (col,col+1): dash.cell(row=rr,column=cc).fill=fill(SOFT)
    for c,w in zip("ABCDEFGHI",[17,14,16,13,18,13,4,26,14]): dash.column_dimensions[c].width=w
    dash["A13"]="Month by month"; dash["A13"].font=Font(name=F,bold=True,size=12,color=ACC)
    hdr(dash,14,["Month","Income","Expenses","Profit"])
    for m in range(1,13):
        r=14+m; c=L(1+m); dash[f"A{r}"]=months[m-1]
        dash[f"B{r}"]=f"='Profit & Loss'!{c}16"; dash[f"C{r}"]=f"='Profit & Loss'!{c}43"; dash[f"D{r}"]=f"=B{r}-C{r}"
        for cc in "ABCD":
            dash[f"{cc}{r}"].border=B; dash[f"{cc}{r}"].font=font()
            if cc!="A": dash[f"{cc}{r}"].number_format=MONEY
    dash.conditional_formatting.add("D15:D26",CellIsRule(operator="lessThan",formula=["0"],font=Font(name=F,bold=True,color="B91C1C")))
    dash["H13"]="Expenses by category (year)"; dash["H13"].font=Font(name=F,bold=True,size=12,color=ACC)
    hdr(dash,14,["Category","Spent"],c1=8)
    for i in range(25):
        r=15+i; dash[f"H{r}"]=f"='Profit & Loss'!A{18+i}"; dash[f"I{r}"]=f"=IF(H{r}=\"\",\"\",'Profit & Loss'!N{18+i})"
        for cc in "HI": dash[f"{cc}{r}"].border=B; dash[f"{cc}{r}"].font=font()
        dash[f"I{r}"].number_format=MONEY
    ch=BarChart(); ch.type="col"; ch.title="Income vs expenses"; ch.height=7.5; ch.width=15
    ch.add_data(Reference(dash,min_col=2,min_row=14,max_col=3,max_row=26),titles_from_data=True); ch.set_categories(Reference(dash,min_col=1,min_row=15,max_row=26))
    ch.y_axis.majorGridlines=None
    for ser,col in zip(ch.series,(ACC,"F59E0B")): ser.graphicalProperties=GraphicalProperties(solidFill=col); ser.graphicalProperties.line.solidFill=col
    dash.add_chart(ch,"A29")

    title(st,"Small Business Bookkeeping","Income, expenses, invoices, profit & loss and sales tax in one file. Works in Excel and Google Sheets.",col="B")
    st.column_dimensions["A"].width=3; st.column_dimensions["B"].width=100
    lines=[("How to start",True),
      ("1.  Settings: type your business name, currency label and year. Edit the income and expense categories and payment methods to fit your business.",False),
      ("2.  Invoices: list each invoice you send (number, customer, dates, total). Status shows Open, Part paid, Paid or Overdue.",False),
      ("3.  Income: log every payment you receive. Type the invoice number when it pays an invoice, and the invoice is marked paid automatically.",False),
      ("4.  Expenses: log every business cost with its category. Mark whether you kept the receipt.",False),
      ("5.  Profit & Loss, Sales Tax and the Dashboard update by themselves.",False),
      ("",False),("Colour key",True),("Yellow cells are inputs: type here.  Grey cells calculate themselves: please don't overwrite them.",False),
      ("",False),("Good to know",True),
      ("•  Room for 1,000 income entries, 1,500 expenses and 300 invoices. New year? Change the year on Settings and keep logging in the same lists.",False),
      ("•  Profit & Loss uses amounts before sales tax; tax is tracked separately on the Sales Tax tab.",False),
      ("•  To use Google Sheets: open drive.google.com, upload this file, then choose Open with > Google Sheets.",False),
      ("•  This template helps you keep records. It is not accounting or tax advice.",False)]
    if sample: lines+=[("",False),("This copy contains sample data",True),("All names and amounts are made-up examples. Use the blank copy in your download for your own business.",False)]
    for i,(t,b) in enumerate(lines,4):
        c=st.cell(row=i,column=2,value=t); c.alignment=Alignment(wrap_text=True,vertical="top"); c.font=Font(name=F,bold=b,size=12 if b else 10,color=ACC if b else INK)
    for ws in wb.worksheets:
        ws.page_setup.orientation="landscape"; ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True); ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    wb.active=1 if sample else 0
    wb.save(path)

META = {
    "name": "Small Business Bookkeeping – Excel & Google Sheets Template",
    "file_base": "Small-Business-Bookkeeping",
    "price": 29,
    "color": "#0E7490",
    "title": "Small Business Bookkeeping",
    "tagline": "Income · expenses · invoices · P&L · sales tax — Excel & Google Sheets",
    "intro": "Simple bookkeeping for freelancers and small businesses. Log income and expenses, track who still owes you, and see your profit every month — no accounting software needed.",
    "features": [
        "Dashboard: income, expenses, net profit, profit margin, unpaid and overdue invoices, with a monthly chart",
        "Invoice tracker: status updates by itself (Open, Part paid, Paid, Overdue) with days overdue",
        "Income log: payments linked to invoices by invoice number",
        "Expense log with 18 ready business categories and a receipt-kept check",
        "Profit & loss by month and category, with profit margin",
        "Sales tax / VAT summary by quarter: collected, paid and net",
        "Any currency, room for 1,000 income entries, 1,500 expenses and 300 invoices",
    ],
    "tags": ["bookkeeping template", "small business", "income expense tracker", "invoice tracker", "profit and loss", "google sheets", "excel template",
             "freelancer finance", "small business budget", "business spreadsheet", "accounting template", "sales tax tracker", "self employed"],
}
