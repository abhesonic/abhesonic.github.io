"""Builds the Rental Property Manager spreadsheet (sample + blank versions)."""
import random, datetime as dt, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as L

F = "Arial"
INK = "1F2A37"; TEAL = "0F766E"; SOFT = "E6F2EF"; INPUT = "FFF8DB"; GREY = "F3F4F6"; LINE = "D7DEE3"
thin = Side(style="thin", color=LINE)
B = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '#,##0.00;(#,##0.00);"-"'
DATE = 'yyyy-mm-dd'
NR = 100  # tenants rows
PR = 600  # payment rows
ER = 500  # expense rows

def font(**k): return Font(name=F, **{"color": INK, **k})
def hdr(ws, row, cols, widths=None):
    for i, t in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=TEAL)
        c.alignment = Alignment(vertical="center", wrap_text=True); c.border = B
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = Font(name=F, size=18, bold=True, color=TEAL)
    ws["A2"] = sub; ws["A2"].font = Font(name=F, size=10, italic=True, color="5B6770")
    ws.sheet_view.showGridLines = False
def style_range(ws, r1, r2, c1, c2, inputs=(), fmt=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c); cell.border = B; cell.font = font()
            if c in inputs: cell.fill = PatternFill("solid", fgColor=INPUT)
            else: cell.fill = PatternFill("solid", fgColor=GREY)
            if fmt and c in fmt: cell.number_format = fmt[c]

def build(sample, path):
    wb = Workbook()
    st = wb.active; st.title = "Start Here"
    dash = wb.create_sheet("Dashboard"); ten = wb.create_sheet("Tenants")
    pay = wb.create_sheet("Rent Payments"); exp = wb.create_sheet("Expenses")
    pl = wb.create_sheet("Property P&L"); s = wb.create_sheet("Settings")

    # ---------- Settings ----------
    title(s, "Settings", "Set these once. Yellow cells are yours to edit.")
    rows = [("Currency label", "USD"), ("Report year", 2026), ("Report month (1-12)", 10)]
    for i, (k, v) in enumerate(rows, 4):
        s.cell(row=i, column=1, value=k).font = font(bold=True)
        c = s.cell(row=i, column=2, value=v); c.font = Font(name=F, color="0000FF"); c.fill = PatternFill("solid", fgColor=INPUT); c.border = B
    s["A7"] = "Month start"; s["B7"] = "=DATE(B5,B6,1)"; s["A8"] = "Month end"; s["B8"] = "=EOMONTH(B7,0)"
    for a in ("A7", "A8"): s[a].font = font(italic=True)
    for a in ("B7", "B8"): s[a].number_format = DATE; s[a].font = font(); s[a].border = B
    s["C6"] = "Dashboard, tenant balances and the P&L follow this year and month."; s["C6"].font = font(italic=True, size=9)
    s["A10"] = "Properties (up to 10)"; s["A10"].font = Font(name=F, bold=True, color=TEAL, size=12)
    hdr(s, 11, ["Property name", "Address", "Units"])
    props = [("Maple Court", "12 Maple Ave", 4), ("Riverside Duplex", "88 River Rd", 2), ("Oak Street House", "5 Oak St", 1)] if sample else []
    for i in range(10):
        r = 12 + i
        if i < len(props):
            for j, v in enumerate(props[i], 1): s.cell(row=r, column=j, value=v)
    style_range(s, 12, 21, 1, 3, inputs=(1, 2, 3))
    s["A22"] = "Total units"; s["A22"].font = font(bold=True); s["C22"] = "=SUM(C12:C21)"; s["C22"].font = font(bold=True)
    s["E10"] = "Expense categories"; s["E10"].font = Font(name=F, bold=True, color=TEAL, size=12)
    hdr_cell = s.cell(row=11, column=5, value="Category"); hdr_cell.font = Font(name=F, bold=True, color="FFFFFF"); hdr_cell.fill = PatternFill("solid", fgColor=TEAL)
    cats = ["Repairs & maintenance", "Property tax", "Insurance", "Mortgage interest", "Utilities", "Management fees", "Cleaning", "Advertising", "Legal & professional", "Supplies", "HOA fees", "Other"]
    for i in range(15):
        c = s.cell(row=12 + i, column=5, value=cats[i] if i < len(cats) else None); c.fill = PatternFill("solid", fgColor=INPUT); c.border = B; c.font = font()
    s["G10"] = "Payment methods"; s["G10"].font = Font(name=F, bold=True, color=TEAL, size=12)
    h = s.cell(row=11, column=7, value="Method"); h.font = Font(name=F, bold=True, color="FFFFFF"); h.fill = PatternFill("solid", fgColor=TEAL)
    for i, m in enumerate(["Bank transfer", "Cash", "Check", "Card", "Online app", "", "", ""]):
        c = s.cell(row=12 + i, column=7, value=m or None); c.fill = PatternFill("solid", fgColor=INPUT); c.border = B; c.font = font()
    s["I10"] = "Lease status"; s["I10"].font = Font(name=F, bold=True, color=TEAL, size=12)
    s["I11"] = "Active"; s["I12"] = "Ended"
    for col, w in zip("ABCDEFGHI", [24, 26, 10, 3, 24, 3, 18, 3, 14]): s.column_dimensions[col].width = w

    PROPS = "Settings!$A$12:$A$21"; CATS = "Settings!$E$12:$E$26"; METH = "Settings!$G$12:$G$19"
    YR = "Settings!$B$5"; MO = "Settings!$B$6"; MS = "Settings!$B$7"; ME = "Settings!$B$8"; CUR = "Settings!$B$4"

    # ---------- Tenants ----------
    title(ten, "Tenants & Leases", "One row per unit or tenant. Yellow columns are yours; grey columns calculate themselves.")
    cols = ["Property", "Unit", "Tenant name", "Phone", "Email", "Lease start", "Lease end", "Monthly rent", "Deposit", "Rent due day", "Status",
            "Tenant ID", "Active this month", "Paid this month", "Balance due", "Rent status", "Days to lease end"]
    hdr(ten, 4, cols, [20, 9, 20, 15, 24, 12, 12, 13, 12, 10, 10, 26, 10, 13, 13, 12, 12])
    style_range(ten, 5, 4 + NR, 1, 17, inputs=range(1, 12), fmt={6: DATE, 7: DATE, 8: MONEY, 9: MONEY, 14: MONEY, 15: MONEY})
    tenants = []
    if sample:
        tenants = [("Maple Court", "1A", "Jordan Lee", "555-0141", "jordan@example.com", dt.date(2025, 3, 1), dt.date(2026, 12, 31), 1250, 1250, 1, "Active"),
                   ("Maple Court", "1B", "Priya Shah", "555-0172", "priya@example.com", dt.date(2025, 9, 1), dt.date(2026, 8, 31), 1180, 1180, 1, "Ended"),
                   ("Maple Court", "1B", "Marco Rossi", "555-0190", "marco@example.com", dt.date(2026, 9, 15), dt.date(2027, 9, 14), 1220, 1220, 15, "Active"),
                   ("Maple Court", "2A", "Emma Brown", "555-0133", "emma@example.com", dt.date(2024, 6, 1), dt.date(2026, 11, 30), 1300, 1300, 1, "Active"),
                   ("Riverside Duplex", "Upper", "Sam Carter", "555-0115", "sam@example.com", dt.date(2025, 1, 1), dt.date(2026, 12, 31), 1650, 1650, 5, "Active"),
                   ("Riverside Duplex", "Lower", "Ava Nguyen", "555-0168", "ava@example.com", dt.date(2026, 2, 1), dt.date(2027, 1, 31), 1550, 1550, 5, "Active"),
                   ("Oak Street House", "House", "The Garcia family", "555-0107", "garcia@example.com", dt.date(2025, 7, 1), dt.date(2027, 6, 30), 2100, 2100, 1, "Active")]
    for i, t in enumerate(tenants):
        for j, v in enumerate(t, 1): ten.cell(row=5 + i, column=j, value=v)
    for r in range(5, 5 + NR):
        ten[f"L{r}"] = f'=IF(OR(A{r}="",C{r}=""),"",A{r}&" | "&B{r}&" | "&C{r})'
        ten[f"M{r}"] = f'=IF(L{r}="",0,IF(AND(K{r}<>"Ended",OR(F{r}="",F{r}<={ME}),OR(G{r}="",G{r}>={MS})),1,0))'
        ten[f"N{r}"] = f"=IF(L{r}=\"\",\"\",SUMIFS('Rent Payments'!$D$5:$D${4+PR},'Rent Payments'!$B$5:$B${4+PR},L{r},'Rent Payments'!$H$5:$H${4+PR},{YR},'Rent Payments'!$I$5:$I${4+PR},{MO}))"
        ten[f"O{r}"] = f'=IF(L{r}="","",MAX(0,M{r}*H{r}-N{r}))'
        ten[f"P{r}"] = (f'=IF(L{r}="","",IF(M{r}=0,"Not active",IF(O{r}<=0,"Paid",IF(TODAY()>DATE({YR},{MO},MIN(MAX(J{r},1),DAY({ME}))),'
                        f'IF(N{r}>0,"Partial - late","Overdue"),IF(N{r}>0,"Partial","Due")))))')
        ten[f"Q{r}"] = f'=IF(OR(L{r}="",G{r}="",K{r}="Ended"),"",G{r}-TODAY())'
    dv = DataValidation(type="list", formula1=f"={PROPS}", allow_blank=True); ten.add_data_validation(dv); dv.add(f"A5:A{4+NR}")
    dv = DataValidation(type="list", formula1="=Settings!$I$11:$I$12", allow_blank=True); ten.add_data_validation(dv); dv.add(f"K5:K{4+NR}")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="31", allow_blank=True); ten.add_data_validation(dv); dv.add(f"J5:J{4+NR}")
    rng = f"P5:P{4+NR}"
    ten.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Paid"'], fill=PatternFill("solid", fgColor="D1FAE5")))
    ten.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(P5,7)="Overdue"'], fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    ten.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("Partial",P5))'], fill=PatternFill("solid", fgColor="FEF3C7")))
    ten.conditional_formatting.add(f"Q5:Q{4+NR}", FormulaRule(formula=['AND(ISNUMBER(Q5),Q5<=60)'], fill=PatternFill("solid", fgColor="FEF3C7"), font=Font(name=F, bold=True, color="92400E")))
    ten.freeze_panes = "D5"

    # ---------- Rent Payments ----------
    title(pay, "Rent Payments", "Log every payment. Pick the tenant from the list; 'Rent for month' is the month the payment covers.")
    hdr(pay, 4, ["Date paid", "Tenant (pick from list)", "Rent for month", "Amount", "Method", "Note", "Property", "Rent yr", "Rent mo", "Paid yr", "Paid mo"],
        [12, 40, 13, 12, 15, 24, 20, 8, 8, 8, 8])
    style_range(pay, 5, 4 + PR, 1, 11, inputs=range(1, 7), fmt={1: DATE, 3: 'mmm yyyy', 4: MONEY})
    for r in range(5, 5 + PR):
        pay[f"G{r}"] = f'=IF(B{r}="","",IFERROR(INDEX(Tenants!$A$5:$A${4+NR},MATCH(B{r},Tenants!$L$5:$L${4+NR},0)),"Unknown tenant"))'
        pay[f"H{r}"] = f'=IF(C{r}="","",YEAR(C{r}))'; pay[f"I{r}"] = f'=IF(C{r}="","",MONTH(C{r}))'
        pay[f"J{r}"] = f'=IF(A{r}="","",YEAR(A{r}))'; pay[f"K{r}"] = f'=IF(A{r}="","",MONTH(A{r}))'
    dv = DataValidation(type="list", formula1=f"=Tenants!$L$5:$L${4+NR}", allow_blank=True); pay.add_data_validation(dv); dv.add(f"B5:B{4+PR}")
    dv = DataValidation(type="list", formula1=f"={METH}", allow_blank=True); pay.add_data_validation(dv); dv.add(f"E5:E{4+PR}")
    pay.freeze_panes = "A5"

    # ---------- Expenses ----------
    title(exp, "Expenses", "Log every cost against a property and a category.")
    hdr(exp, 4, ["Date", "Property", "Category", "Description", "Amount", "Vendor / payee", "Year", "Month"], [12, 20, 22, 30, 12, 20, 8, 8])
    style_range(exp, 5, 4 + ER, 1, 8, inputs=range(1, 7), fmt={1: DATE, 5: MONEY})
    for r in range(5, 5 + ER):
        exp[f"G{r}"] = f'=IF(A{r}="","",YEAR(A{r}))'; exp[f"H{r}"] = f'=IF(A{r}="","",MONTH(A{r}))'
    dv = DataValidation(type="list", formula1=f"={PROPS}", allow_blank=True); exp.add_data_validation(dv); dv.add(f"B5:B{4+ER}")
    dv = DataValidation(type="list", formula1=f"={CATS}", allow_blank=True); exp.add_data_validation(dv); dv.add(f"C5:C{4+ER}")
    exp.freeze_panes = "A5"

    if sample:
        random.seed(7); prow = 5
        key = lambda t: f"{t[0]} | {t[1]} | {t[2]}"
        for t in tenants:
            for m in range(1, 11):
                ms = dt.date(2026, m, 1)
                start = t[5]; end = t[6]
                if ms < start.replace(day=1) or ms > end: continue
                if (t[2] == "Emma Brown" and m == 10) or (t[2] == "Sam Carter" and m == 10): continue
                amt = t[7] if not (t[2] == "Ava Nguyen" and m == 10) else 800
                paid = dt.date(2026, m, min(t[9] + random.randint(-3, 3), 28) if t[9] > 3 else random.randint(1, 4))
                pay[f"A{prow}"] = paid; pay[f"B{prow}"] = key(t); pay[f"C{prow}"] = ms; pay[f"D{prow}"] = amt
                pay[f"E{prow}"] = random.choice(["Bank transfer", "Bank transfer", "Online app", "Check"]); prow += 1
        exps = [("Maple Court", "Repairs & maintenance", "Boiler service", 420, "HeatPro"), ("Maple Court", "Property tax", "Q1 property tax", 1850, "County"),
                ("Maple Court", "Insurance", "Annual building policy", 2400, "SafeHome Ins."), ("Maple Court", "Cleaning", "Hallway cleaning", 160, "BrightClean"),
                ("Riverside Duplex", "Repairs & maintenance", "Leaking faucet", 140, "Joe's Plumbing"), ("Riverside Duplex", "Property tax", "Q1 property tax", 1320, "County"),
                ("Riverside Duplex", "Utilities", "Shared water bill", 95, "City Water"), ("Oak Street House", "Mortgage interest", "Monthly interest", 610, "First Bank"),
                ("Oak Street House", "Repairs & maintenance", "Gutter cleaning", 180, "RoofCare"), ("Maple Court", "Advertising", "Listing for unit 1B", 60, "RentList"),
                ("Maple Court", "Management fees", "Leasing fee unit 1B", 610, "Agent"), ("Riverside Duplex", "Insurance", "Annual policy", 1450, "SafeHome Ins.")]
        erow = 5
        for m in range(1, 11):
            for e in random.sample(exps, 3):
                exp[f"A{erow}"] = dt.date(2026, m, random.randint(2, 27)); exp[f"B{erow}"] = e[0]; exp[f"C{erow}"] = e[1]
                exp[f"D{erow}"] = e[2]; exp[f"E{erow}"] = e[3]; exp[f"F{erow}"] = e[4]; erow += 1

    # ---------- Dashboard ----------
    title(dash, "Rental Property Dashboard", "Everything here calculates from your other tabs. Change the year and month on the Settings tab.")
    dash["A3"] = '="Year "&Settings!B5&"  ·  Month "&TEXT(Settings!B7,"mmmm")&"  ·  Amounts in "&Settings!B4'
    dash["A3"].font = font(bold=True, size=11)
    kpis = [("Rent collected (year)", f"=SUMIFS('Rent Payments'!$D$5:$D${4+PR},'Rent Payments'!$J$5:$J${4+PR},{YR})", MONEY),
            ("Expenses (year)", f"=SUMIFS(Expenses!$E$5:$E${4+ER},Expenses!$G$5:$G${4+ER},{YR})", MONEY),
            ("Net income (year)", "=A6-C6", MONEY),
            ("Occupancy (this month)", f"=IFERROR(SUM(Tenants!$M$5:$M${4+NR})/Settings!$C$22,0)", '0%'),
            ("Rent outstanding (this month)", f"=SUM(Tenants!$O$5:$O${4+NR})", MONEY),
            ("Leases ending in 60 days", f'=COUNTIFS(Tenants!$Q$5:$Q${4+NR},">=0",Tenants!$Q$5:$Q${4+NR},"<=60")', '0')]
    for i, (lab, f, fmt) in enumerate(kpis):
        col = 1 + (i % 3) * 2; row = 5 + (i // 3) * 3
        lc = dash.cell(row=row, column=col, value=lab); lc.font = Font(name=F, size=9, color="5B6770", bold=True)
        vc = dash.cell(row=row + 1, column=col, value=f); vc.font = Font(name=F, size=18, bold=True, color=TEAL); vc.number_format = fmt
        for rr in (row, row + 1):
            for cc in (col, col + 1):
                dash.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor=SOFT)
    for c, w in zip("ABCDEFGH", [17, 15, 16, 13, 16, 13, 16, 14]): dash.column_dimensions[c].width = w
    dash["A12"] = "Month by month"; dash["A12"].font = Font(name=F, bold=True, size=12, color=TEAL)
    for i, h in enumerate(["Month", "Rent collected", "Expenses", "Net"], 1):
        c = dash.cell(row=13, column=i, value=h); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=TEAL); c.border = B
    for m in range(1, 13):
        r = 13 + m
        dash[f"A{r}"] = dt.date(2026, m, 1).strftime("%b")
        dash[f"B{r}"] = f"=SUMIFS('Rent Payments'!$D$5:$D${4+PR},'Rent Payments'!$J$5:$J${4+PR},{YR},'Rent Payments'!$K$5:$K${4+PR},{m})"
        dash[f"C{r}"] = f"=SUMIFS(Expenses!$E$5:$E${4+ER},Expenses!$G$5:$G${4+ER},{YR},Expenses!$H$5:$H${4+ER},{m})"
        dash[f"D{r}"] = f"=B{r}-C{r}"
        for c in "ABCD":
            dash[f"{c}{r}"].border = B; dash[f"{c}{r}"].font = font()
            if c != "A": dash[f"{c}{r}"].number_format = MONEY
    dash["A26"] = "Total"; dash["A26"].font = font(bold=True)
    for c in "BCD":
        dash[f"{c}26"] = f"=SUM({c}14:{c}25)"; dash[f"{c}26"].font = font(bold=True); dash[f"{c}26"].number_format = MONEY; dash[f"{c}26"].border = B
    dash.row_dimensions[13].height = 28
    ch = BarChart(); ch.type = "col"; ch.title = "Rent vs expenses"; ch.height = 7.5; ch.width = 16
    ch.add_data(Reference(dash, min_col=2, min_row=13, max_col=3, max_row=25), titles_from_data=True)
    ch.set_categories(Reference(dash, min_col=1, min_row=14, max_row=25)); ch.y_axis.majorGridlines = None
    from openpyxl.chart.shapes import GraphicalProperties
    for ser, col in zip(ch.series, ("0F766E", "F59E0B")):
        ser.graphicalProperties = GraphicalProperties(solidFill=col); ser.graphicalProperties.line.solidFill = col
    dash.add_chart(ch, "F12")
    dash["A29"] = "By property (year)"; dash["A29"].font = Font(name=F, bold=True, size=12, color=TEAL)
    heads = ["Property", "Units", "Occupied now", "Occupancy", "Rent collected", "Expenses", "Net", "Outstanding now"]
    for i, h in enumerate(heads, 1):
        c = dash.cell(row=30, column=i, value=h); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=TEAL); c.border = B
        c.alignment = Alignment(wrap_text=True)
    for i in range(10):
        r = 31 + i; sr = 12 + i
        dash[f"A{r}"] = f'=IF(Settings!A{sr}="","",Settings!A{sr})'
        dash[f"B{r}"] = f'=IF(A{r}="","",Settings!C{sr})'
        dash[f"C{r}"] = f'=IF(A{r}="","",SUMIFS(Tenants!$M$5:$M${4+NR},Tenants!$A$5:$A${4+NR},A{r}))'
        dash[f"D{r}"] = f'=IF(A{r}="","",IFERROR(C{r}/B{r},0))'
        dash[f"E{r}"] = f"=IF(A{r}=\"\",\"\",SUMIFS('Rent Payments'!$D$5:$D${4+PR},'Rent Payments'!$G$5:$G${4+PR},A{r},'Rent Payments'!$J$5:$J${4+PR},{YR}))"
        dash[f"F{r}"] = f'=IF(A{r}="","",SUMIFS(Expenses!$E$5:$E${4+ER},Expenses!$B$5:$B${4+ER},A{r},Expenses!$G$5:$G${4+ER},{YR}))'
        dash[f"G{r}"] = f'=IF(A{r}="","",E{r}-F{r})'
        dash[f"H{r}"] = f'=IF(A{r}="","",SUMIFS(Tenants!$O$5:$O${4+NR},Tenants!$A$5:$A${4+NR},A{r}))'
        for j in range(1, 9):
            c = dash.cell(row=r, column=j); c.border = B; c.font = font()
            c.number_format = {4: '0%', 5: MONEY, 6: MONEY, 7: MONEY, 8: MONEY}.get(j, 'General')
    dash["A42"] = "Tip: tenants marked Overdue or Partial are highlighted on the Tenants tab, and leases ending within 60 days are flagged in yellow."
    dash["A42"].font = font(italic=True, size=9)

    # ---------- Property P&L ----------
    title(pl, "Property Profit & Loss", "Pick a property in the yellow cell to see its 12-month P&L for the report year.")
    pl["A4"] = "Property"; pl["A4"].font = font(bold=True)
    pl["B4"] = props[0][0] if props else None
    pl["B4"].fill = PatternFill("solid", fgColor=INPUT); pl["B4"].border = B; pl["B4"].font = Font(name=F, color="0000FF", bold=True)
    dv = DataValidation(type="list", formula1=f"={PROPS}", allow_blank=True); pl.add_data_validation(dv); dv.add("B4")
    pl["D4"] = '="Year "&Settings!B5'; pl["D4"].font = font(bold=True)
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    hdr(pl, 6, ["Line"] + months + ["Total"], [24] + [10] * 12 + [12])
    pl["A7"] = "Rent income"; pl["A7"].font = font(bold=True)
    for m in range(1, 13):
        pl.cell(row=7, column=1 + m, value=f"=SUMIFS('Rent Payments'!$D$5:$D${4+PR},'Rent Payments'!$G$5:$G${4+PR},$B$4,'Rent Payments'!$J$5:$J${4+PR},{YR},'Rent Payments'!$K$5:$K${4+PR},{m})")
    pl["A8"] = "Expenses"; pl["A8"].font = Font(name=F, bold=True, color=TEAL)
    for i in range(15):
        r = 9 + i
        pl[f"A{r}"] = f'=IF(Settings!E{12+i}="","",Settings!E{12+i})'
        for m in range(1, 13):
            pl.cell(row=r, column=1 + m, value=f'=IF($A{r}="","",SUMIFS(Expenses!$E$5:$E${4+ER},Expenses!$B$5:$B${4+ER},$B$4,Expenses!$C$5:$C${4+ER},$A{r},Expenses!$G$5:$G${4+ER},{YR},Expenses!$H$5:$H${4+ER},{m}))')
    pl["A24"] = "Total expenses"; pl["A25"] = "Net income"
    for m in range(1, 13):
        col = L(1 + m)
        pl[f"{col}24"] = f"=SUM({col}9:{col}23)"; pl[f"{col}25"] = f"={col}7-{col}24"
    for r in range(7, 26):
        pl[f"N{r}"] = f'=IF(A{r}="","",SUM(B{r}:M{r}))' if r != 8 else None
        for c in range(1, 15):
            cell = pl.cell(row=r, column=c); cell.border = B
            if r != 8 and c > 1: cell.number_format = MONEY
            if not cell.font.bold: cell.font = font(bold=(r in (7, 24, 25) or c == 14))
    for r in (24, 25):
        pl[f"A{r}"].font = font(bold=True)
        for c in range(1, 15): pl.cell(row=r, column=c).fill = PatternFill("solid", fgColor=SOFT)

    # ---------- Start Here ----------
    title(st, "Rental Property Manager", "Track tenants, rent, expenses and profit for up to 10 properties. Works in Excel and Google Sheets.")
    st.column_dimensions["A"].width = 4; st.column_dimensions["B"].width = 100
    lines = [("How to start", True),
             ("1.  Settings tab: type your currency label, the year and month you want to report on, and your properties with their number of units.", False),
             ("2.  Tenants tab: add one row per tenant (or per unit). Status 'Ended' keeps the history but stops rent from being expected.", False),
             ("3.  Rent Payments tab: log each payment. Pick the tenant from the drop-down list and set the month the payment covers.", False),
             ("4.  Expenses tab: log each cost with its property and category.", False),
             ("5.  Dashboard and Property P&L update on their own. Change the month on the Settings tab to see any month.", False),
             ("", False),
             ("Colour key", True),
             ("Yellow cells are inputs: type here.  Grey cells calculate themselves: please don't overwrite them.", False),
             ("Rent status on the Tenants tab: Paid (green), Partial (amber), Overdue (red). Leases ending within 60 days are flagged amber.", False),
             ("", False),
             ("Good to know", True),
             ("•  Rent collected and expenses on the Dashboard count payments by the date they were paid (cash basis).", False),
             ("•  'Paid this month' and 'Balance due' on the Tenants tab use the 'Rent for month' column, so late payments still count toward the right month.", False),
             ("•  To use Google Sheets: open drive.google.com, upload this file, then choose Open with > Google Sheets.", False),
             ("•  Room for 100 tenants, 600 payments and 500 expenses. Copy a row down if you need more.", False)]
    if sample:
        lines += [("", False), ("This copy contains sample data", True),
                  ("The tenants, payments and expenses are made-up examples so you can see how everything works. Delete the yellow cells on the Tenants, Rent Payments and Expenses tabs (and the properties on Settings) before entering your own, or use the blank copy included in your download.", False)]
    st["B1"] = st["A1"].value; st["B1"].font = Font(name=F, size=18, bold=True, color=TEAL); st["A1"] = None
    st["B2"] = st["A2"].value; st["B2"].font = Font(name=F, size=10, italic=True, color="5B6770"); st["A2"] = None
    for i, (t, b) in enumerate(lines, 4):
        c = st.cell(row=i, column=2, value=t); c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name=F, bold=b, size=12 if b else 10, color=TEAL if b else INK)
    from openpyxl.worksheet.properties import PageSetupProperties
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    wb.active = 1 if sample else 0
    wb.save(path)

META = {
    "name": "Rental Property Manager – Excel & Google Sheets Template",
    "file_base": "Rental-Property-Manager",
    "price": 29,
    "color": "#0F766E",
    "title": "Rental Property Manager",
    "tagline": "Tenants · rent tracking · expenses · P&L — Excel & Google Sheets",
    "intro": "Run your rentals from one spreadsheet. Track tenants and leases, log rent and expenses, and see what each property really earns — no software subscription.",
    "features": [
        "Dashboard: rent collected, expenses, net income, occupancy, rent outstanding and leases ending soon, with a month-by-month chart",
        "Tenants & leases: rent status for any month (Paid, Partial, Overdue) and leases ending within 60 days flagged automatically",
        "Rent payments log: pick the tenant from a list; payments are matched to the right month even when they're late",
        "Expenses log with 12 ready categories (repairs, property tax, insurance, mortgage interest and more)",
        "Property P&L: 12-month profit and loss for each property, by expense category",
        "Up to 10 properties, 100 tenants, 600 payments and 500 expenses, in any currency",
    ],
    "tags": ["rental property", "landlord spreadsheet", "property management", "rent tracker", "tenant tracker", "google sheets", "excel template",
             "real estate investor", "rental income", "landlord template", "property expenses", "rental ledger", "airbnb host"],
}
