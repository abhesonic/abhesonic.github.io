"""Etsy Seller Profit Calculator & Sales Tracker (sample + blank versions).

Settings (editable Etsy fee rates) + Products (cost per item: materials, packaging, labour) -> fees, profit per sale,
margin, break-even price and suggested price for a target margin.
Orders log (qty, price, discount, shipping charged, sales tax, offsite ad, label cost) -> every Etsy fee per order and net profit.
Expenses log (ads, supplies, subscriptions) -> Dashboard: monthly revenue, fees, net profit, fee breakdown, best products.
"""
import random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, Reference
from _common import Kit, F, B, DATE, font, fill, style_range, input_cell, finish

MAIN = "C2410C"; SOFT = "FFF1E8"; ACC = "FDBA74"
K = Kit(MAIN, SOFT)
NP = 200; P0 = 7; P1 = P0 + NP - 1
NO = 3000; O0 = 6; O1 = O0 + NO - 1
NE = 500; E0 = 6; E1 = E0 + NE - 1
CUR = '#,##0.00;-#,##0.00;"-"'; QTY = '#,##0;-#,##0;"-"'; PCT = '0.0%;-0.0%;"-"'
def od(c): return f"Orders!${c}${O0}:${c}${O1}"
def pr(c): return f"Products!${c}${P0}:${c}${P1}"
def ex(c): return f"Expenses!${c}${E0}:${c}${E1}"
# Settings cells
LIST, TRX, PRC, PFX, OFF, OCAP, OTH, HR = ("Settings!$B$6", "Settings!$B$7", "Settings!$B$8", "Settings!$B$9", "Settings!$B$10",
                                           "Settings!$B$11", "Settings!$B$12", "Settings!$B$13")
AS = "Settings!$B$5"; YS = "Settings!$B$17"; MS = "Settings!$B$18"

PRODUCTS = [  # sku, name, category, materials, packaging, labour min, price, shipping charged, est label cost, target margin
    ("NK-INIT", "Personalized Initial Necklace", "Jewelry", 4.10, 0.90, 20, 32.00, 0, 4.50, 0.35),
    ("NK-BIRTH", "Birthstone Bar Necklace", "Jewelry", 6.80, 0.90, 30, 42.00, 0, 4.50, 0.35),
    ("ER-HOOP", "Hammered Gold Hoops", "Jewelry", 3.20, 0.70, 15, 24.00, 0, 4.50, 0.35),
    ("BR-STACK", "Stacking Bracelet Set of 3", "Jewelry", 5.50, 0.90, 25, 36.00, 0, 4.50, 0.35),
    ("PR-A4", "Botanical Art Print A4", "Prints", 1.40, 1.10, 5, 18.00, 4.95, 4.20, 0.40),
    ("PR-A3", "Botanical Art Print A3", "Prints", 2.60, 1.60, 5, 28.00, 5.95, 5.40, 0.40),
    ("DG-PLAN", "Digital Weekly Planner (PDF)", "Digital", 0, 0, 0, 7.50, 0, 0, 0.60),
    ("DG-SVG", "Wildflower SVG Bundle", "Digital", 0, 0, 0, 5.00, 0, 0, 0.60),
    ("CR-MUG", "Custom Name Mug", "Home", 5.90, 2.40, 12, 22.00, 6.50, 7.80, 0.30),
    ("CR-TOTE", "Embroidered Canvas Tote", "Home", 6.20, 1.20, 35, 34.00, 5.50, 5.20, 0.30),
]
CATS = ["Jewelry", "Prints", "Digital", "Home"]


def orders():
    random.seed(7)
    rows = []; d = dt.date(2025, 11, 2); n = 2310
    w = [10, 7, 8, 5, 6, 3, 9, 7, 4, 3]
    while d <= dt.date(2026, 10, 30):
        boost = 1.0 + (0.6 if d.month in (2, 5, 12) else 0) + ((d.month - 1) * 0.06 if d.year == 2026 else 0)
        for _ in range(sum(1 for _ in range(6) if random.random() < 0.42 * boost)):
            p = random.choices(PRODUCTS, weights=w)[0]
            q = 1 if random.random() < 0.85 else 2
            disc = round(p[6] * q * 0.15, 2) if random.random() < 0.15 else None
            ship = p[7] if p[7] else None
            tax = round((p[6] * q - (disc or 0) + p[7]) * random.choice([0, 0, 0.06, 0.0725, 0.0825]), 2) or None
            ad = "Yes" if random.random() < 0.08 else "No"
            label = p[8] if p[8] else None
            rows.append((d, f"{n}", p[0], q, None, disc, ship, tax, ad, label, "Repeat buyer" if random.random() < 0.05 else None))
            n += 1
        d += dt.timedelta(days=1)
    return rows


EXPENSES = []
def expenses():
    random.seed(3); out = []
    for y, m in [(2025, 11), (2025, 12)] + [(2026, k) for k in range(1, 11)]:
        out.append((dt.date(y, m, 1), "Etsy Ads", round(random.uniform(25, 60), 2), "Monthly ad budget"))
        out.append((dt.date(y, m, 5), "Subscriptions", 10.00, "Etsy Plus"))
        out.append((dt.date(y, m, 12), "Supplies (not per item)", round(random.uniform(15, 45), 2), "Tools, tape, printer ink"))
        if y == 2026 and m in (3, 9): out.append((dt.date(y, m, 20), "Equipment", 89.00 if m == 3 else 149.00, "Ring mandrel / heat press"))
    return out


def build(sample, path):
    wb = Workbook()
    st = wb.active; st.title = "Start Here"
    dash = wb.create_sheet("Dashboard"); ps = wb.create_sheet("Products"); osh = wb.create_sheet("Orders")
    es = wb.create_sheet("Expenses"); cc = wb.create_sheet("Calc"); s = wb.create_sheet("Settings")

    # ---------- Settings ----------
    K.title(s, "Settings", "Yellow cells are yours to edit. Fee rates are Etsy's standard US rates – change them if yours differ.")
    rows = [("Currency label", "USD", None, ""),
            ("Today / as-of date", dt.date(2026, 10, 30) if sample else "=TODAY()", DATE, "Keep =TODAY() so 'this month' stays current."),
            ("Listing fee (per item sold)", 0.20, CUR, "Etsy charges 0.20 USD per listing, and again for every extra item sold in one order."),
            ("Transaction fee %", 0.065, PCT, "6.5% of item price + shipping + gift wrap."),
            ("Payment processing %", 0.03, PCT, "Etsy Payments, US: 3% + 0.25. Other countries differ (check your Payment account)."),
            ("Payment processing fixed fee", 0.25, CUR, "Per order."),
            ("Offsite Ads fee %", 0.15, PCT, "15% (12% if you sold more than 10,000 USD in the last 12 months). Only on orders from an offsite ad."),
            ("Offsite Ads max fee per order", 100, CUR, "Etsy caps the Offsite Ads fee at 100 USD per order."),
            ("Other % fee (currency conversion, regulatory…)", 0, PCT, "e.g. 2.5% currency conversion if you list in another currency, or a regulatory operating fee."),
            ("Your hourly rate (for labour cost)", 15, CUR, "What your time is worth. Used to price labour on the Products tab."),
            ("Shop name", "Juniper & Thread" if sample else None, None, "")]
    for i, (k_, v, fmt, note) in enumerate(rows, 4):
        s.cell(row=i, column=1, value=k_).font = font(bold=True)
        c = s.cell(row=i, column=2); input_cell(c, v)
        if fmt: c.number_format = fmt
        s.cell(row=i, column=3, value=note).font = font(italic=True, size=9)
    for i, (k_, f, fmt) in enumerate([("Year starts", f"=DATE(YEAR({AS}),1,1)", DATE), ("This month starts", f"=DATE(YEAR({AS}),MONTH({AS}),1)", DATE)], 17):
        s.cell(row=i, column=1, value=k_).font = font(italic=True)
        c = s.cell(row=i, column=2, value=f); c.number_format = fmt; c.border = B; c.fill = fill("F3F4F6"); c.font = font()
    K.hdr(s, 4, ["Categories"], col0=6); K.hdr(s, 4, ["Expense types"], col0=8)
    etypes = ["Etsy Ads", "Subscriptions", "Supplies (not per item)", "Equipment", "Shipping supplies", "Photography", "Other"]
    for i in range(15):
        input_cell(s.cell(row=5 + i, column=6, value=CATS[i] if (sample and i < len(CATS)) else ("General" if (not sample and i == 0) else None)))
        input_cell(s.cell(row=5 + i, column=8, value=etypes[i] if i < len(etypes) else None))
    for c, w in zip("ABCDEFGH", [44, 14, 80, 3, 3, 16, 3, 24]): s.column_dimensions[c].width = w
    s["A16"] = "Fee rates checked against Etsy's Fees & Payments Policy (US sellers), October 2026."; s["A16"].font = font(italic=True, size=9)

    # ---------- Products ----------
    K.title(ps, "Products – pricing calculator", "One row per product. Type the yellow cells; the grey ones show your fees, profit and the price you should charge.")
    ps["A3"] = '="Amounts in "&Settings!B4&" · profit per sale = price + shipping charged − Etsy fees − cost per item − shipping label"'
    ps["A3"].font = font(bold=True)
    ps["A4"] = "Break-even price = the lowest price that still covers every cost. Suggested price = price that gives your target margin (profit ÷ what the buyer pays)."
    ps["A4"].font = font(italic=True, size=9)
    heads = ["SKU", "Product", "Category", "Materials per item", "Packaging per item", "Labour (minutes)", "Your price", "Shipping charged",
             "Est. shipping label", "Target margin",
             "Labour cost", "Cost per item", "Etsy fees per sale", "Profit per sale", "Margin", "Profit if from Offsite Ad", "Break-even price",
             "Suggested price", "Units sold", "Revenue", "Net profit (orders)", "Rank key"]
    K.hdr(ps, 6, heads, [11, 30, 11, 10, 10, 9, 9, 9, 9, 8, 9, 9, 9, 9, 8, 10, 10, 10, 8, 11, 11, 7])
    style_range(ps, P0, P1, 1, 22, inputs=range(1, 11), fmt={4: CUR, 5: CUR, 6: "0", 7: CUR, 8: CUR, 9: CUR, 10: '0%', 11: CUR, 12: CUR,
                                                            13: CUR, 14: CUR, 15: PCT, 16: CUR, 17: CUR, 18: CUR, 19: QTY, 20: CUR, 21: CUR, 22: "0.000"})
    if sample:
        for i, p in enumerate(PRODUCTS):
            for j, v in enumerate(p, 1):
                ps.cell(row=P0 + i, column=j, value=v)
    pct = f"({TRX}+{PRC}+{OTH})"
    for r in range(P0, P1 + 1):
        on = f'A{r}=""'
        ps[f"K{r}"] = f'=IF({on},"",N(F{r})/60*{HR})'
        ps[f"L{r}"] = f'=IF({on},"",N(D{r})+N(E{r})+N(K{r}))'
        ps[f"M{r}"] = f'=IF({on},"",{LIST}+(N(G{r})+N(H{r}))*{pct}+{PFX})'
        ps[f"N{r}"] = f'=IF({on},"",N(G{r})+N(H{r})-M{r}-L{r}-N(I{r}))'
        ps[f"O{r}"] = f'=IF(OR({on},N(G{r})+N(H{r})=0),"",N{r}/(N(G{r})+N(H{r})))'
        ps[f"P{r}"] = f'=IF({on},"",N{r}-MIN({OCAP},{OFF}*(N(G{r})+N(H{r}))))'
        ps[f"Q{r}"] = f'=IF({on},"",({LIST}+{PFX}+L{r}+N(I{r}))/(1-{pct})-N(H{r}))'
        ps[f"R{r}"] = f'=IF({on},"",IF(1-{pct}-N(J{r})<=0,"",({LIST}+{PFX}+L{r}+N(I{r}))/(1-{pct}-N(J{r}))-N(H{r})))'
        ps[f"S{r}"] = f'=IF({on},0,SUMIFS({od("D")},{od("C")},A{r}))'
        ps[f"T{r}"] = f'=IF({on},0,SUMIFS({od("N")},{od("C")},A{r}))'
        ps[f"U{r}"] = f'=IF({on},0,SUMIFS({od("W")},{od("C")},A{r}))'
        ps[f"V{r}"] = f'=IF(OR({on},S{r}=0),0,U{r}+10000+ROW()/100000)'
    dv = DataValidation(type="list", formula1="=Settings!$F$5:$F$19", allow_blank=True); ps.add_data_validation(dv); dv.add(f"C{P0}:C{P1}")
    ps.conditional_formatting.add(f"N{P0}:N{P1}", CellIsRule(operator="lessThan", formula=["0"], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    ps.conditional_formatting.add(f"O{P0}:O{P1}", FormulaRule(formula=[f'AND(O{P0}<>"",J{P0}<>"",O{P0}<J{P0})'], fill=fill("FEF3C7"), font=Font(name=F, bold=True, color="92400E")))
    ps.conditional_formatting.add(f"O{P0}:O{P1}", FormulaRule(formula=[f'AND(O{P0}<>"",J{P0}<>"",O{P0}>=J{P0})'], fill=fill("D1FAE5"), font=Font(name=F, color="166534")))
    ps.conditional_formatting.add(f"A{P0}:A{P1}", FormulaRule(formula=[f'AND(A{P0}<>"",COUNTIF($A${P0}:$A${P1},A{P0})>1)'], fill=fill("FCA5A5")))
    ps.freeze_panes = "C7"

    # ---------- Orders ----------
    K.title(osh, "Orders", "One row per product in an order. Leave 'Item price' empty to use the price on the Products tab.")
    osh["A3"] = "Sales tax: Etsy collects and pays it in most places, so it is NOT your income – but Etsy's payment processing fee is charged on it, so type it in."
    osh["A3"].font = font(bold=True, size=9)
    heads = ["Date", "Order #", "SKU", "Qty", "Item price (optional)", "Discount (total)", "Shipping charged", "Sales tax collected",
             "From Offsite Ad?", "Shipping label cost", "Note",
             "Product", "Product #", "Revenue", "Listing fees", "Transaction fee", "Processing fee", "Offsite Ads fee", "Other fees",
             "Total Etsy fees", "Cost of items", "Label", "Net profit", "Margin", "Month", "New order"]
    K.hdr(osh, 5, heads, [11, 9, 11, 6, 9, 9, 9, 9, 8, 9, 14, 28, 6, 10, 8, 9, 9, 9, 8, 10, 10, 8, 10, 8, 9, 7])
    style_range(osh, O0, O1, 1, 26, inputs=range(1, 12), fmt={1: DATE, 4: QTY, 5: CUR, 6: CUR, 7: CUR, 8: CUR, 10: CUR, 13: "0", 14: CUR, 15: CUR,
                                                              16: CUR, 17: CUR, 18: CUR, 19: CUR, 20: CUR, 21: CUR, 22: CUR, 23: CUR, 24: PCT, 25: "mmm yy;;"})
    if sample:
        for i, row in enumerate(orders()):
            for j, v in enumerate(row, 1):
                if v is not None: osh.cell(row=O0 + i, column=j, value=v)
    for r in range(O0, O1 + 1):
        osh[f"M{r}"] = f'=IF(C{r}="",0,IFERROR(MATCH(C{r},{pr("A")},0),0))'
        osh[f"L{r}"] = f'=IF(C{r}="","",IF(M{r}=0,"SKU not found",INDEX({pr("B")},M{r})))'
        price = f"IF(N(E{r})>0,E{r},IF(M{r}=0,0,N(INDEX({pr('G')},M{r}))))"
        osh[f"N{r}"] = f'=IF(C{r}="",0,N(D{r})*{price}-N(F{r})+N(G{r}))'
        osh[f"O{r}"] = f'=IF(C{r}="",0,N(D{r})*{LIST})'
        osh[f"P{r}"] = f"=N{r}*{TRX}"
        osh[f"Z{r}"] = f'=IF(C{r}="",0,IF(B{r}="",1,IF(COUNTIF($B${O0}:B{r},B{r})=1,1,0)))'
        osh[f"Q{r}"] = f'=IF(C{r}="",0,Z{r}*{PFX}+(N{r}+N(H{r}))*{PRC})'
        osh[f"R{r}"] = f'=IF(I{r}="Yes",MIN({OCAP},N{r}*{OFF}),0)'
        osh[f"S{r}"] = f"=N{r}*{OTH}"
        osh[f"T{r}"] = f"=SUM(O{r}:S{r})"
        osh[f"U{r}"] = f'=IF(M{r}=0,0,N(D{r})*N(INDEX({pr("L")},M{r})))'
        osh[f"V{r}"] = f"=N(J{r})"
        osh[f"W{r}"] = f"=N{r}-T{r}-U{r}-V{r}"
        osh[f"X{r}"] = f'=IF(N{r}=0,"",W{r}/N{r})'
        osh[f"Y{r}"] = f'=IF(N(A{r})=0,0,DATE(YEAR(A{r}),MONTH(A{r}),1))'
    osh["B4"] = "Same order # on several rows = the 0.25 fixed fee is charged once."; osh["B4"].font = font(italic=True, size=9)
    dv = DataValidation(type="list", formula1=f"={pr('A')}", allow_blank=True); osh.add_data_validation(dv); dv.add(f"C{O0}:C{O1}")
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True); osh.add_data_validation(dv); dv.add(f"I{O0}:I{O1}")
    osh.conditional_formatting.add(f"L{O0}:L{O1}", FormulaRule(formula=[f'L{O0}="SKU not found"'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    osh.conditional_formatting.add(f"W{O0}:W{O1}", CellIsRule(operator="lessThan", formula=["0"], fill=fill("FEE2E2"), font=Font(name=F, color="B91C1C")))
    osh.freeze_panes = "D6"

    # ---------- Expenses ----------
    K.title(es, "Other Expenses", "Costs that are not per item: Etsy Ads, subscriptions, tools, photography…")
    K.hdr(es, 5, ["Date", "Type", "Amount", "Note", "Month"], [11, 24, 11, 34, 9])
    style_range(es, E0, E1, 1, 5, inputs=range(1, 5), fmt={1: DATE, 3: CUR, 5: "mmm yy;;"})
    if sample:
        for i, row in enumerate(expenses()):
            for j, v in enumerate(row, 1): es.cell(row=E0 + i, column=j, value=v)
    for r in range(E0, E1 + 1):
        es[f"E{r}"] = f'=IF(N(A{r})=0,0,DATE(YEAR(A{r}),MONTH(A{r}),1))'
    dv = DataValidation(type="list", formula1="=Settings!$H$5:$H$19", allow_blank=True); es.add_data_validation(dv); dv.add(f"B{E0}:B{E1}")
    es.freeze_panes = "A6"

    # ---------- Calc ----------
    K.title(cc, "Calc", "Numbers for the dashboard charts. Calculated – nothing to type.")
    K.hdr(cc, 4, ["Month", "Label", "Revenue", "Etsy fees", "Profit from orders", "Other expenses", "Net profit", "Orders", "Units"],
          [11, 9, 12, 11, 12, 12, 12, 8, 8])
    for i in range(12):
        r = 5 + i
        cc[f"A{r}"] = f"=DATE(YEAR({AS}),MONTH({AS})-{11 - i},1)"
        cc[f"B{r}"] = f'=TEXT(A{r},"mmm yy")'
        cc[f"C{r}"] = f'=SUMIFS({od("N")},{od("Y")},A{r})'
        cc[f"D{r}"] = f'=SUMIFS({od("T")},{od("Y")},A{r})'
        cc[f"E{r}"] = f'=SUMIFS({od("W")},{od("Y")},A{r})'
        cc[f"F{r}"] = f'=SUMIFS({ex("C")},{ex("E")},A{r})'
        cc[f"G{r}"] = f"=E{r}-F{r}"
        cc[f"H{r}"] = f'=SUMIFS({od("Z")},{od("Y")},A{r})'
        cc[f"I{r}"] = f'=SUMIFS({od("D")},{od("Y")},A{r})'
    style_range(cc, 5, 16, 1, 9, fmt={1: "mmm yyyy", 3: CUR, 4: CUR, 5: CUR, 6: CUR, 7: CUR, 8: QTY, 9: QTY})

    # ---------- Dashboard ----------
    K.title(dash, "Etsy Profit Dashboard", "Updates by itself from the Products, Orders and Expenses tabs.")
    dash["A3"] = f'=IF(Settings!B14="","",Settings!B14&"  ·  ")&"As of "&TEXT({AS},"dd mmm yyyy")&"  ·  amounts in "&Settings!B4'
    dash["A3"].font = font(bold=True)
    inm = f'{od("A")},">="&{MS},{od("A")},"<="&{AS}'; iny = f'{od("A")},">="&{YS},{od("A")},"<="&{AS}'
    exm = f'{ex("A")},">="&{MS},{ex("A")},"<="&{AS}'; exy = f'{ex("A")},">="&{YS},{ex("A")},"<="&{AS}'
    kp = [("Revenue this month", f'=SUMIFS({od("N")},{inm})', "#,##0.00"),
          ("Etsy fees this month", f'=SUMIFS({od("T")},{inm})', "#,##0.00"),
          ("Net profit this month", f'=SUMIFS({od("W")},{inm})-SUMIFS({ex("C")},{exm})', "#,##0.00"),
          ("Revenue this year", f'=SUMIFS({od("N")},{iny})', "#,##0.00"),
          ("Etsy fees this year", f'=SUMIFS({od("T")},{iny})', "#,##0.00"),
          ("Net profit this year", f'=SUMIFS({od("W")},{iny})-SUMIFS({ex("C")},{exy})', "#,##0.00"),
          ("Units sold this year", f'=SUMIFS({od("D")},{iny})', "#,##0"),
          ("Fees as % of revenue (year)", "=IFERROR(C9/A9,0)", "0.0%"),
          ("Net margin (year)", "=IFERROR(E9/A9,0)", "0.0%")]
    for i, (lab, f, fmt) in enumerate(kp):
        K.kpi(dash, 5 + (i // 3) * 3, 1 + (i % 3) * 2, lab, f, fmt)
        dash.cell(row=6 + (i // 3) * 3, column=1 + (i % 3) * 2).alignment = Alignment(horizontal="left")
    for c, w in zip("ABCDEFGHIJKL", [30, 12, 20, 12, 20, 12, 3, 30, 12, 12, 12, 10]): dash.column_dimensions[c].width = w
    K.section(dash, "A15", "Where your money goes (this year)")
    K.hdr(dash, 16, ["Item", "Amount", "% of revenue"])
    br = [("Revenue (items + shipping)", f'=SUMIFS({od("N")},{iny})'), ("Listing fees", f'=-SUMIFS({od("O")},{iny})'),
          ("Transaction fees", f'=-SUMIFS({od("P")},{iny})'), ("Payment processing", f'=-SUMIFS({od("Q")},{iny})'),
          ("Offsite Ads fees", f'=-SUMIFS({od("R")},{iny})'), ("Other fees", f'=-SUMIFS({od("S")},{iny})'),
          ("Cost of items", f'=-SUMIFS({od("U")},{iny})'), ("Shipping labels", f'=-SUMIFS({od("V")},{iny})'),
          ("Other expenses", f'=-SUMIFS({ex("C")},{exy})'), ("Net profit", "=SUM(B17:B25)")]
    for i, (lab, f) in enumerate(br):
        r = 17 + i
        dash[f"A{r}"] = lab; dash[f"B{r}"] = f; dash[f"C{r}"] = f"=IFERROR(B{r}/$B$17,0)"
    style_range(dash, 17, 26, 1, 3, fmt={2: CUR, 3: PCT})
    for c in "ABC": dash[f"{c}26"].font = Font(name=F, bold=True, color=MAIN)
    K.section(dash, "H15", "Best products by net profit (all orders)")
    K.hdr(dash, 16, ["Product", "Units", "Revenue", "Net profit", "Margin"], col0=8)
    for n in range(10):
        r = 17 + n
        m = f"MATCH(LARGE({pr('V')},{n + 1}),{pr('V')},0)"
        dash[f"H{r}"] = f'=IF(LARGE({pr("V")},{n + 1})=0,"",INDEX({pr("B")},{m}))'
        for c, src in zip("IJK", "STU"):
            dash[f"{c}{r}"] = f'=IF($H{r}="","",INDEX({pr(src)},{m}))'
        dash[f"L{r}"] = f'=IF(OR($H{r}="",N(J{r})=0),"",K{r}/J{r})'
    style_range(dash, 17, 26, 8, 12, fmt={9: QTY, 10: CUR, 11: CUR, 12: PCT})
    dash.conditional_formatting.add("K17:K26", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, bold=True, color="B91C1C")))
    c1 = BarChart(); c1.type = "col"; c1.title = "Revenue, Etsy fees and net profit by month"; c1.height = 7.5; c1.width = 17
    c1.add_data(Reference(cc, min_col=3, max_col=4, min_row=4, max_row=16), titles_from_data=True)
    c1.add_data(Reference(cc, min_col=7, max_col=7, min_row=4, max_row=16), titles_from_data=True)
    c1.set_categories(Reference(cc, min_col=2, min_row=5, max_row=16))
    for ser, col in zip(c1.series, (MAIN, ACC, "15803D")): ser.graphicalProperties.solidFill = col
    c1.y_axis.majorGridlines = None; c1.y_axis.number_format = "#,##0"
    dash.add_chart(c1, "A29")
    c2 = LineChart(); c2.title = "Orders per month"; c2.height = 7.5; c2.width = 14; c2.legend = None
    c2.add_data(Reference(cc, min_col=8, min_row=4, max_row=16), titles_from_data=True)
    c2.set_categories(Reference(cc, min_col=2, min_row=5, max_row=16))
    c2.series[0].graphicalProperties.line.solidFill = MAIN; c2.series[0].graphicalProperties.line.width = 28000; c2.series[0].smooth = False
    c2.y_axis.majorGridlines = None; c2.y_axis.scaling.min = 0
    dash.add_chart(c2, "H29")

    # ---------- Start Here ----------
    lines = [("How to start", True),
             ("1.  Settings tab: check the Etsy fee rates (US standard rates are filled in), set your hourly rate and your categories.", False),
             ("2.  Products tab: one row per product – materials, packaging, minutes of work, your price, shipping you charge and what the label costs you.", False),
             ("     You instantly see Etsy fees, profit per sale, margin, the break-even price and a suggested price for your target margin.", False),
             ("3.  Orders tab: add one row per product sold (copy from your Etsy Orders or the monthly CSV). Every fee and the net profit are calculated.", False),
             ("4.  Expenses tab: Etsy Ads, subscriptions, tools and other costs that are not per item.", False),
             ("5.  Dashboard: revenue, fees and net profit this month and this year, where your money goes, best products and charts.", False),
             ("", False), ("How fees are calculated", True),
             ("Listing fee 0.20 per item sold · Transaction fee 6.5% of item price + shipping · Payment processing 3% of the order total (incl. sales tax) + 0.25 per order", False),
             ("Offsite Ads 15% of the order (max 100) only when you mark the order 'From Offsite Ad = Yes'. All rates can be changed on the Settings tab.", False),
             ("Margin = net profit ÷ what the buyer paid (excluding sales tax). A red profit means you lose money on that product.", False),
             ("", False), ("Colour key", True),
             ("Yellow cells are inputs: type here.  Grey cells calculate themselves: please don't overwrite them.", False),
             ("", False), ("Good to know", True),
             ("•  Room for 200 products, 3,000 order lines and 500 expenses. Each SKU must be unique (duplicates turn red).", False),
             ("•  Selling outside the US? Change the payment processing rate on Settings to your country's rate.", False),
             ("•  Not tax advice – keep Etsy's own monthly statement as your official record.", False),
             ("•  Google Sheets: open drive.google.com, upload this file, then choose Open with > Google Sheets.", False)]
    if sample:
        lines += [("", False), ("This copy contains sample data", True),
                  ("A made-up jewelry, print and gift shop with 10 products and 12 months of orders, shown as of 30 Oct 2026. Use the blank copy for your own shop.", False)]
    K.start(st, "Etsy Seller Profit Calculator", "See your real profit after every Etsy fee – per product, per order and per month. Works in Excel and Google Sheets.", lines)
    finish(wb, sample, path)


META = {
    "name": "Etsy Seller Profit Calculator & Sales Tracker – Excel & Google Sheets",
    "file_base": "Etsy-Seller-Profit-Calculator",
    "price": 17,
    "color": "#C2410C",
    "title": "Etsy Seller Profit Calculator",
    "tagline": "Every Etsy fee · real profit per product & order — Excel & Google Sheets",
    "intro": "Know exactly what you keep from every Etsy sale. Enter your costs once, log your orders, and the calculator takes off every Etsy fee (listing, transaction, payment processing, Offsite Ads) and shows your real profit per product, per order and per month.",
    "features": [
        "Pricing calculator: Etsy fees, profit per sale, margin, break-even price and a suggested price for your target margin",
        "Cost per item from materials, packaging and your time (hourly rate)",
        "Orders log for up to 3,000 lines: listing, transaction, processing and Offsite Ads fees calculated for every order",
        "Sales tax and discounts handled correctly; multi-item orders charged the fixed fee once",
        "Expenses tab for Etsy Ads, subscriptions and supplies",
        "Dashboard: revenue, Etsy fees and net profit this month and this year, fee breakdown, best products and monthly charts",
        "Fee rates are editable – works for any country and future Etsy fee changes",
    ],
    "tags": ["etsy seller", "etsy profit calculator", "etsy fee calculator", "pricing calculator", "small business", "google sheets",
             "excel template", "etsy shop", "profit tracker", "handmade business", "sales tracker", "etsy spreadsheet", "product pricing"],
}
