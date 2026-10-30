"""Habit Tracker – yearly & monthly heat-map (sample + blank versions).

Habits tab (up to 15 habits, target days per week, optional start date) -> Daily Log (one row per day of the year,
type x when done) -> Monthly View (one month grid), Year Heatmap (one habit, 12 x 31 calendar), Dashboard
(today, this week, this month, year score, streaks, monthly heat-map by habit, charts).
Score = days done / goal days (days elapsed x target per week / 7), capped at 100%.
"""
import random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as L
from _common import Kit, F, INK, B, DATE, font, fill, style_range, input_cell, finish

MAIN = "7C3AED"; SOFT = "F3EEFF"; DONE = "A78BFA"; MISS = "F1F1F4"
K = Kit(MAIN, SOFT)
NH = 15                    # habits
H0 = 6; H1 = H0 + NH - 1   # Habits rows 6..20
D0 = 6; D1 = D0 + 365      # Daily Log rows 6..371 (366 days)
LOGC = [L(3 + k) for k in range(NH)]     # C..Q habit columns
STRC = [L(21 + k) for k in range(NH)]    # U..AI streak helpers
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
PCT = '0%'
DL = "'Daily Log'"
Y, YS, PE, AS = "Settings!$B$4", "Settings!$B$8", "Settings!$B$10", "Settings!$B$5"
GRID = '"✓";"";"·"'


def sample_log():
    random.seed(7)
    habits = [("Drink 8 glasses of water", "Health", 7, None, 0.82), ("Read 20 minutes", "Learning", 7, None, 0.74),
              ("Workout", "Fitness", 4, None, 0.48), ("Meditate 10 minutes", "Mind", 7, None, 0.66),
              ("No phone after 10 pm", "Mind", 7, None, 0.58), ("Walk 8,000 steps", "Fitness", 5, None, 0.62),
              ("Journal", "Mind", 5, None, 0.55), ("Learn Spanish (15 min)", "Learning", 5, dt.date(2026, 3, 1), 0.6),
              ("No added sugar", "Health", 7, None, 0.62), ("In bed by 11 pm", "Health", 7, dt.date(2026, 6, 1), 0.7)]
    end = dt.date(2026, 10, 30); log = {}
    for k, (_, _, _, start, p) in enumerate(habits):
        d = dt.date(2026, 1, 1); prev = False
        while d <= end:
            if start is None or d >= start:
                trend = min(0.12, (d - dt.date(2026, 1, 1)).days / 2400)   # habits get stronger over the year
                q = min(0.93, p + trend + (0.05 if prev else -0.10))
                prev = random.random() < q
                if prev: log[(d, k)] = "x"
            d += dt.timedelta(days=1)
    return habits, log


def build(sample, path):
    wb = Workbook()
    st = wb.active; st.title = "Start Here"
    dash = wb.create_sheet("Dashboard"); hb = wb.create_sheet("Habits"); lg = wb.create_sheet("Daily Log")
    mv = wb.create_sheet("Monthly View"); yh = wb.create_sheet("Year Heatmap"); cc = wb.create_sheet("Calc"); s = wb.create_sheet("Settings")

    # ---------- Settings ----------
    K.title(s, "Settings", "Yellow cells are yours to edit.")
    rows = [("Year", 2026 if sample else "=YEAR(TODAY())", "0"),
            ("Today / as-of date", dt.date(2026, 10, 30) if sample else "=TODAY()", DATE)]
    for i, (k_, v, fmt) in enumerate(rows, 4):
        s.cell(row=i, column=1, value=k_).font = font(bold=True)
        c = s.cell(row=i, column=2); input_cell(c, v); c.number_format = fmt
    s["C4"] = "Starts as the current year. For a new year, save a copy of the file and change the year."
    s["C5"] = "Keep =TODAY() so the dashboard always shows today." + (" (The sample uses a fixed date.)" if sample else "")
    for a in ("C4", "C5"): s[a].font = font(italic=True, size=9)
    calc = [("Year start", "=DATE(B4,1,1)", DATE), ("Year end", "=DATE(B4,12,31)", DATE),
            ("Counted up to", "=MAX(B8-1,MIN(B5,B9))", DATE), ("Days in year", "=B9-B8+1", "0"),
            ("Week starts (Monday)", "=B5-WEEKDAY(B5,2)+1", DATE), ("Day of year (as-of)", "=B5-B8+1", "0"),
            ("Active habits", f"=SUM(Habits!$G${H0}:$G${H1})", "0"), ("As-of month in this year", '=IF(YEAR(B5)=B4,MONTH(B5),0)', "0")]
    for i, (k_, f, fmt) in enumerate(calc, 8):
        s.cell(row=i, column=1, value=k_).font = font(italic=True)
        c = s.cell(row=i, column=2, value=f); c.number_format = fmt; c.border = B; c.fill = fill("F3F4F6"); c.font = font()
    K.hdr(s, 4, ["Categories"], col0=8)
    cats = ["Health", "Fitness", "Mind", "Learning", "Work", "Money", "Home", "Social", "Creative", "Other"]
    for i in range(12):
        c = s.cell(row=5 + i, column=8, value=cats[i] if i < len(cats) else None); input_cell(c)
    K.hdr(s, 4, ["Months"], col0=10)
    for i, m in enumerate(MONTHS):
        c = s.cell(row=5 + i, column=10, value=m); c.font = font(size=9); c.border = B
    for c, w in zip("ABCDEFGHIJ", [24, 14, 3, 3, 3, 3, 3, 16, 3, 13]): s.column_dimensions[c].width = w
    s["A18"] = "Week starts on Monday. The Monthly score uses your 'target days per week' so a 4x-a-week habit can still reach 100%."
    s["A18"].font = font(italic=True, size=9)

    # ---------- Habits ----------
    K.title(hb, "Your Habits", "Up to 15 habits. Yellow = type here. Grey columns are calculated.")
    heads = ["#", "Habit", "Category", "Target days per week", "Start date (optional)", "Why it matters / note",
             "Active", "Counting from", "Days so far", "Target / week", "Days done", "Goal days", "Year score",
             "Current streak", "Longest streak", "Done this week", "Done this month", "Month score"]
    K.hdr(hb, 5, heads, [4, 28, 12, 10, 12, 28, 7, 12, 8, 8, 8, 8, 9, 9, 9, 9, 9, 9])
    style_range(hb, H0, H1, 1, 18, inputs=(2, 3, 4, 5, 6), fmt={5: DATE, 8: DATE, 12: "0.0", 13: PCT, 18: PCT})
    habits, log = sample_log() if sample else ([], {})
    for k in range(NH):
        r = H0 + k; lc = LOGC[k]; sc = STRC[k]
        hb[f"A{r}"] = k + 1
        if k < len(habits):
            name, cat, tgt, start, _ = habits[k]
            hb[f"B{r}"] = name; hb[f"C{r}"] = cat; hb[f"D{r}"] = tgt
            if start: hb[f"E{r}"] = start
        dates = f"{DL}!$A${D0}:$A${D1}"; col = f"{DL}!${lc}${D0}:${lc}${D1}"
        hb[f"G{r}"] = f'=IF(B{r}<>"",1,0)'
        hb[f"H{r}"] = f'=IF(G{r}=1,MAX({YS},N(E{r})),"")'
        hb[f"I{r}"] = f"=IF(G{r}=1,MAX(0,{PE}-H{r}+1),0)"
        hb[f"J{r}"] = f"=IF(G{r}=1,MIN(7,MAX(1,IF(N(D{r})=0,7,ROUND(N(D{r}),0)))),0)"
        hb[f"K{r}"] = f'=IF(G{r}=1,COUNTIFS({dates},">="&H{r},{dates},"<="&{PE},{col},"<>"),0)'
        hb[f"L{r}"] = f"=I{r}*J{r}/7"
        hb[f"M{r}"] = f'=IF(L{r}>0,MIN(1,K{r}/L{r}),"")'
        st_ = f"{DL}!${sc}${D0}:${sc}${D1}"; idx = "Settings!$B$13"
        hb[f"N{r}"] = (f'=IF(OR(G{r}=0,{idx}<1,{idx}>Settings!$B$11),0,IF(INDEX({st_},{idx})>0,INDEX({st_},{idx}),'
                       f'IF({idx}>1,INDEX({st_},{idx}-1),0)))')
        hb[f"O{r}"] = f"=IF(G{r}=1,MAX({st_}),0)"
        hb[f"P{r}"] = f'=IF(G{r}=1,COUNTIFS({dates},">="&Settings!$B$12,{dates},"<="&{AS},{col},"<>"),0)'
        hb[f"Q{r}"] = f'=IF(G{r}=1,COUNTIFS({dates},">="&MAX(H{r},DATE(YEAR({AS}),MONTH({AS}),1)),{dates},"<="&{AS},{col},"<>"),0)'
        hb[f"R{r}"] = f'=IF(OR(G{r}=0,Settings!$B$15=0),"",Calc!{L(2 + 0)}1)'   # placeholder, replaced below
        hb[f"R{r}"] = f'=IF(OR(G{r}=0,Settings!$B$15=0),"",INDEX(Calc!$B${42 + k}:$M${42 + k},Settings!$B$15))'
    dv = DataValidation(type="list", formula1="=Settings!$H$5:$H$16", allow_blank=True); hb.add_data_validation(dv); dv.add(f"C{H0}:C{H1}")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="7", allow_blank=True,
                        error="Type a number from 1 to 7", showErrorMessage=True); hb.add_data_validation(dv); dv.add(f"D{H0}:D{H1}")
    for c in ("M", "R"):
        hb.conditional_formatting.add(f"{c}{H0}:{c}{H1}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=1, end_color=DONE))
    hb["A22"] = "Target days per week: 7 = every day, 3 = three times a week. Leave the start date empty to count from 1 January."
    hb["A23"] = "Year score = days done ÷ goal days so far (capped at 100%). A streak counts days in a row; today still counts if you haven't ticked it yet."
    for a in ("A22", "A23"): hb[a].font = font(italic=True, size=9)
    hb.freeze_panes = "C6"

    # ---------- Daily Log ----------
    K.title(lg, "Daily Log", "Type x (or anything) in a habit's column on the days you did it. One row per day of the year.")
    lg["A3"] = '="Year "&Settings!B4&"  ·  Today is highlighted"'; lg["A3"].font = font(bold=True)
    heads = ["Date", "Day"] + [f"Habit {k+1}" for k in range(NH)] + ["Done", "Done %", ""] + [f"Streak {k+1}" for k in range(NH)]
    K.hdr(lg, 5, heads, [11, 6] + [7] * NH + [7, 8, 2] + [6] * NH)
    for k in range(NH):
        lg[f"{LOGC[k]}5"] = f'=IF(Habits!$B${H0+k}<>"",Habits!$B${H0+k},"Habit {k+1}")'
        lg[f"{STRC[k]}4"] = None
    lg["U4"] = "Streak helpers (calculated, days in a row)"; lg["U4"].font = font(italic=True, size=9)
    lg.row_dimensions[5].height = 66
    for c in LOGC: lg[f"{c}5"].alignment = Alignment(text_rotation=90, horizontal="center", vertical="bottom", wrap_text=True)
    style_range(lg, D0, D1, 1, 19, inputs=range(3, 18), fmt={1: "ddd dd mmm", 19: PCT})
    style_range(lg, D0, D1, 21, 35)
    order = {d: i for i, d in enumerate([dt.date(2026, 1, 1) + dt.timedelta(days=n) for n in range(366)])}
    for r in range(D0, D1 + 1):
        n = r - D0
        lg[f"A{r}"] = f'=IF({n + 1}<=Settings!$B$11,Settings!$B$8+{n},"")'
        lg[f"B{r}"] = f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r},2),"Mon","Tue","Wed","Thu","Fri","Sat","Sun"))'
        lg[f"R{r}"] = f'=IF(A{r}="","",COUNTA(C{r}:Q{r}))'
        lg[f"S{r}"] = f'=IF(OR(A{r}="",Settings!$B$14=0),"",R{r}/Settings!$B$14)'
        for k in range(NH):
            lg[f"{STRC[k]}{r}"] = f'=IF({LOGC[k]}{r}<>"",N({STRC[k]}{r-1})+1,0)'
            lg[f"{LOGC[k]}{r}"].alignment = Alignment(horizontal="center")
    for (d, k), v in log.items():
        lg[f"{LOGC[k]}{D0 + order[d]}"] = v
    lg.conditional_formatting.add(f"C{D0}:Q{D1}", FormulaRule(formula=[f'C{D0}<>""'], fill=fill(DONE), font=Font(name=F, bold=True, color="FFFFFF")))
    lg.conditional_formatting.add(f"A{D0}:B{D1}", FormulaRule(formula=[f"$A{D0}={AS}"], fill=fill("FDE68A"), font=Font(name=F, bold=True)))
    lg.conditional_formatting.add(f"A{D0}:B{D1}", FormulaRule(formula=[f'AND($A{D0}<>"",$A{D0}>{AS})'], font=Font(name=F, color="9CA3AF")))
    lg.freeze_panes = "C6"

    # ---------- Calc (month x habit) ----------
    K.title(cc, "Calc", "Monthly numbers for the dashboard. Calculated – nothing to type.")
    blocks = [(5, "Days done"), (23, "Days counted"), (41, "Score (done ÷ goal)")]
    for top, name in blocks:
        cc.cell(row=top - 1, column=1, value=name).font = font(bold=True)
        K.hdr(cc, top, ["Habit"] + [m[:3] for m in MONTHS] + ["Year"])
    # block rows: done 6..20, days 24..38, score 42..56 ; overall row 57
    for k in range(NH):
        hr = H0 + k; lc = LOGC[k]
        dates = f"{DL}!$A${D0}:$A${D1}"; col = f"{DL}!${lc}${D0}:${lc}${D1}"
        for top in (5, 23, 41):
            cc.cell(row=top + 1 + k, column=1, value=f'=IF(Habits!$B${hr}<>"",Habits!$B${hr},"")')
        for m in range(1, 13):
            c = L(1 + m); ms = f"DATE({Y},{m},1)"; me = f"MIN(EOMONTH({ms},0),{PE})"
            cc[f"{c}{6+k}"] = f'=IF(Habits!$G${hr}=0,0,COUNTIFS({dates},">="&MAX({ms},Habits!$H${hr}),{dates},"<="&{me},{col},"<>"))'
            cc[f"{c}{24+k}"] = f"=IF(Habits!$G${hr}=0,0,MAX(0,{me}-MAX({ms},Habits!$H${hr})+1))"
            cc[f"{c}{42+k}"] = f'=IF({c}{24+k}=0,"",MIN(1,{c}{6+k}/({c}{24+k}*Habits!$J${hr}/7)))'
        cc[f"N{6+k}"] = f"=SUM(B{6+k}:M{6+k})"; cc[f"N{24+k}"] = f"=SUM(B{24+k}:M{24+k})"
        cc[f"N{42+k}"] = f"=Habits!M{hr}"
    cc["A57"] = "All habits"
    for m in range(1, 14):
        c = L(1 + m)
        goal = f"SUMPRODUCT({c}24:{c}38,Habits!$J${H0}:$J${H1})/7"
        cc[f"{c}57"] = f'=IF({goal}=0,"",MIN(1,SUM({c}6:{c}20)/({goal})))'
    cc["A59"] = "Chart: month"; cc["B59"] = "Score"
    for m in range(12):
        cc[f"A{60+m}"] = MONTHS[m][:3]; cc[f"B{60+m}"] = f'=IF({L(2+m)}57="",0,{L(2+m)}57)'
    style_range(cc, 6, 20, 1, 14); style_range(cc, 24, 38, 1, 14); style_range(cc, 42, 57, 1, 14, fmt={c: PCT for c in range(2, 15)})
    style_range(cc, 60, 71, 1, 2, fmt={2: PCT})
    cc.column_dimensions["A"].width = 26

    # ---------- Monthly View ----------
    K.title(mv, "Monthly View", "Pick a month in the yellow cell. ✓ = done, · = missed. Calculated from the Daily Log.")
    mv["A3"] = "Month:"; mv["A3"].font = font(bold=True)
    input_cell(mv["B3"], "=INDEX(Settings!$J$5:$J$16,MONTH(Settings!$B$5))")
    dv = DataValidation(type="list", formula1="=Settings!$J$5:$J$16", allow_blank=False); mv.add_data_validation(dv); dv.add("B3")
    mv["D3"] = "Month #"; mv["E3"] = "=IFERROR(MATCH(B3,Settings!$J$5:$J$16,0),1)"
    mv["F3"] = "Days"; mv["G3"] = f"=DAY(EOMONTH(DATE({Y},E3,1),0))"
    mv["I3"] = f'=B3&" "&{Y}'; mv["I3"].font = Font(name=F, bold=True, size=12, color=MAIN)
    for a in ("D3", "F3"): mv[a].font = font(italic=True, size=9)
    K.hdr(mv, 6, ["Habit", "Target / week"] + [str(d) for d in range(1, 32)] + ["Done", "Days", "Score"])
    mv.column_dimensions["A"].width = 26; mv.column_dimensions["B"].width = 8
    for d in range(1, 32):
        c = L(2 + d); mv.column_dimensions[c].width = 3.6
        mv[f"{c}6"] = f'=IF({d}<=$G$3,{d},"")'; mv[f"{c}6"].alignment = Alignment(horizontal="center")
        mv[f"{c}5"] = f'=IF({d}<=$G$3,MID("MTWTFSS",WEEKDAY(DATE({Y},$E$3,{d}),2),1),"")'
        mv[f"{c}5"].font = font(size=8, color="6B7280"); mv[f"{c}5"].alignment = Alignment(horizontal="center")
    for c, w in zip(("AH", "AI", "AJ"), (6, 6, 7)): mv.column_dimensions[c].width = w
    LOGR = f"{DL}!$C${D0}:$Q${D1}"
    for k in range(NH):
        r = 7 + k; hr = H0 + k
        mv[f"A{r}"] = f'=IF(Habits!$B${hr}<>"",Habits!$B${hr},"")'
        mv[f"B{r}"] = f'=IF(A{r}="","",Habits!$J${hr})'
        for d in range(1, 32):
            c = L(2 + d); dd = f"DATE({Y},$E$3,{d})"
            mv[f"{c}{r}"] = (f'=IF(OR($A{r}="",{d}>$G$3),"",IF(OR({dd}>{AS},{dd}<Habits!$H${hr}),"",'
                             f'IF(INDEX({LOGR},{dd}-{YS}+1,{k+1})<>"",1,0)))')
        mv[f"AH{r}"] = f'=IF(A{r}="","",SUM(C{r}:AG{r}))'
        mv[f"AI{r}"] = f'=IF(A{r}="","",COUNT(C{r}:AG{r}))'
        mv[f"AJ{r}"] = f'=IF(OR(A{r}="",N(AI{r})=0),"",MIN(1,AH{r}/(AI{r}*B{r}/7)))'
    mv["A22"] = "All habits – done %"
    for d in range(1, 32):
        c = L(2 + d); mv[f"{c}22"] = f'=IF(COUNT({c}7:{c}21)=0,"",SUM({c}7:{c}21)/COUNT({c}7:{c}21))'
        mv[f"{c}22"].number_format = '0%'; mv[f"{c}22"].font = font(size=7)
    mv["AH22"] = "=SUM(AH7:AH21)"; mv["AI22"] = "=SUM(AI7:AI21)"
    style_range(mv, 7, 22, 1, 36, fmt={**{c: GRID for c in range(3, 34)}, 36: PCT})
    for r in range(7, 23):
        for d in range(1, 32): mv[f"{L(2+d)}{r}"].alignment = Alignment(horizontal="center")
        if r == 22:
            for d in range(1, 32): mv[f"{L(2+d)}22"].number_format = '0%'; mv[f"{L(2+d)}22"].font = font(size=7)
    mv.conditional_formatting.add("C7:AG21", FormulaRule(formula=["C7=1"], fill=fill(DONE), font=Font(name=F, bold=True, color="FFFFFF")))
    mv.conditional_formatting.add("C7:AG21", FormulaRule(formula=['AND(C7<>"",C7=0)'], fill=fill(MISS), font=Font(name=F, color="9CA3AF")))
    mv.conditional_formatting.add("AJ7:AJ21", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=1, end_color=DONE))
    mv.conditional_formatting.add("C22:AG22", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=1, end_color=DONE))
    mv.freeze_panes = "C7"

    # ---------- Year Heatmap ----------
    K.title(yh, "Year Heatmap", "Pick one habit in the yellow cell to see its whole year. Calculated from the Daily Log.")
    yh["A3"] = "Habit:"; yh["A3"].font = font(bold=True)
    input_cell(yh["B3"], f"=Habits!B{H0}")
    dv = DataValidation(type="list", formula1=f"=Habits!$B${H0}:$B${H1}", allow_blank=False); yh.add_data_validation(dv); dv.add("B3")
    yh.merge_cells("B3:N3")
    yh["P3"] = "Habit #"; yh["Q3"] = f"=IFERROR(MATCH(B3,Habits!$B${H0}:$B${H1},0),0)"
    yh["S3"] = "From"; yh["T3"] = f'=IF(Q3=0,{YS},INDEX(Habits!$H${H0}:$H${H1},Q3))'; yh["T3"].number_format = DATE
    yh["V3"] = "Target"; yh["W3"] = f"=IF(Q3=0,7,INDEX(Habits!$J${H0}:$J${H1},Q3))"
    for a in ("P3", "S3", "V3"): yh[a].font = font(italic=True, size=9)
    K.hdr(yh, 5, ["Month"] + [str(d) for d in range(1, 32)] + ["Done", "Days", "Score"])
    yh.column_dimensions["A"].width = 12
    for d in range(1, 32): yh.column_dimensions[L(1 + d)].width = 3.6
    for m in range(1, 13):
        r = 5 + m; yh[f"A{r}"] = MONTHS[m - 1]
        for d in range(1, 32):
            c = L(1 + d); dd = f"DATE({Y},{m},{d})"
            yh[f"{c}{r}"] = (f'=IF(OR($Q$3=0,{d}>DAY(EOMONTH(DATE({Y},{m},1),0))),"",IF(OR({dd}>{AS},{dd}<$T$3),"",'
                             f'IF(INDEX({LOGR},{dd}-{YS}+1,$Q$3)<>"",1,0)))')
        yh[f"AG{r}"] = f"=SUM(B{r}:AF{r})"; yh[f"AH{r}"] = f"=COUNT(B{r}:AF{r})"
        yh[f"AI{r}"] = f'=IF(AH{r}=0,"",MIN(1,AG{r}/(AH{r}*$W$3/7)))'
    yh["A18"] = "Year"; yh["AG18"] = "=SUM(AG6:AG17)"; yh["AH18"] = "=SUM(AH6:AH17)"; yh["AI18"] = '=IF(AH18=0,"",MIN(1,AG18/(AH18*$W$3/7)))'
    style_range(yh, 6, 18, 1, 35, fmt={**{c: GRID for c in range(2, 33)}, 35: PCT})
    for r in range(6, 18):
        for d in range(1, 32): yh[f"{L(1+d)}{r}"].alignment = Alignment(horizontal="center")
    yh.conditional_formatting.add("B6:AF17", FormulaRule(formula=["B6=1"], fill=fill(DONE), font=Font(name=F, bold=True, color="FFFFFF")))
    yh.conditional_formatting.add("B6:AF17", FormulaRule(formula=['AND(B6<>"",B6=0)'], fill=fill(MISS), font=Font(name=F, color="9CA3AF")))
    yh.conditional_formatting.add("AI6:AI18", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=1, end_color=DONE))
    yh["A20"] = "Longest streak:"; yh["E20"] = f'=IF(Q3=0,"",INDEX(Habits!$O${H0}:$O${H1},Q3))'
    yh["A21"] = "Current streak:"; yh["E21"] = f'=IF(Q3=0,"",INDEX(Habits!$N${H0}:$N${H1},Q3))'
    for a in ("A20", "A21", "E20", "E21"): yh[a].font = font(bold=True)
    yh.freeze_panes = "B6"

    # ---------- Dashboard ----------
    K.title(dash, "Habit Tracker Dashboard", "Updates by itself. Add habits on the Habits tab and tick them on the Daily Log.")
    dash["A3"] = f'="Year "&{Y}&"  ·  as of "&TEXT({AS},"dd mmm yyyy")&"  ·  "&Settings!B14&" active habits"'; dash["A3"].font = font(bold=True)
    idx = "Settings!$B$13"
    today_done = f"IF(OR({idx}<1,{idx}>Settings!$B$11),0,INDEX({DL}!$R${D0}:$R${D1},{idx}))"
    wk_goal = f"SUM(Habits!$J${H0}:$J${H1})"
    kp = [("Done today", f'={today_done}&" of "&Settings!B14', "@"),
          ("This week vs weekly goal", f'=IF({wk_goal}=0,"",MIN(1,SUM(Habits!$P${H0}:$P${H1})/{wk_goal}))', PCT),
          ("This month score", '=IF(Settings!B15=0,"",INDEX(Calc!$B$57:$M$57,Settings!B15))', PCT),
          ("Year score so far", f'=IF(SUM(Habits!$L${H0}:$L${H1})=0,"",MIN(1,SUM(Habits!$K${H0}:$K${H1})/SUM(Habits!$L${H0}:$L${H1})))', PCT),
          ("Best current streak (days)", f"=MAX(Habits!$N${H0}:$N${H1})", "0"),
          ("Longest streak this year", f"=MAX(Habits!$O${H0}:$O${H1})", "0")]
    for i, (lab, f, fmt) in enumerate(kp):
        K.kpi(dash, 5 + (i // 3) * 3, 1 + (i % 3) * 2, lab, f, fmt)
        dash.cell(row=6 + (i // 3) * 3, column=1 + (i % 3) * 2).alignment = Alignment(horizontal="left")
    best = f'=IF(MAX(Habits!$N${H0}:$N${H1})=0,"",INDEX(Habits!$B${H0}:$B${H1},MATCH(MAX(Habits!$N${H0}:$N${H1}),Habits!$N${H0}:$N${H1},0)))'
    longest = f'=IF(MAX(Habits!$O${H0}:$O${H1})=0,"",INDEX(Habits!$B${H0}:$B${H1},MATCH(MAX(Habits!$O${H0}:$O${H1}),Habits!$O${H0}:$O${H1},0)))'
    dash["C10"] = best; dash["E10"] = longest
    for a in ("C10", "E10"): dash[a].font = font(italic=True, size=9, color="5B6770")
    widths = [28, 12, 11, 11, 11, 11, 11, 11, 11, 11, 11, 11, 11, 11]
    for i, w in enumerate(widths): dash.column_dimensions[L(1 + i)].width = w
    K.section(dash, "A11", "Your habits")
    K.hdr(dash, 12, ["Habit", "Category", "Target / week", "Done this week", "Weekly goal met", "Month score", "Year score",
                     "Current streak", "Longest streak", "Total days done"])
    for k in range(NH):
        r = 13 + k; hr = H0 + k
        dash[f"A{r}"] = f'=IF(Habits!$B${hr}<>"",Habits!$B${hr},"")'
        refs = {"B": "C", "C": "J", "D": "P", "F": "R", "G": "M", "H": "N", "I": "O", "J": "K"}
        for dc, hc in refs.items():
            dash[f"{dc}{r}"] = f'=IF($A{r}="","",Habits!${hc}${hr})'
        dash[f"E{r}"] = f'=IF($A{r}="","",IF(D{r}>=C{r},"Yes ✓",C{r}-D{r}&" to go"))'
    style_range(dash, 13, 12 + NH, 1, 10, fmt={6: PCT, 7: PCT})
    for r in range(13, 13 + NH):
        for c in "CDEFGHIJ": dash[f"{c}{r}"].alignment = Alignment(horizontal="center")
    dash.conditional_formatting.add(f"F13:G{12+NH}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=1, end_color=DONE))
    dash.conditional_formatting.add(f"E13:E{12+NH}", FormulaRule(formula=['LEFT(E13,3)="Yes"'], font=Font(name=F, bold=True, color="15803D")))
    K.section(dash, "A30", "Monthly heat-map (score = days done ÷ goal days)")
    K.hdr(dash, 31, ["Habit"] + [m[:3] for m in MONTHS] + ["Year"])
    for k in range(NH):
        r = 32 + k
        dash[f"A{r}"] = f"=Calc!A{42+k}"
        for m in range(13): dash[f"{L(2+m)}{r}"] = f"=Calc!{L(2+m)}{42+k}"
    dash["A47"] = "All habits"
    for m in range(13): dash[f"{L(2+m)}47"] = f"=Calc!{L(2+m)}57"
    style_range(dash, 32, 47, 1, 14, fmt={c: PCT for c in range(2, 15)})
    for r in range(32, 48):
        for m in range(13): dash[f"{L(2+m)}{r}"].alignment = Alignment(horizontal="center")
    for c in range(1, 15): dash.cell(row=47, column=c).font = font(bold=True)
    dash.conditional_formatting.add("B32:N47", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                              mid_type="num", mid_value=0.6, mid_color="DDD6FE", end_type="num", end_value=1, end_color="7C3AED"))
    ch = BarChart(); ch.type = "col"; ch.title = "All habits – score by month"; ch.height = 7.5; ch.width = 16; ch.legend = None
    ch.add_data(Reference(cc, min_col=2, min_row=59, max_row=71), titles_from_data=True)
    ch.set_categories(Reference(cc, min_col=1, min_row=60, max_row=71))
    ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1; ch.y_axis.number_format = "0%"; ch.y_axis.majorGridlines = None
    ch.series[0].graphicalProperties.solidFill = MAIN
    dash.add_chart(ch, "A50")
    ch2 = BarChart(); ch2.type = "bar"; ch2.title = "Year score by habit"; ch2.height = 7.5; ch2.width = 14; ch2.legend = None
    ch2.add_data(Reference(cc, min_col=14, min_row=41, max_row=41 + NH), titles_from_data=True)
    ch2.set_categories(Reference(cc, min_col=1, min_row=42, max_row=41 + NH))
    ch2.x_axis.scaling.orientation = "maxMin"; ch2.y_axis.scaling.min = 0; ch2.y_axis.scaling.max = 1; ch2.y_axis.number_format = "0%"
    ch2.y_axis.majorGridlines = None; ch2.series[0].graphicalProperties.solidFill = DONE
    dash.add_chart(ch2, "H50")

    # ---------- Start Here ----------
    lines = [("How to start", True),
             ("1.  Habits tab: type up to 15 habits, choose a category and set how many days per week you aim for (7 = every day).", False),
             ("2.  Daily Log tab: each day, type x in the column of every habit you did. Today's row is highlighted in yellow.", False),
             ("3.  Dashboard: today's progress, this week vs your weekly goal, month and year scores, streaks and a monthly heat-map.", False),
             ("4.  Monthly View: pick a month to see a ✓ / · grid of every habit, day by day.", False),
             ("5.  Year Heatmap: pick one habit to see its whole year as a calendar heat-map.", False),
             ("", False), ("How the score works", True),
             ("Score = days done ÷ goal days. Goal days = days counted × target per week ÷ 7. So a habit you aim to do 4 times a week reaches 100% at 4 days a week.", False),
             ("Streak = days in a row. Your current streak doesn't break until the day is over, so today counts even if you haven't ticked it yet.", False),
             ("", False), ("Colour key", True),
             ("Yellow cells are inputs: type here.  Grey cells calculate themselves: please don't overwrite them.  Purple = done.", False),
             ("", False), ("Good to know", True),
             ("•  Adding a habit later in the year? Type its start date on the Habits tab so earlier days don't lower its score.", False),
             ("•  New year: save a copy of the file, clear the x's on the Daily Log and change the year on the Settings tab.", False),
             ("•  To use Google Sheets: open drive.google.com, upload this file, then choose Open with > Google Sheets. Works great on your phone too.", False)]
    if sample:
        lines += [("", False), ("This copy contains sample data", True),
                  ("The habits and ticks are made-up examples for 2026, shown as of 30 Oct 2026. Use the blank copy in your download for your own habits.", False)]
    K.start(st, "Habit Tracker", "Track up to 15 habits every day, see streaks, scores and a heat-map of your year. Works in Excel and Google Sheets.", lines)
    finish(wb, sample, path)


META = {
    "name": "Habit Tracker with Yearly & Monthly Heat-Map – Excel & Google Sheets",
    "file_base": "Habit-Tracker-Heatmap",
    "price": 12,
    "color": "#7C3AED",
    "title": "Habit Tracker & Heat-Map",
    "tagline": "Daily ticks · streaks · monthly & yearly heat-map — Excel & Google Sheets",
    "intro": "Build habits that stick. Tick off up to 15 habits each day and the tracker shows your streaks, weekly goals, monthly scores and a colour heat-map of your whole year – automatically.",
    "features": [
        "Dashboard: done today, this week vs weekly goal, month score, year score, best and longest streaks",
        "Monthly heat-map for every habit, January to December, plus charts of your progress",
        "Daily Log for the whole year: just type x – today's row is highlighted",
        "Monthly View: a ✓ / · grid of all habits for any month you pick",
        "Year Heatmap: one habit, every day of the year, as a calendar heat-map",
        "Weekly targets (e.g. gym 4x a week), optional start dates, categories and notes",
        "Up to 15 habits, works for any year, phone-friendly in Google Sheets",
    ],
    "tags": ["habit tracker", "habit tracker spreadsheet", "google sheets", "excel template", "daily habit tracker", "habit heatmap",
             "streak tracker", "goal tracker", "self care tracker", "routine tracker", "2027 habit tracker", "productivity", "digital planner"],
}
