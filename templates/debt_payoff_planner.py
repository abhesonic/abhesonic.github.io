"""Debt Snowball & Avalanche Payoff Planner (sample + blank versions).

Month-by-month payoff simulation in plain formulas (no macros):
every month each debt gets interest, then its minimum payment; whatever is left of the monthly budget
(minimums of all debts + extra + one-time payments) goes to the debts in priority order
(snowball = smallest balance first, avalanche = highest APR first, or a custom order).
Minimums of paid-off debts roll into the budget automatically because the budget stays the same.
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter as L
from _common import Kit, F, INK, B, MONEY, DATE, font, fill, style_range, input_cell, finish

MAIN = "15803D"; SOFT = "ECFDF3"
K = Kit(MAIN, SOFT)
ND = 10          # debts
MONTHS = 360     # 30 years
LUMPS = 24       # one-time extra payments
FIRST = 12; LAST = FIRST + MONTHS - 1   # plan rows for months 1..360 (row 11 = start)
METHODS = ["Snowball – smallest balance first", "Avalanche – highest interest first", "Custom order (set on the Debts tab)"]
PLANS = [("Snowball Plan", "K"), ("Avalanche Plan", "L"), ("Custom Plan", "M")]   # sheet, rank column on Debts
MON = 'mmm yyyy'

def cols(start, n=ND):
    i = start
    return [L(i + k) for k in range(n)]
REM = cols(4)            # D..M  remaining after interest + minimum
END = cols(15)           # O..X  end-of-month balance
PAY = cols(28)           # AB..AK payment per debt


def q(name):
    return f"'{name}'"


def build_plan(ws, title):
    sheet, rank_col = title
    K.title(ws, sheet, "Month-by-month simulation. Everything here is calculated – nothing to type.")
    for c, w in zip("ABC", [26, 14, 16]): ws.column_dimensions[c].width = w
    for c in REM + END + PAY: ws.column_dimensions[c].width = 12
    ws.column_dimensions["N"].width = 13
    for c in ("Y", "Z", "AA"): ws.column_dimensions[c].width = 13
    # summary
    labels = [("Debt-free after (months)", f'=IF(AA11<=0.005,0,IF(COUNTIF(AA{FIRST}:AA{LAST},">0.005")>={MONTHS},"Over 30 yrs",COUNTIF(AA{FIRST}:AA{LAST},">0.005")+1))', "0"),
              ("Debt-free month", '=IF(ISNUMBER(B5),IF(B5=0,"",EDATE(Settings!$B$5,B5-1)),"")', MON),
              ("Total interest", f"=SUM(Y{FIRST}:Y{LAST})", MONEY),
              ("Total paid", f"=SUM(Z{FIRST}:Z{LAST})", MONEY)]
    for i, (lab, f, fmt) in enumerate(labels, 5):
        ws.cell(row=i, column=1, value=lab).font = font(bold=True)
        c = ws.cell(row=i, column=2, value=f); c.number_format = fmt; c.font = Font(name=F, bold=True, color=MAIN); c.border = B
    for r, lab in ((5, "Payoff order"), (6, "Monthly rate"), (7, "Minimum payment")):
        ws.cell(row=r, column=3, value=lab).font = font(italic=True, size=9)
    ws["N8"] = "Paid off in month #"; ws["N9"] = "Interest paid"
    for a in ("N8", "N9"): ws[a].font = font(italic=True, size=9)
    for k in range(ND):
        d = 5 + k   # Debts row
        name = f'=IF(Debts!$G${d}=1,Debts!$A${d},"")'
        for col in (REM[k], END[k], PAY[k]):
            ws[f"{col}4"] = name; ws[f"{col}4"].font = font(bold=True, size=9); ws[f"{col}4"].alignment = Alignment(wrap_text=True)
        ws[f"{REM[k]}5"] = f"=Debts!${rank_col}${d}"
        ws[f"{REM[k]}6"] = f"=Debts!$H${d}"; ws[f"{REM[k]}6"].number_format = "0.000%"
        ws[f"{REM[k]}7"] = f"=IF(Debts!$G${d}=1,N(Debts!$E${d}),0)"; ws[f"{REM[k]}7"].number_format = MONEY
        e = END[k]
        ws[f"{e}8"] = f'=IF(Debts!$G${d}=0,"",IF(COUNTIF({e}${FIRST}:{e}${LAST},">0.005")>={MONTHS},"Over 30 yrs",COUNTIF({e}${FIRST}:{e}${LAST},">0.005")+1))'
        ws[f"{e}9"] = f'=IF(Debts!$G${d}=0,"",{REM[k]}$6*({e}$11+SUM({e}${FIRST}:{e}${LAST-1})))'
        ws[f"{e}9"].number_format = MONEY
        for r in (8, 9): ws[f"{e}{r}"].font = font(bold=True, size=9)
    heads = ["Month #", "Month", "Budget this month"] + [f"After minimum – debt {k+1}" for k in range(ND)] + ["Extra to priority debts"] + \
            [f"Balance – debt {k+1}" for k in range(ND)] + ["Interest", "Paid this month", "Total balance"] + [f"Payment – debt {k+1}" for k in range(ND)]
    K.hdr(ws, 10, heads)
    ws["A11"] = 0; ws["B11"] = "Start"
    for k in range(ND):
        ws[f"{END[k]}11"] = f"=IF(Debts!$G${5+k}=1,N(Debts!$C${5+k}),0)"
    ws["AA11"] = f"=SUM({END[0]}11:{END[-1]}11)"
    for r in range(FIRST, LAST + 1):
        p = r - 1
        ws[f"A{r}"] = f"=A{p}+1"
        ws[f"B{r}"] = f"=EDATE(Settings!$B$5,A{r}-1)"
        ws[f"C{r}"] = f"=Settings!$B$10+SUMIFS(Settings!$B$15:$B${14+LUMPS},Settings!$D$15:$D${14+LUMPS},A{r})"
        for k in range(ND):
            ws[f"{REM[k]}{r}"] = f"=MAX(0,{END[k]}{p}*(1+{REM[k]}$6)-{REM[k]}$7)"
        ws[f"Y{r}"] = f"=SUMPRODUCT({END[0]}{p}:{END[-1]}{p},$D$6:$M$6)"
        ws[f"N{r}"] = f"=MAX(0,C{r}-(SUM({END[0]}{p}:{END[-1]}{p})+Y{r}-SUM(D{r}:M{r})))"
        for k in range(ND):
            dk = REM[k]
            ws[f"{END[k]}{r}"] = f'={dk}{r}-MIN({dk}{r},MAX(0,$N{r}-SUMIFS($D{r}:$M{r},$D$5:$M$5,"<"&{dk}$5)))'
            ws[f"{PAY[k]}{r}"] = f"={END[k]}{p}*(1+{dk}$6)-{END[k]}{r}"
        ws[f"Z{r}"] = f"=SUM({END[0]}{p}:{END[-1]}{p})+Y{r}-SUM({END[0]}{r}:{END[-1]}{r})"
        ws[f"AA{r}"] = f"=SUM({END[0]}{r}:{END[-1]}{r})"
    fmt = {2: MON, 3: MONEY}
    for c in range(4, 38): fmt[c] = MONEY
    style_range(ws, 11, LAST, 1, 37, fmt=fmt)
    ws.freeze_panes = "D12"


def build(sample, path):
    wb = Workbook()
    st = wb.active; st.title = "Start Here"
    dash = wb.create_sheet("Dashboard"); debts = wb.create_sheet("Debts"); sched = wb.create_sheet("Payment Schedule")
    plan_ws = [wb.create_sheet(p[0]) for p in PLANS]
    cd = wb.create_sheet("Chart Data"); s = wb.create_sheet("Settings")

    # ---------- Settings ----------
    K.title(s, "Settings", "Yellow cells are yours to edit.")
    rows = [("Currency label", "USD", None), ("First payment month", "=DATE(YEAR(TODAY()),MONTH(TODAY())+1,1)", MON),
            ("Payoff method", METHODS[0], None), ("Extra payment every month", 250 if sample else None, MONEY)]
    for i, (k_, v, fmt) in enumerate(rows, 4):
        s.cell(row=i, column=1, value=k_).font = font(bold=True)
        c = s.cell(row=i, column=2); input_cell(c, v)
        if fmt: c.number_format = fmt
    s["C5"] = "Type any date in the month of your first payment. It starts as next month."; s["C5"].font = font(italic=True, size=9)
    s["C7"] = "Added on top of all minimum payments, every month."; s["C7"].font = font(italic=True, size=9)
    calc = [("Plan number (1-3)", "=IFERROR(MATCH(B6,$H$5:$H$7,0),1)", "0"),
            ("Total minimum payments", "=SUMIFS(Debts!$E$5:$E$14,Debts!$G$5:$G$14,1)", MONEY),
            ("Monthly debt budget", "=B9+N(B7)", MONEY)]
    for i, (k_, f, fmt) in enumerate(calc, 8):
        s.cell(row=i, column=1, value=k_).font = font(italic=True)
        c = s.cell(row=i, column=2, value=f); c.number_format = fmt; c.border = B; c.fill = fill("F3F4F6"); c.font = font()
    s["H4"] = "Payoff methods"; s["H4"].font = font(bold=True, size=9)
    for i, m in enumerate(METHODS): s.cell(row=5 + i, column=8, value=m).font = font(size=9)
    s["J4"] = "Debt types"; s["J4"].font = font(bold=True, size=9)
    types = ["Credit card", "Store card", "Car loan", "Student loan", "Personal loan", "Medical bill", "Mortgage", "Line of credit", "Buy now pay later", "Other"]
    for i, t in enumerate(types): s.cell(row=5 + i, column=10, value=t).font = font(size=9)
    dv = DataValidation(type="list", formula1="=$H$5:$H$7", allow_blank=False); s.add_data_validation(dv); dv.add("B6")
    K.section(s, "A13", "One-time extra payments (tax refund, bonus, side income...)")
    K.hdr(s, 14, ["Date paid", "Amount", "Note", "Month #"])
    style_range(s, 15, 14 + LUMPS, 1, 4, inputs=(1, 2, 3), fmt={1: DATE, 2: MONEY})
    for r in range(15, 15 + LUMPS):
        s[f"D{r}"] = f'=IF(OR(A{r}="",N(B{r})=0),"",(YEAR(A{r})-YEAR($B$5))*12+MONTH(A{r})-MONTH($B$5)+1)'
    if sample:
        s["A15"] = dt.date(2027, 4, 15); s["B15"] = 1200; s["C15"] = "Tax refund"
        s["A16"] = dt.date(2027, 12, 20); s["B16"] = 800; s["C16"] = "Year-end bonus"
    for c, w in zip("ABCDEFGHIJ", [30, 16, 26, 10, 3, 3, 3, 34, 3, 18]): s.column_dimensions[c].width = w

    # ---------- Debts ----------
    K.title(debts, "Your Debts", "One row per debt. Type the APR as a percent (24.9%). Grey columns calculate themselves.")
    heads = ["Debt name", "Type", "Current balance", "Interest rate (APR)", "Minimum monthly payment", "Custom order (optional)",
             "Included", "Monthly rate", "APR key", "Custom key", "Snowball order", "Avalanche order", "Custom order", "Check",
             "Months on minimums only", "Interest on minimums only"]
    K.hdr(debts, 4, heads, [24, 15, 15, 12, 14, 12, 9, 10, 8, 9, 10, 10, 10, 30, 13, 15])
    style_range(debts, 5, 14, 1, 16, inputs=(1, 2, 3, 4, 5, 6), fmt={3: MONEY, 4: "0.00%", 5: MONEY, 8: "0.000%", 9: "0.00%", 16: MONEY})
    if sample:
        data = [("Visa credit card", "Credit card", 4200, 0.249, 126, None), ("Store card", "Store card", 850, 0.2999, 35, None),
                ("Car loan", "Car loan", 11500, 0.069, 325, None), ("Student loan", "Student loan", 18200, 0.055, 210, None),
                ("Personal loan", "Personal loan", 3100, 0.115, 145, None), ("Hospital bill", "Medical bill", 640, 0.0, 50, None)]
        for i, row in enumerate(data):
            for j, v in enumerate(row, 1):
                if v is not None: debts.cell(row=5 + i, column=j, value=v)
    for r in range(5, 15):
        debts[f"G{r}"] = f'=IF(AND(A{r}<>"",N(C{r})>0),1,0)'
        debts[f"H{r}"] = f"=IF(G{r}=1,N(D{r})/12,0)"
        debts[f"I{r}"] = f'=IF(G{r}=1,N(D{r}),"")'
        debts[f"J{r}"] = f'=IF(G{r}=1,IF(N(F{r})>0,N(F{r}),1000+ROW()),"")'
        debts[f"K{r}"] = f'=IF(G{r}=0,"",COUNTIFS($G$5:$G$14,1,$C$5:$C$14,"<"&C{r})+COUNTIFS($G$5:G{r},1,$C$5:C{r},C{r}))'
        debts[f"L{r}"] = f'=IF(G{r}=0,"",COUNTIFS($G$5:$G$14,1,$I$5:$I$14,">"&I{r})+COUNTIFS($G$5:G{r},1,$I$5:I{r},I{r}))'
        debts[f"M{r}"] = f'=IF(G{r}=0,"",COUNTIFS($G$5:$G$14,1,$J$5:$J$14,"<"&J{r})+COUNTIFS($G$5:G{r},1,$J$5:J{r},J{r}))'
        debts[f"N{r}"] = (f'=IF(G{r}=0,IF(A{r}<>"","Add the current balance",""),IF(N(D{r})>1,"APR looks too high – type it as a percent (24.9%)",'
                          f'IF(N(E{r})<=0,"Add the minimum payment",IF(N(E{r})<=C{r}*H{r},"Minimum doesn\'t cover the interest","OK"))))')
        debts[f"O{r}"] = (f'=IF(G{r}=0,"",IF(N(E{r})<=C{r}*H{r},"Never",IF(H{r}=0,ROUNDUP(C{r}/E{r}-1E-9,0),'
                          f'ROUNDUP(NPER(H{r},-E{r},C{r})-1E-9,0))))')
        debts[f"P{r}"] = f'=IF(OR(G{r}=0,O{r}="Never"),"",(O{r}-1)*E{r}+(-FV(H{r},O{r}-1,-E{r},C{r}))*(1+H{r})-C{r})'
    debts["A15"] = "Total"; debts["C15"] = "=SUM(C5:C14)"; debts["E15"] = "=SUMIFS(E5:E14,G5:G14,1)"
    debts["O15"] = '=IF(COUNTIF(O5:O14,"Never")>0,"Never",MAX(O5:O14))'; debts["P15"] = '=IF(O15="Never","",SUM(P5:P14))'
    for c in "ACEOP":
        x = debts[f"{c}15"]; x.font = font(bold=True); x.border = B; x.fill = fill(SOFT)
        if c in "CEP": x.number_format = MONEY
    for ref, f1 in (("B5:B14", "=Settings!$J$5:$J$14"),):
        dv = DataValidation(type="list", formula1=f1, allow_blank=True); debts.add_data_validation(dv); dv.add(ref)
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                        error="Type the APR as a percent, for example 24.9%", showErrorMessage=True)
    debts.add_data_validation(dv); dv.add("D5:D14")
    debts.conditional_formatting.add("N5:N14", FormulaRule(formula=['AND(N5<>"",N5<>"OK")'], fill=fill("FEE2E2"), font=Font(name=F, bold=True, color="B91C1C")))
    debts.conditional_formatting.add("N5:N14", CellIsRule(operator="equal", formula=['"OK"'], fill=fill("D1FAE5")))
    debts["A17"] = "Custom order: number your debts 1, 2, 3… in the order you want to attack them, then choose 'Custom order' on the Settings tab."
    debts["A18"] = "'Months / interest on minimums only' shows what happens if you pay only each minimum and never pay extra."
    for a in ("A17", "A18"): debts[a].font = font(italic=True, size=9)
    debts.freeze_panes = "B5"

    for ws, p in zip(plan_ws, PLANS): build_plan(ws, p)

    PN = "Settings!$B$8"
    def pick(ref):   # ref like "B5" or "O8" -> chosen plan's cell
        return f"IF({PN}=1,{q(PLANS[0][0])}!{ref},IF({PN}=2,{q(PLANS[1][0])}!{ref},{q(PLANS[2][0])}!{ref}))"

    # ---------- Payment Schedule (chosen plan) ----------
    K.title(sched, "Payment Schedule", "What to pay on each debt every month, following the method chosen on the Settings tab. Paid-off months are green.")
    sched["A3"] = '="Method: "&Settings!B6'; sched["A3"].font = font(bold=True)
    sched["P3"] = "Debt-free after (months):"; sched["P3"].font = font(italic=True, size=9)
    sched["Q3"] = "=" + pick("$B$5"); sched["Q3"].font = font(bold=True)
    sched["B4"] = "Paid off in month #"; sched["B4"].font = font(italic=True, size=9)
    K.hdr(sched, 5, ["Month #", "Month"] + [f"Debt {k+1}" for k in range(ND)] + ["Total payment", "Total balance left"],
          [9, 11] + [13] * ND + [14, 16])
    for k in range(ND):
        c = L(3 + k)
        sched[f"{c}5"] = f'=IF(Debts!$G${5+k}=1,Debts!$A${5+k},"Debt {k+1}")'
        sched[f"{c}4"] = "=" + pick(f"{END[k]}$8"); sched[f"{c}4"].font = font(bold=True, size=9); sched[f"{c}4"].alignment = Alignment(horizontal="right")
    for i in range(MONTHS):
        r = 6 + i; pr = FIRST + i
        sched[f"A{r}"] = i + 1
        sched[f"B{r}"] = f'=IF(A{r}>$Q$3,"",EDATE(Settings!$B$5,A{r}-1))'
        for k in range(ND):
            sched[f"{L(3+k)}{r}"] = f'=IF($A{r}>$Q$3,"",{pick(PAY[k] + str(pr))})'
        sched[f"M{r}"] = f'=IF($A{r}>$Q$3,"",SUM(C{r}:L{r}))'
        sched[f"N{r}"] = f'=IF($A{r}>$Q$3,"",{pick("AA" + str(pr))})'
    fmt = {2: MON}
    for c in range(3, 15): fmt[c] = MONEY
    style_range(sched, 6, 5 + MONTHS, 1, 14, fmt=fmt)
    sched.conditional_formatting.add(f"C6:L{5+MONTHS}", FormulaRule(formula=['AND(ISNUMBER(C$4),$A6=C$4)'], fill=fill("BBF7D0"), font=Font(name=F, bold=True, color="166534")))
    sched.freeze_panes = "C6"

    # ---------- Chart Data ----------
    K.title(cd, "Chart Data", "Balance every 6 months: paying minimums only vs your plan. Calculated – nothing to type.")
    K.hdr(cd, 4, ["Month #", "Label", "Minimum payments only", "Your plan"] + [f"Min-only debt {k+1}" for k in range(ND)], [9, 9, 16, 14] + [12] * ND)
    for i in range(21):
        r = 5 + i
        cd[f"A{r}"] = f"=(ROW()-5)*6"
        cd[f"B{r}"] = f'=IF(A{r}=0,"Now",A{r}/12&"y")'
        for k in range(ND):
            d = 5 + k
            cd[f"{L(5+k)}{r}"] = f"=IF(Debts!$G${d}=1,MAX(0,-FV(Debts!$H${d},$A{r},-N(Debts!$E${d}),Debts!$C${d})),0)"
        cd[f"C{r}"] = f"=SUM(E{r}:N{r})"
        cd[f"D{r}"] = "=" + "IF({pn}=1,INDEX({a}!$AA$11:$AA${last},A{r}+1),IF({pn}=2,INDEX({b}!$AA$11:$AA${last},A{r}+1),INDEX({c}!$AA$11:$AA${last},A{r}+1)))".format(
            pn=PN, a=q(PLANS[0][0]), b=q(PLANS[1][0]), c=q(PLANS[2][0]), last=LAST, r=r)
    style_range(cd, 5, 25, 1, 14, fmt={c: MONEY for c in range(3, 15)})

    # ---------- Dashboard ----------
    K.title(dash, "Debt Payoff Dashboard", "Updates by itself. Add your debts on the Debts tab and pick a method on the Settings tab.")
    dash["A3"] = '="Method: "&Settings!B6&"  ·  Amounts in "&Settings!B4'; dash["A3"].font = font(bold=True)
    dash["L1"] = f"={PN}"; dash["L1"].font = Font(name=F, size=8, color="FFFFFF")
    pr = "13+$L$1"
    kp = [("Total debt today", "=Debts!C15", MONEY),
          ("Monthly debt payment", "=Settings!B10", MONEY),
          ("Debt-free month", f'=IF(INDEX(C14:C16,$L$1)="",IF(INDEX(B14:B16,$L$1)=0,"No debt",INDEX(B14:B16,$L$1)),INDEX(C14:C16,$L$1))', MON),
          ("Months until debt-free", "=INDEX(B14:B16,$L$1)", "0"),
          ("Interest you will pay", "=INDEX(D14:D16,$L$1)", MONEY),
          ("Interest saved vs minimums", '=IF(INDEX(E14:E16,$L$1)="","—",INDEX(E14:E16,$L$1))', MONEY)]
    for i, (lab, f, fmt) in enumerate(kp):
        K.kpi(dash, 5 + (i // 3) * 3, 1 + (i % 3) * 2, lab, f, fmt)
    for c, w in zip("ABCDEFGHIJ", [30, 15, 15, 15, 17, 13, 12, 13, 3, 3]): dash.column_dimensions[c].width = w
    K.section(dash, "A11", "Compare the methods")
    K.hdr(dash, 12, ["Plan", "Debt-free in (months)", "Debt-free month", "Total interest", "Interest saved vs minimums"])
    dash["A13"] = "Minimum payments only"; dash["B13"] = "=Debts!O15"
    dash["C13"] = '=IF(ISNUMBER(B13),IF(B13=0,"",EDATE(Settings!$B$5,B13-1)),"")'; dash["D13"] = '=IF(Debts!P15="","",Debts!P15)'; dash["E13"] = "—"
    for i, (name, _) in enumerate(PLANS):
        r = 14 + i
        dash[f"A{r}"] = ["Snowball (smallest balance first)", "Avalanche (highest interest first)", "Custom order"][i]
        dash[f"B{r}"] = f"={q(name)}!B5"; dash[f"C{r}"] = f"={q(name)}!B6"; dash[f"D{r}"] = f"={q(name)}!B7"
        dash[f"E{r}"] = f'=IF(AND(ISNUMBER($D$13),ISNUMBER(D{r})),$D$13-D{r},"")'
    style_range(dash, 13, 16, 1, 5, fmt={3: MON, 4: MONEY, 5: MONEY})
    for r in range(13, 17): dash[f"B{r}"].alignment = Alignment(horizontal="right")
    dash.conditional_formatting.add("A14:E16", FormulaRule(formula=["ROW()=13+$L$1"], fill=fill("BBF7D0"), font=Font(name=F, bold=True, color="166534")))
    dash["A17"] = "Your chosen method is highlighted. Avalanche usually saves the most interest; snowball gives quicker wins."
    dash["A17"].font = font(italic=True, size=9)
    K.section(dash, "A19", "Your debts with the chosen method")
    K.hdr(dash, 20, ["Debt", "Balance", "APR", "Minimum", "Pay-off order", "Paid off in month #", "Paid off", "Interest paid"])
    for k in range(ND):
        r = 21 + k; d = 5 + k; e = END[k]
        dash[f"A{r}"] = f'=IF(Debts!$G${d}=1,Debts!$A${d},"")'
        dash[f"B{r}"] = f'=IF(A{r}="","",Debts!$C${d})'
        dash[f"C{r}"] = f'=IF(A{r}="","",N(Debts!$D${d}))'
        dash[f"D{r}"] = f'=IF(A{r}="","",N(Debts!$E${d}))'
        dash[f"E{r}"] = f'=IF(A{r}="","",IF($L$1=1,Debts!$K${d},IF($L$1=2,Debts!$L${d},Debts!$M${d})))'
        dash[f"F{r}"] = f'=IF(A{r}="","",{pick(e + "$8")})'
        dash[f"G{r}"] = f'=IF(ISNUMBER(F{r}),EDATE(Settings!$B$5,F{r}-1),"")'
        dash[f"H{r}"] = f'=IF(A{r}="","",{pick(e + "$9")})'
    style_range(dash, 21, 30, 1, 8, fmt={2: MONEY, 3: "0.00%", 4: MONEY, 7: MON, 8: MONEY})
    dash["A31"] = "Total"; dash["B31"] = "=SUM(B21:B30)"; dash["D31"] = "=SUM(D21:D30)"; dash["H31"] = "=SUM(H21:H30)"
    for c in "ABDH":
        x = dash[f"{c}31"]; x.font = font(bold=True); x.border = B; x.fill = fill(SOFT)
        if c != "A": x.number_format = MONEY
    for r in range(21, 31):
        for c in "EF": dash[f"{c}{r}"].alignment = Alignment(horizontal="center")
    ch = LineChart(); ch.title = "Debt balance over time"; ch.height = 8; ch.width = 18
    ch.add_data(Reference(cd, min_col=3, min_row=4, max_col=4, max_row=25), titles_from_data=True)
    ch.set_categories(Reference(cd, min_col=2, min_row=5, max_row=25))
    ch.y_axis.majorGridlines = None
    for ser, col in zip(ch.series, ("9CA3AF", MAIN)):
        ser.graphicalProperties.line.solidFill = col; ser.graphicalProperties.line.width = 28000; ser.smooth = False
    dash.add_chart(ch, "A34")

    # ---------- Start Here ----------
    lines = [("How to start", True),
             ("1.  Debts tab: type each debt – name, current balance, interest rate (APR, as 24.9%) and minimum monthly payment. Up to 10 debts.", False),
             ("2.  Settings tab: choose your first payment month, the payoff method (Snowball, Avalanche or Custom order) and how much extra you can pay each month.", False),
             ("3.  Optional: add one-time extra payments on the Settings tab (tax refund, bonus). They go straight to your priority debt.", False),
             ("4.  Dashboard: your debt-free date, total interest, interest saved, and a side-by-side comparison of all methods.", False),
             ("5.  Payment Schedule: exactly what to pay on each debt every month. The month each debt is paid off turns green.", False),
             ("", False), ("How the plan works", True),
             ("Every month you pay the minimum on every debt. Your extra money goes to one priority debt. When a debt is paid off, its minimum payment rolls over to the next debt, so your total monthly payment stays the same and the payoff speeds up.", False),
             ("Snowball = smallest balance first (quick wins). Avalanche = highest interest rate first (least interest). Custom = your own order.", False),
             ("", False), ("Colour key", True),
             ("Yellow cells are inputs: type here.  Grey cells calculate themselves: please don't overwrite them.", False),
             ("", False), ("Good to know", True),
             ("•  Interest is calculated monthly (APR ÷ 12). Real statements can differ slightly because of daily interest, fees or promotional rates.", False),
             ("•  The plan runs for up to 30 years. If a minimum payment doesn't even cover the interest, the Debts tab warns you.", False),
             ("•  To use Google Sheets: open drive.google.com, upload this file, then choose Open with > Google Sheets.", False),
             ("•  This planner is a calculator to help you plan. It is not financial advice.", False)]
    if sample:
        lines += [("", False), ("This copy contains sample data", True),
                  ("The debts and payments are made-up examples. Use the blank copy in your download for your own plan.", False)]
    K.start(st, "Debt Payoff Planner", "Snowball, avalanche or your own order: see your debt-free date and what to pay each month. Works in Excel and Google Sheets.", lines)
    finish(wb, sample, path)


META = {
    "name": "Debt Snowball & Avalanche Payoff Planner – Excel & Google Sheets",
    "file_base": "Debt-Payoff-Planner",
    "price": 17,
    "color": "#15803D",
    "title": "Debt Snowball & Avalanche Planner",
    "tagline": "Debt-free date · interest saved · monthly payment plan — Excel & Google Sheets",
    "intro": "See exactly when you'll be debt-free. Enter your debts once and the planner builds a month-by-month payoff plan using the snowball, avalanche or your own order – and shows how much interest you save.",
    "features": [
        "Dashboard: total debt, monthly payment, debt-free month, total interest and interest saved vs paying minimums only",
        "Side-by-side comparison: minimum payments only vs snowball vs avalanche vs your custom order",
        "Payment schedule: what to pay on each debt every month, with each payoff month highlighted",
        "Rollover built in: when a debt is paid off, its payment moves to the next debt automatically",
        "Extra monthly payment plus one-time extra payments (tax refund, bonus) applied to your priority debt",
        "Per-debt results: payoff order, payoff month and interest paid; chart of your balance over time",
        "Up to 10 debts and 30 years, any currency, warnings when a minimum doesn't cover the interest",
    ],
    "tags": ["debt snowball", "debt avalanche", "debt payoff", "debt tracker", "debt free", "google sheets", "excel template",
             "payoff planner", "credit card payoff", "budget template", "debt calculator", "loan payoff", "personal finance"],
}
